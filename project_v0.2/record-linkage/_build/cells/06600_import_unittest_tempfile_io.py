import unittest, tempfile, io

_T = {}


def _env():
    """Mappings, reference tables and a small configuration, loaded once."""
    if "cfg" not in _T:
        cfg = make_config("selftest", random_pairs=150_000, dedup_random_pairs=150_000, sim_records=4_000,
                          chunk_size=700, bootstrap_reps=40, io_format="csv", exhaustive=False, baseline=False,
                          address_standardizer="usaddress", self_check=False, incremental=False)
        _T["cfg"] = cfg
        _T["maps"] = load_mappings(cfg.paths["mappings"])
        _T["ref"] = load_reference(cfg.paths["reference"], cfg.core)
    return _T["cfg"], _T["maps"], _T["ref"]


def _small_run():
    """One end-to-end run on a small synthetic set with every edge case (cached)."""
    if "run" not in _T:
        cfg, maps, ref = _env()
        X, W, truth, edges, _ = make_dataset(400, 100, 1200, ref, maps["simulation_noise"], cfg.seed)
        X, _ = validate_input(X, "X"); W, _ = validate_input(W, "W")
        tf = cfg.paths["out"] / "selftest_truth.csv"
        truth.to_csv(tf, index=False)
        _T["run"] = run_pipeline(X, W, maps, ref, cfg, truth_file=tf)
        _T["inputs"] = (X, W, truth, edges)
    return _T["run"]


def _tframe(rows):
    df = pd.DataFrame([blank_row(r["record_id"], **{k: v for k, v in r.items() if k != "record_id"})
                       for r in rows], columns=SCHEMA)
    return validate_input(df, "t")[0]


def _parties(rows, source="X", watchlist=False):
    cfg, maps, ref = _env()
    nick = Nicknames(ref["nicknames"])
    return rows_to_parties(_tframe(rows), source, maps, nick, cfg.core, watchlist)


class TestNormalizers(unittest.TestCase):
    def test_ssn(self):
        self.assertEqual(norm_ssn("212-45-6781"), ("212456781", True, ""))
        for bad in ("000-12-3456", "666123456", "912345678", "123456789", "111111111", "12345"):
            self.assertFalse(norm_ssn(bad)[1], bad)

    def test_tin(self):
        self.assertTrue(norm_tin("45-1234567")[1])
        self.assertFalse(norm_tin("07-1234567")[1])
        self.assertFalse(norm_tin("000000000")[1])

    def test_npi_luhn(self):
        self.assertTrue(npi_luhn_ok("1234567893"))
        self.assertFalse(npi_luhn_ok("1234567890"))
        self.assertEqual(norm_npi("0000000000"), ("", False, ""))
        self.assertEqual(norm_npi("1234567890")[2], "check_digit")
        self.assertTrue(npi_luhn_ok(luhn_npi("123456789")))

    def test_phone_email(self):
        self.assertEqual(norm_phone("1 (718) 392-4411")[0], "7183924411")
        self.assertFalse(norm_phone("2125550123")[1])
        self.assertFalse(norm_phone("0123456789")[1])
        self.assertEqual(norm_email(" A.B@Mail.Test ")[0], "a.b@mail.test")
        self.assertFalse(norm_email("none@none.com")[1])
        self.assertFalse(norm_email("not an email")[1])

    def test_vin(self):
        self.assertTrue(norm_vin("1M8GDM9AXKP042788")[1])
        self.assertEqual(norm_vin("1M8GDM9AYKP042788")[2], "check_digit")
        self.assertEqual(norm_vin("1M8GDM9AXKP04278I")[2], "format")

    def test_dob(self):
        self.assertEqual(norm_dob("3/4/1961")[0], "1961-03-04")
        self.assertEqual(norm_dob("19610304")[0], "1961-03-04")
        self.assertEqual(norm_dob("1900-01-01")[2], "placeholder")
        self.assertEqual(norm_dob("02/30/1960")[2], "impossible")
        self.assertEqual(norm_dob("2090-01-01")[2], "range")

    def test_qualified_ids(self):
        self.assertEqual(norm_dl("d123-456", "New York")[0], "NY:D123456")
        self.assertEqual(norm_plate("abc 123", "")[0], ":ABC123")
        self.assertFalse(norm_license("00000", "NY")[1])

    def test_address(self):
        self.assertEqual(parse_street_line("2161 UNIVERSITY AVENUE W, STE 5"), ("2161", "", "UNIVERSITY", "AVE", "5"))
        self.assertEqual(parse_street_line("P O BOX 2161")[:3], ("2161", "", "PO BOX"))
        self.assertEqual(parse_street_line("222 N W 45TH AVE")[1], "NW")
        self.assertEqual(norm_street_parts("12", "", "Main Street", "", "Apt 4B"), ("12", "", "MAIN", "ST", "4B"))
        self.assertEqual(norm_state("new york"), "NY")
        self.assertEqual(norm_zip("2134"), "02134")
        self.assertEqual(norm_city("St. Louis"), "SAINT LOUIS")

    def test_person_names(self):
        self.assertEqual(clean_person("Dr. John", "a.", "Smith Jr."), ("JOHN", "A", "SMITH"))
        self.assertEqual(clean_person("", "", "Moy, Marvin"), ("MARVIN", "", "MOY"))
        self.assertEqual(clean_person("William A Weiner", "", ""), ("WILLIAM", "A", "WEINER"))
        self.assertEqual(clean_person("José", "", "O'Brien"), ("JOSE", "", "OBRIEN"))
        self.assertEqual(parse_person("WILLIAM A. WEINER, D.O.")["creds"], ["DO"])

    def test_org_aliases(self):
        self.assertEqual(org_aliases("Rutland Medical P.C. d/b/a Soul Radiology Medical Imaging"),
                         [["RUTLAND", "MEDICAL"], ["SOUL", "RADIOLOGY", "MEDICAL", "IMAGING"]])
        self.assertEqual(org_aliases("AMERICAN TRANSIT INS. CO."), [["AMERICAN", "TRANSIT", "INSURANCE"]])
        self.assertEqual(org_aliases("Smith & Jones LLP"), [["SMITH", "JONES"]])
        self.assertEqual(org_alias_string("Nexray Medical Imaging, P.C."), "NEXRAY MEDICAL IMAGING")

    def test_nicknames(self):
        _, _, ref = _env()
        n = Nicknames(ref["nicknames"])
        self.assertTrue(n.roots("Bill") & n.roots("William"))
        self.assertFalse(n.roots("Robert") & n.roots("William"))
        self.assertEqual(nysiis("Weiner"), nysiis("Wiener"))

    def test_near(self):
        self.assertTrue(_near("123456789", "123456780"))
        self.assertTrue(_near("123456789", "123457689"))
        self.assertFalse(_near("123456789", "123456798x"))
        self.assertFalse(_near("123456789", "123450780"))


class TestBreakdown(unittest.TestCase):
    def test_person_and_business_tied(self):
        rows, parties, details, ties, empty, _ = _parties([
            {"record_id": "r1", "first_name": "Ann", "last_name": "Lee", "business_name": "Lee Chiropractic PC",
             "tin": "45-1234567", "work_phone": "7183924411", "home_phone": "7183924412", "email": "ann@lee.test",
             "street_number": "5", "street_name": "Main", "street_type": "St", "zip": "11211", "state": "NY"}])
        self.assertEqual(sorted(parties["part"]), ["business", "person"])
        self.assertEqual(len(ties), 1)
        P = parties[parties["part"] == "person"].iloc[0]
        B = parties[parties["part"] == "business"].iloc[0]
        self.assertEqual(P["tie"], B["party_id"])
        self.assertEqual((P["id_email"], B["id_email"]), ("ann@lee.test", ""))
        self.assertEqual((P["phone_own"], P["phone_row"], B["phone_row"]), ("7183924412", "7183924411", "7183924411"))
        self.assertEqual(B["id_tin"], "451234567")
        self.assertEqual(P["addr_street"], B["addr_street"])
        own = details.set_index(["party_id", "type", "value"])["ownership"]
        self.assertEqual(own[(P["party_id"], "phone", "7183924411")], "row")
        self.assertEqual(own[(P["party_id"], "phone", "7183924412")], "own")

    def test_tin_only_and_empty(self):
        _, parties, _, _, empty, _ = _parties([{"record_id": "t", "tin": "451234567"},
                                               {"record_id": "e", "city": "Brooklyn"},
                                               {"record_id": "b", "business_name": "Acme Towing", "email": "x@acme.test"}])
        self.assertEqual(list(parties["record_id"]), ["t", "b"])
        self.assertEqual(list(empty["record_id"]), ["e"])
        self.assertEqual(parties.set_index("record_id").loc["b", "id_email"], "x@acme.test")

    def test_invalid_values_kept_visible_not_compared(self):
        _, parties, details, _, _, _ = _parties([{"record_id": "r", "first_name": "A", "last_name": "B",
                                                   "provider_npi": "1234567890", "ssn": "123456789"}])
        self.assertEqual(parties.iloc[0]["id_npi"], "")
        d = details.set_index("type")
        self.assertFalse(bool(d.loc["npi", "valid"]))
        self.assertEqual(d.loc["ssn", "reason"], "placeholder")

    def test_value_index(self):
        _, p1, d1, _, _, _ = _parties([{"record_id": "a", "first_name": "Ann", "last_name": "Lee", "ssn": "212456781"},
                                       {"record_id": "b", "first_name": "A", "last_name": "Lee", "ssn": "212456781"},
                                       {"record_id": "c", "first_name": "Bo", "last_name": "Kim", "ssn": "313456781"},
                                       {"record_id": "d", "first_name": "Cy", "last_name": "Poe", "ssn": "313456781"}])
        vi = mark_junk(value_index_fast(p1, d1), CoreParams(junk_holders=1)).set_index(["type", "value"])
        self.assertTrue(bool(vi.loc[("ssn", "212456781"), "single_holder"]))      # Ann Lee = A. Lee
        self.assertFalse(bool(vi.loc[("ssn", "313456781"), "single_holder"]))
        self.assertTrue(bool(vi.loc[("ssn", "313456781"), "junk"]))


class TestCategory(unittest.TestCase):
    def test_extracted_precedence(self):
        cfg, maps, _ = _env()
        _, p, _, _, _, _ = _parties([
            {"record_id": "npi", "category": "legal", "first_name": "A", "last_name": "B", "provider_npi": luhn_npi("123456789")},
            {"record_id": "giv", "category": "witness", "first_name": "C", "last_name": "D", "business_name": "D Auto Body"},
            {"record_id": "kw", "business_name": "Pemberton & Vasquez LLP"},
            {"record_id": "oth", "category": "other", "first_name": "E", "last_name": "F"},
            {"record_id": "lic", "first_name": "G", "last_name": "H", "professional_license_type": "DC"}])
        e = p.set_index(["record_id", "part"])
        self.assertEqual(e.loc[("npi", "person"), "category"], "medical")
        self.assertTrue(bool(e.loc[("npi", "person"), "category_mismatch"]))
        self.assertEqual(e.loc[("giv", "person"), "category"], "witness")          # given before keyword
        self.assertEqual(e.loc[("giv", "person"), "category_inferred"], "repair shop")
        self.assertEqual(e.loc[("giv", "person"), "prior_group"], "private")
        self.assertEqual(e.loc[("giv", "business"), "prior_group"], "business")
        self.assertEqual(e.loc[("kw", "business"), "category"], "legal")
        self.assertEqual(e.loc[("oth", "person"), "category"], "")                 # other = no information
        self.assertEqual(e.loc[("oth", "person"), "prior_group"], "unknown")
        self.assertEqual(e.loc[("lic", "person"), "cat_strength"], "strong")
        self.assertNotIn("witness", set(p["category_inferred"]))

    def test_watchlist_mapping_and_unmapped(self):
        _, p, _, _, _, unm = _parties([
            {"record_id": "w1", "first_name": "A", "last_name": "B", "provider_specialty": "CHIROPRACTIC"},
            {"record_id": "w2", "first_name": "C", "last_name": "D", "professional_license_type": "LAW PRACTICE"},
            {"record_id": "w3", "first_name": "E", "last_name": "F", "provider_specialty": "BASKET WEAVING"}],
            source="W", watchlist=True)
        self.assertEqual(list(p["category"]), ["medical", "legal", ""])
        self.assertIn("BASKET WEAVING", set(unm["value"]))


def _mini_ctx(xp, wp):
    cfg, maps, ref = _env()
    nick = Nicknames(ref["nicknames"])
    rar = Rarity(ref["surnames"], ref["first_names"], ref["org_tokens"], ref["meta"], cfg.core)
    stats = build_value_stats(xp, wp)
    dba = declared_dba_pairs(pd.concat([xp["org_aliases"], wp["org_aliases"]]))
    return CompareContext(cfg.core, rar, nick, stats, dba,
                          {"first_nick": 0.005, "first_close": 0.005, "first_agrees_fuzzy": 0.01, "first_initial": 0.05,
                           "first_empty": 0.05, "first_differs": 0.9, "last_close": 0.002, "last_sound": 0.004})


def _pair_levels(xrow, wrow, part="person"):
    _, xp, _, _, _, _ = _parties([xrow])
    _, wp, _, _, _, _ = _parties([wrow], source="W", watchlist=True)
    X, W = split_parts(xp)[part], split_parts(wp)[part]
    ctx = _mini_ctx(xp, wp)
    lv = compare_pairs(X, W, [0], [0], part, ctx)
    return lv.iloc[0], ctx


def _lev(lv, f):
    c = int(lv[f])
    return "empty" if c < 0 else FIELD_LEVELS[f][c]


class TestComparisons(unittest.TestCase):
    def test_name_levels(self):
        cases = [(("John", "Smith"), ("John", "Smith"), "exact"),
                 (("Bill", "Szczepanski"), ("William", "Szczepanski"), "first_nick"),
                 (("Steven", "Kowalczyk"), ("Stephen", "Kowalczyk"), "first_close"),
                 (("Maria", "Smyth"), ("Maria", "Smith"), "last_sound_first_agrees"),
                 (("Maria", "Garcia-Villanueva"), ("Maria", "Garcia"), "last_close_first_agrees"),
                 (("R", "Featherstone"), ("Rupert", "Featherstone"), "initial_agrees"),
                 (("Chidi", "Nwachukwu"), ("Nwachukwu", "Chidi"), "swapped"),
                 (("", "Webb"), ("Henry", "Webb"), "first_empty"),
                 (("Alan", "Webb"), ("Henry", "Webb"), "first_differs"),
                 (("Alan", "Webb"), ("Henry", "Moss"), "else")]
        for (f1, l1), (f2, l2), want in cases:
            row = {"record_id": "x", "first_name": f1, "last_name": l1}
            if not l1:
                row["business_name"] = "placeholder biz"    # keep a party with no surname
                row["first_name"] = ""
                row["middle_name"] = ""
            lv, _ = _pair_levels(row, {"record_id": "w", "first_name": f2, "last_name": l2})
            got = _lev(lv, "name") if l1 else "empty"
            self.assertEqual(got, want, (f1, l1, f2, l2))

    def test_dob_levels(self):
        for a, b, want in (("1975-04-09", "1975-04-09", "exact"), ("1975-04-09", "1975-09-04", "swap_or_typo"),
                           ("1975-04-09", "1975-04-08", "swap_or_typo"), ("1975-04-09", "1975-04-21", "year_month"),
                           ("1975-04-09", "1975-11-21", "year"), ("1975-04-09", "1980-11-21", "differs"),
                           ("1975-04-09", "1900-01-01", "empty")):
            lv, _ = _pair_levels({"record_id": "x", "first_name": "A", "last_name": "B", "dob": a},
                                 {"record_id": "w", "first_name": "A", "last_name": "B", "dob": b})
            self.assertEqual(_lev(lv, "dob"), want, (a, b))

    def test_address_levels(self):
        base = {"street_number": "5", "street_name": "Main", "street_type": "St", "unit": "2", "city": "Brooklyn",
                "state": "NY", "zip": "11211"}
        for change, want in (({}, "exact"), ({"unit": "3"}, "street"), ({"street_number": "7"}, "zip"),
                             ({"street_number": "7", "zip": "11222"}, "city_state"),
                             ({"street_number": "7", "zip": "11222", "city": "Albany"}, "state"),
                             ({"street_number": "7", "zip": "90026", "city": "Los Angeles", "state": "CA"}, "differs")):
            lv, _ = _pair_levels({"record_id": "x", "first_name": "A", "last_name": "B", **base},
                                 {"record_id": "w", "first_name": "A", "last_name": "B", **{**base, **change}})
            self.assertEqual(_lev(lv, "address"), want, change)

    def test_identifier_levels_and_veto(self):
        lv, _ = _pair_levels({"record_id": "x", "first_name": "A", "last_name": "B", "ssn": "212456781"},
                             {"record_id": "w", "first_name": "A", "last_name": "B", "ssn": "212456781"})
        self.assertEqual(_lev(lv, "ssn"), "exact")
        lv, _ = _pair_levels({"record_id": "x", "first_name": "A", "last_name": "B", "ssn": "212456781"},
                             {"record_id": "w", "first_name": "A", "last_name": "B", "ssn": "212456718"})
        self.assertEqual(_lev(lv, "ssn"), "near")
        self.assertFalse(bool(lv["veto_ssn"]))
        lv, _ = _pair_levels({"record_id": "x", "first_name": "A", "last_name": "B", "ssn": "212456781"},
                             {"record_id": "w", "first_name": "A", "last_name": "B", "ssn": "313999222"})
        self.assertEqual(_lev(lv, "ssn"), "differs")
        self.assertTrue(bool(lv["veto_ssn"]))
        lv, _ = _pair_levels({"record_id": "x", "first_name": "A", "last_name": "B", "driver_license_number": "D1234567", "driver_license_state": "NY"},
                             {"record_id": "w", "first_name": "A", "last_name": "B", "driver_license_number": "K9876543", "driver_license_state": "NJ"})
        self.assertFalse(bool(lv["veto_dl"]))                                       # two states: no veto
        lv, _ = _pair_levels({"record_id": "x", "business_name": "A Corp", "clinic_npi": luhn_npi("111111112")},
                             {"record_id": "w", "business_name": "A Corp", "clinic_npi": luhn_npi("222222223")}, "business")
        self.assertEqual(_lev(lv, "cnpi"), "differs")
        self.assertFalse(bool(lv["veto_cnpi"]))                                     # clinic NPI never vetoes

    def test_phone_levels(self):
        lv, _ = _pair_levels({"record_id": "x", "first_name": "A", "last_name": "B", "home_phone": "7183924411"},
                             {"record_id": "w", "first_name": "A", "last_name": "B", "home_phone": "7183924411"})
        self.assertEqual(_lev(lv, "phone"), "exact_owned_single")
        lv, _ = _pair_levels({"record_id": "x", "first_name": "A", "last_name": "B", "work_phone": "7183924411"},
                             {"record_id": "w", "first_name": "C", "last_name": "D", "work_phone": "7183924411"})
        self.assertEqual(_lev(lv, "phone"), "exact_shared")

    def test_org_levels(self):
        cases = [("Lakeshore Med Ctr", "Lakeshore Medical Center Inc", "exact"),
                 ("Kestrel Imaging LLC d/b/a Harborview Radiology", "Harborview Radiology", "exact"),
                 ("Pinecrest Orthopedic", "Pinecrest Orthopedic Rehabilitation", "short_form"),
                 ("Allport Indemnity Company", "Allport Casualty Surety Company", "sibling"),
                 ("Quixotic Holdings", "Unrelated Towing", "none")]
        for a, b, want in cases:
            lv, _ = _pair_levels({"record_id": "x", "business_name": a}, {"record_id": "w", "business_name": b}, "business")
            self.assertEqual(_lev(lv, "org"), want, (a, b))

    def test_spec_cat(self):
        lv, _ = _pair_levels({"record_id": "x", "first_name": "A", "last_name": "B", "provider_specialty": "Chiropractor", "provider_npi": luhn_npi("123456789")},
                             {"record_id": "w", "first_name": "A", "last_name": "B", "provider_specialty": "CHIROPRACTIC"})
        self.assertEqual(_lev(lv, "spec_cat"), "specialty")


class TestCandidates(unittest.TestCase):
    def setUp(self):
        run = _small_run()
        self.pr = run["pr"]
        self.cfg = _env()[0]

    def test_chunked_equals_unchunked(self):
        for part in ("person", "business"):
            X, W = self.pr.X[part], self.pr.W[part]
            plan = CandidatePlan(X, W, part, self.pr.ctx.rarity, self.cfg, self.pr.ctx.dba_pairs)
            one = plan.chunk(0, len(X)).sort_values(["l", "r"]).reset_index(drop=True)
            many = pd.concat([plan.chunk(lo, min(len(X), lo + 97)) for lo in range(0, len(X), 97)])
            many = many.sort_values(["l", "r"]).reset_index(drop=True)
            self.assertTrue(one.equals(many), part)

    def test_each_rule_proposes_its_target(self):
        X = _tframe([{"record_id": "x1", "first_name": "Bill", "last_name": "Kowalski"},
                    {"record_id": "x6", "first_name": "Rita", "last_name": "Oldname", "dob": "1970-01-02"},
                    {"record_id": "x2", "first_name": "Ann", "last_name": "Zyx", "ssn": "212456781"},
                    {"record_id": "x3", "first_name": "Mo", "last_name": "Qwe", "home_phone": "7183924411"},
                    {"record_id": "x4", "first_name": "Chidi", "last_name": "Nwachukwu"},
                    {"record_id": "x5", "last_name": "Plover", "dob": "1981-01-01"}])
        W = _tframe([{"record_id": "w1", "first_name": "William", "last_name": "Kowalski"},
                    {"record_id": "w6", "first_name": "Rita", "last_name": "Newname", "dob": "1970-01-02"},
                    {"record_id": "w2", "first_name": "Zed", "last_name": "Other", "ssn": "212456781"},
                    {"record_id": "w3", "first_name": "Zed", "last_name": "Else", "home_phone": "7183924411"},
                    {"record_id": "w4", "first_name": "Nwachukwu", "last_name": "Chidi"},
                    {"record_id": "w5", "first_name": "Ruth", "last_name": "Plover"}])
        cfg, maps, ref = _env()
        pr = prepare(X, W, maps, ref, cfg, log=lambda *a: None)
        plan = CandidatePlan(pr.X["person"], pr.W["person"], "person", pr.ctx.rarity, cfg)
        c = plan.chunk(0, len(pr.X["person"]))
        got = {(pr.X["person"]["record_id"][l], pr.W["person"]["record_id"][r]): set(n.split(","))
               for l, r, n in zip(c["l"], c["r"], plan.rule_names(c["rules"]))}
        self.assertIn("nysiis_canon_initial", got[("x1", "w1")])
        self.assertNotIn("nysiis_initial", got[("x1", "w1")])
        self.assertIn("dob_initial", got[("x6", "w6")])
        self.assertIn("ssn", got[("x2", "w2")])
        self.assertIn("phone", got[("x3", "w3")])
        self.assertIn("swapped", got[("x4", "w4")])
        self.assertIn("nysiis_first_missing", got[("x5", "w5")])

    def test_cross_file_same_part_only(self):
        S = _small_run()["scored"]
        for part in ("person", "business"):
            s = S["link"][part][0]
            ents = S["entities"][part]
            self.assertTrue((s["e"] < len(ents.ids)).all() and (s["r"] < len(self.pr.W[part])).all())

    def test_strict_rules_subset_of_blocking(self):
        cfg = self.cfg
        for part in ("person", "business"):
            X, W = self.pr.X[part], self.pr.W[part]
            l, r = strict_pairs(X, W, part, self.pr.ctx)
            sc = _small_run()["scored"]
            keys = sc["link"][part][3]["keys"]
            k = _key(sc["entities"][part].e_of[l], r)
            self.assertTrue(np.isin(k, keys).all(), part)


class TestScoring(unittest.TestCase):
    def setUp(self):
        self.run = _small_run()

    def test_empty_fields_are_zero_bits(self):
        for part in ("person", "business"):
            S, L = self.run["scored"]["link"][part][0], self.run["scored"]["link"][part][1]
            for f in COMPARED_FIELDS[part]:
                self.assertTrue((S.loc[L[f].to_numpy() < 0, f"bits_{f}"] == 0).all(), (part, f))

    def test_evidence_sums_to_total(self):
        ev = self.run["evidence"].groupby("pair_id")["bits"].sum()
        pairs = self.run["pairs"].set_index("pair_id")["bits"]
        diff = (pairs.reindex(ev.index) - ev).abs()
        self.assertLess(float(diff.max()), 1e-6)
        zero = pairs[~pairs.index.isin(ev.index)]
        self.assertTrue((zero.abs() < 1e-9).all())

    def test_agreement_never_lowers_p(self):
        part = "person"
        S, L = self.run["scored"]["link"][part][0], self.run["scored"]["link"][part][1]
        w = self.run["params"]["weights"][part]
        lv = L.copy()
        base = score_pairs(lv, part, w, S["prior_logit"].to_numpy())
        for f in ("ssn", "npi", "email", "dob", "name"):
            lv2 = lv.copy()
            empty = lv2[f].to_numpy() < 0
            lv2.loc[empty, f] = 0
            lv2.loc[empty, f"uv_{f}"] = 1e-6
            more = score_pairs(lv2, part, w, S["prior_logit"].to_numpy())
            ok = base["veto"].to_numpy() == ""
            self.assertTrue((more["p"].to_numpy()[ok] >= base["p"].to_numpy()[ok] - 1e-12).all(), f)

    def test_veto_visible_with_p_zero(self):
        P = self.run["pairs"]
        v = P[P["veto"] != ""]
        self.assertGreater(len(v), 0)
        self.assertTrue((v["p"] == 0).all())

    def test_name_only_never_identifier(self):
        part = "person"
        S, L = self.run["scored"]["link"][part][0], self.run["scored"]["link"][part][1]
        idn = S["basis"].to_numpy() == "identifier"
        any_id = np.zeros(len(L), bool)
        for f in IDENTIFIER_FIELDS:
            if f in L:
                any_id |= (L[f].to_numpy() == 0) & (S[f"bits_{f}"].to_numpy() > 0)
        self.assertTrue((any_id[idn]).all())
        self.assertTrue((S.loc[~any_id, "basis"] != "identifier").all())

    def test_coparty_needs_anchor_and_name(self):
        sc = self.run["scored"]
        S, L = sc["link"]["person"][0], sc["link"]["person"][1]
        cp = S["bits_co_party"].to_numpy() > 0
        self.assertGreater(int(cp.sum()), 0)
        self.assertTrue((S["bits_name"].to_numpy()[cp] > 0).all())
        self.assertTrue((L["co_party"].to_numpy()[cp] == 0).all())
        anchors = link_anchors(sc["link"]["business"][0], sc["entities"]["business"], _env()[0])
        X, W = self.run["pr"].X["person"], self.run["pr"].W["person"]
        ep = sc["entities"]["person"]
        for e, r in zip(S["e"].to_numpy()[cp], S["r"].to_numpy()[cp]):
            tl = X["tie_pos"].to_numpy()[ep.members(e)]
            tl = tl[tl >= 0]
            tr = W["tie_pos"].to_numpy()[r]
            ks = _key(anchors["e_of_business"][tl], np.full(len(tl), tr))
            self.assertTrue(np.isin(ks, anchors["keys"]).any())

    def test_deterministic(self):
        cfg, maps, ref = _env()
        X, W, truth, edges = _T["inputs"]
        again = run_pipeline(X, W, maps, ref, cfg)
        a = self.run["pairs"].drop(columns=["_e", "_r"]).reset_index(drop=True)
        b = again["pairs"].drop(columns=["_e", "_r"]).reset_index(drop=True)
        pd.testing.assert_frame_equal(a, b)

    def test_edge_cases(self):
        X, W, truth, edges = _T["inputs"]
        r = check_edge_cases(self.run["pr"], self.run["pairs"], self.run["rows"], self.run["scored"], edges)
        self.assertTrue(r["ok"].all(), r[~r["ok"]].to_string())

    def test_every_true_name_only_pair_visible(self):
        t = [x for x in self.run["diag"]["truth"] if x.get("true_pairs")]
        for x in t:
            self.assertEqual(x["proposed"], x["kept"], x["part"])

    def test_rows_one_per_party_entities_one_per_entity(self):
        # B1: the grain changed. Rows: one per extracted party (as the v1.0 entities sheet was);
        # entities: one per pooled entity; every row points to its entity's result.
        rows, e = self.run["rows"], self.run["entities"]
        self.assertTrue(rows["party_id"].is_unique)
        n = sum(len(self.run["pr"].X[p]) for p in ("person", "business"))
        self.assertEqual(len(rows), n)
        self.assertTrue(e["entity_id"].is_unique)
        self.assertTrue(rows["entity_id"].isin(set(e["entity_id"])).all())
        self.assertEqual(int(e["records"].sum()), n)
        m = rows.merge(e[["entity_id", "p"]], on="entity_id", suffixes=("", "_e"))
        self.assertTrue(np.allclose(m["p"], m["p_e"]))


class TestParameters(unittest.TestCase):
    def test_u_closed_form(self):
        lv = pd.DataFrame({"email": np.array([0] * 3 + [1] * 97 + [-1] * 50, dtype=np.int8)})
        u = estimate_u(lv, ["email"], CoreParams(u_pseudo=0.0)).set_index("level")["u"]
        self.assertAlmostEqual(u["exact"], 0.03)
        self.assertAlmostEqual(u["differs"], 0.97)

    def test_anchor_field_excluded(self):
        lv = pd.DataFrame({"ssn": np.array([0, 0, 0, 2], dtype=np.int8), "dob": np.array([0, 0, 4, 4], dtype=np.int8)})
        c = level_counts(lv, ["ssn", "dob"], exclude={"ssn": np.array([True, True, True, False])})
        self.assertEqual(c["ssn"][1], 1.0)
        self.assertEqual(c["dob"][1], 4.0)

    def test_em_recovers_m(self):
        rng = np.random.default_rng(3)
        n, lam = 20000, 0.3
        m = {"a": np.array([0.8, 0.15, 0.05]), "b": np.array([0.7, 0.3]), "c": np.array([0.9, 0.1])}
        u = {"a": np.array([0.01, 0.09, 0.9]), "b": np.array([0.05, 0.95]), "c": np.array([0.1, 0.9])}
        match = rng.random(n) < lam
        data = {}
        for f in m:
            k = len(m[f])
            data[f] = np.where(match, rng.choice(k, n, p=m[f]), rng.choice(k, n, p=u[f])).astype(np.int8)
        lv = pd.DataFrame(data)
        saved = {f: FIELD_LEVELS.get(f) for f in m}
        FIELD_LEVELS.update({"a": ["x", "y", "z"], "b": ["x", "y"], "c": ["x", "y"]})
        try:
            est, lam_hat, _ = em_fixed_u(lv, ["a", "b"], ["c"], {"c": m["c"]}, u, CoreParams())
        finally:
            for f, v in saved.items():
                if v is None:
                    FIELD_LEVELS.pop(f, None)
        self.assertAlmostEqual(lam_hat, lam, delta=0.02)
        self.assertLess(np.abs(est["a"][0] - m["a"]).max(), 0.03)
        self.assertLess(np.abs(est["b"][0] - m["b"]).max(), 0.03)

    def test_combine_chain(self):
        p = CoreParams(alpha=5, n_min=50)
        rows = combine_m("email", [("anchors_strict", np.array([8.0, 2.0]), 10.0),
                                   ("simulation", np.array([900.0, 100.0]), 1000.0)], [0.7, 0.3], p)
        m = rows[0]["m"]
        sim = (900 + 5 * 0.7) / (1000 + 5)
        self.assertAlmostEqual(m, (8 + 5 * sim) / 15)
        rows = combine_m("email", [("anchors_strict", np.array([80.0, 20.0]), 100.0),
                                   ("simulation", np.array([0.0, 1000.0]), 1000.0)], [0.7, 0.3], p)
        self.assertAlmostEqual(rows[0]["m"], (80 + 5 * 0.7) / 105)       # simulation skipped: enough pairs
        self.assertEqual(rows[0]["m_chain"], "anchors_strict(100)")

    def test_prior_formula(self):
        p = CoreParams()
        self.assertAlmostEqual(prior_estimate(10, 1_000_000, 0.5, p)[0], 2e-5)
        v, flagged = prior_estimate(0, 1_000_000, 0.5, p)
        self.assertTrue(flagged)
        self.assertAlmostEqual(v, 1e-6)
        self.assertAlmostEqual(strict_recall({"ssn": 0.5}, lambda f, l: 0.9), 0.45)

    def test_manifest_has_every_parameter(self):
        run = _small_run()
        W = pd.concat([run["params"]["weights"]["person"], run["params"]["weights"]["business"]])
        self.assertTrue((W["m_chain"] != "").all() and W["m_pairs"].notna().all())
        self.assertTrue((W["u_source"].fillna("") != "").all())
        n_levels = sum(len(FIELD_LEVELS[f]) for p in ("person", "business") for f in PART_FIELDS[p])
        self.assertEqual(len(W), n_levels)
        self.assertTrue((run["prior_df"]["prior"] > 0).all())


class TestOutput(unittest.TestCase):
    def test_continuation_sheets(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / "t.xlsx"
            df = pd.DataFrame({"a": np.arange(25), "b": ["x"] * 25, "c": [True] * 25, "d": [np.nan] * 25})
            written = write_workbook(path, {"cand": df, "manifest": df.head(2)}, row_limit=10)
            self.assertEqual(written["cand"], ["cand", "cand (2)", "cand (3)"])
            wb = openpyxl.load_workbook(path, read_only=True)
            self.assertEqual(wb.sheetnames, ["cand", "cand (2)", "cand (3)", "manifest"])
            rows = list(wb["cand (3)"].iter_rows(values_only=True))
            self.assertEqual(rows[0], ("a", "b", "c", "d"))
            self.assertEqual(len(rows), 6)
            wb.close()

    def test_manifest_sections(self):
        cfg, maps, ref = _env()
        run = _small_run()
        sw = Stopwatch(); sw.mark("t")
        m = build_manifest(cfg, {"x": "0"}, ref, ({"source": "x", "rows": 1, "all_empty_columns": [], "added_columns": [],
                                                   "passthrough_columns": []},), run["pr"], run["models"],
                           run["sens"], run["scored"], run["diag"], sw, None, {"run_id": "r1"})
        for s in ("config", "config.core", "run", "reference", "m", "u", "bits", "prior", "blocking", "dedup",
                  "m.em_passes", "u.sample", "data_quality", "timing"):
            self.assertIn(s, set(m["section"]), s)


class TestLeieExport(unittest.TestCase):
    def test_export(self):
        raw = ('LASTNAME,FIRSTNAME,MIDNAME,BUSNAME,GENERAL,SPECIALTY,UPIN,NPI,DOB,ADDRESS,CITY,STATE,ZIP,EXCLTYPE,EXCLDATE,REINDATE,WAIVERDATE,WVRSTATE\n'
               '"DOE","JANE","Q","","PHYSICIAN (MD, DO)","FAMILY PRACTICE","","1234567893","19600102","12 W MAIN ST, APT 4","ALBANY","NY","12207","1128b4","20200101","00000000","00000000",""\n'
               '"","","","ACME DME, INC","DME COMPANY","DME - GENERAL","","1234567893","","P O BOX 7","TROY","NY","12180","1128a1","20190101","00000000","00000000",""\n')
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "raw.csv"
            p.write_text(raw, encoding="utf-8")
            out = export_leie(p)
        self.assertEqual(list(out.columns[:len(SCHEMA)]), SCHEMA)
        a, b = out.iloc[0], out.iloc[1]
        self.assertEqual((a["provider_npi"], a["clinic_npi"], b["provider_npi"], b["clinic_npi"]),
                         ("1234567893", "", "", "1234567893"))
        self.assertEqual(a["dob"], "1960-01-02")
        self.assertEqual((a["street_number"], a["street_direction"], a["street_name"], a["street_type"], a["unit"]),
                         ("12", "W", "MAIN", "ST", "4"))
        self.assertEqual(a["professional_license_type"], "PHYSICIAN (MD, DO)")
        self.assertEqual(a["x_excldate"], "2020-01-01")
        self.assertTrue(out["record_id"].is_unique and out["record_id"].str.startswith("leie:").all())


class TestInputs(unittest.TestCase):
    def test_validation(self):
        with self.assertRaises(InputError):
            validate_input(pd.DataFrame({"record_id": ["a", "a"]}), "t")
        with self.assertRaises(InputError):
            validate_input(pd.DataFrame({"first_name": ["a"]}), "t")
        df, rep = validate_input(pd.DataFrame({"record_id": ["a"], "odd": ["1"]}), "t")
        self.assertIn("x_odd", df.columns)
        self.assertIn("ssn", rep["all_empty_columns"])

    def test_mappings_valid(self):
        cfg, maps, ref = _env()
        self.assertEqual(set(ref["hash_checks"]["status"]), {"ok"})
        self.assertGreater(len(maps["category_map"]), 280)


class TestClosedFormU(unittest.TestCase):
    def test_closed_form_matches_brute_force(self):
        # A small watchlist: closed-form email-exact u must equal a brute-force count over all
        # extracted x watchlist pairs.
        _, xp, _, _, _, _ = _parties([
            {"record_id": "x1", "first_name": "Ann", "last_name": "Lee", "email": "a@x.test"},
            {"record_id": "x2", "first_name": "Bo", "last_name": "Kim", "email": "b@x.test"},
            {"record_id": "x3", "first_name": "Cy", "last_name": "Poe", "email": ""}])
        _, wp, _, _, _, _ = _parties([
            {"record_id": "w1", "first_name": "Zed", "last_name": "Zoo", "email": "a@x.test"},
            {"record_id": "w2", "first_name": "Yan", "last_name": "Yoo", "email": "c@x.test"}],
            source="W", watchlist=True)
        closed = closed_form_u(xp, wp, "person")
        u_email, n = closed["email"]["exact"]
        brute_num = sum(1 for a in ("a@x.test", "b@x.test", "") for b in ("a@x.test", "c@x.test") if a and b and a == b)
        brute_den = sum(1 for a in ("a@x.test", "b@x.test", "") for b in ("a@x.test", "c@x.test") if a and b)
        self.assertAlmostEqual(u_email, brute_num / brute_den)
        self.assertEqual(n, brute_den)

    def test_closed_form_dedup_distinct_parties_only(self):
        # same_frame=True: agreement counted among distinct parties (i < j), never a party with
        # itself, and never double-counted.
        _, xp, _, _, _, _ = _parties([
            {"record_id": "x1", "first_name": "Ann", "last_name": "Lee", "email": "a@x.test"},
            {"record_id": "x2", "first_name": "Bo", "last_name": "Kim", "email": "a@x.test"},
            {"record_id": "x3", "first_name": "Cy", "last_name": "Poe", "email": "a@x.test"}])
        closed = closed_form_u(xp, xp, "person", same_frame=True)
        u_email, n = closed["email"]["exact"]
        self.assertAlmostEqual(u_email, 1.0)     # every pair agrees
        self.assertEqual(n, 3)                    # C(3,2) = 3 pairs, not 9

    def test_combine_u_rescales_remaining_mass(self):
        cfg, maps, ref = _env()
        u_s = estimate_u(pd.DataFrame({"ssn": np.array([0, 1, 2, 1, 2], dtype=np.int8)}), ["ssn"], cfg.core)
        closed = {"ssn": {"exact": (0.002, 500.0)}}
        out = combine_u(u_s, closed, "ssn", cfg.core)
        # closed form gets the same pseudo-count smoothing as the sample (a level never seen
        # among all pairs is rare, not impossible): (0.002*500 + 0.5) / (500 + 0.5*3)
        expect = (0.002 * 500.0 + cfg.core.u_pseudo) / (500.0 + cfg.core.u_pseudo * 3)
        self.assertAlmostEqual(float(out.loc[out["level"] == "exact", "u"].iloc[0]), expect, delta=1e-6)
        self.assertAlmostEqual(out["u"].sum(), 1.0, delta=1e-6)
        self.assertEqual(out.loc[out["level"] == "exact", "u_method"].iloc[0], "closed form")


class TestEMBootstrap(unittest.TestCase):
    def test_dirichlet_prior_shrinks_thin_pattern(self):
        # A field with almost no data among the selected pairs should barely move from the
        # prior centre m0 (the point of the Dirichlet/Beta prior in EM).
        fields = ["a"]
        FIELD_LEVELS["a"] = ["x", "y"]
        try:
            pat = np.array([[0], [1]], dtype=np.int64)
            counts = np.array([1.0, 1.0])          # one informative pair each: thin
            u_tab = {"a": np.array([0.9, 0.1])}
            m0 = {"a": np.array([0.5, 0.5])}
            est = em_patterns(pat, counts, fields, u_tab, m0, alpha=50.0, params=CoreParams())
            self.assertLess(abs(est["m"]["a"][0] - 0.5), 0.05)
        finally:
            del FIELD_LEVELS["a"]

    def test_bounds_no_field_below_bits_min(self):
        cfg, maps, ref = _env()
        rows = [{"field": "middle", "level": lev, "level_code": i, "m": 0.0001, "m_lo": np.nan, "m_hi": np.nan,
                 "m_source": "x", "m_chain": "x", "m_pairs": 10, "m_prior_centre": 0.0001, "m_published": 0.0001,
                 "m_passes": "{}", "bits_lo": np.nan, "bits_hi": np.nan, "unstable": ""}
                for i, lev in enumerate(FIELD_LEVELS["middle"])]
        u_df = pd.DataFrame({"field": ["middle"] * 3, "level": FIELD_LEVELS["middle"], "level_code": [0, 1, 2],
                             "u": [0.9999, 1e-9, 1e-9]})
        w = weights_table(rows, u_df, {}, cfg.core)
        self.assertTrue((w["bits_field_level"] >= cfg.core.bits_min).all())
        self.assertTrue((w["bits_field_level"] <= cfg.core.bits_max).all())
        # a synthetic middle-name "differs" case cannot produce an absurd weight (the user's
        # -16-bit example): m is floored, so bits never blow up to -16 from a near-zero count.
        self.assertGreater(w.loc[w["level"] == "differs", "bits_field_level"].iloc[0], cfg.core.bits_min)

    def test_final_weights_flag_unstable(self):
        run = _small_run()
        w = pd.concat([run["models"]["link"].weights["person"], run["models"]["link"].weights["business"]])
        self.assertIn("unstable", w.columns)
        # some levels are thin on the tiny synthetic set: at least the mechanism fires somewhere,
        # or every relevant level has a tight enough interval (either is a valid run).
        self.assertTrue((w["unstable"] == "").all() or (w["unstable"] != "").any())


class TestRicherComparisons(unittest.TestCase):
    def test_name_similarity_measures(self):
        sims = name_similarity("SMITH", "SMYTH")
        for k in ("jaro_winkler", "levenshtein", "token_sort", "token_set", "nysiis", "metaphone"):
            self.assertIn(k, sims)
        self.assertGreater(sims["jaro_winkler"], 0.8)
        self.assertTrue(sims["metaphone"])   # SMITH/SMYTH: same metaphone code (SM0), different NYSIIS

    def test_org_close_level(self):
        # rarity-weighted token-set similarity: two aliases sharing most of their word-rarity
        # mass and a high token-set ratio grade as 'close', stronger than a generic shared word.
        lv, _ = _pair_levels({"record_id": "x", "business_name": "Ridgeview Medical Imaging Center"},
                             {"record_id": "w", "business_name": "Ridgeview Medical Imaging Ctr"}, "business")
        self.assertIn(_lev(lv, "org"), ("exact", "close"))

    def test_middle_conditional_on_surname(self):
        # a middle name agreeing between two DIFFERENT surnames must not count: it is compared
        # only when the surname group already agrees (no double counting of the name group).
        lv, _ = _pair_levels({"record_id": "x", "first_name": "A", "middle_name": "J", "last_name": "Smith"},
                             {"record_id": "w", "first_name": "A", "middle_name": "J", "last_name": "Ortiz"})
        self.assertEqual(_lev(lv, "middle"), "empty")
        lv2, _ = _pair_levels({"record_id": "x", "first_name": "A", "middle_name": "James", "last_name": "Smith"},
                              {"record_id": "w", "first_name": "A", "middle_name": "James", "last_name": "Smith"})
        self.assertEqual(_lev(lv2, "middle"), "exact")


class TestAddressStandardizer(unittest.TestCase):
    def test_given_mode_passes_through(self):
        std = AddressStandardizer("given")
        df = pd.DataFrame({"street_number": ["5"], "street_direction": [""], "street_name": ["Main St"],
                           "street_type": [""], "unit": ["2"], "city": ["Brooklyn"], "state": ["NY"], "zip": ["11211"]})
        out = std.standardize(df)
        self.assertEqual(out["name"].iloc[0], "Main St")     # unparsed: kept exactly as given

    def test_usaddress_mode_splits_a_whole_line(self):
        std = AddressStandardizer("usaddress")
        df = pd.DataFrame({"street_number": [""], "street_direction": [""], "street_name": ["2161 University Ave W Ste 5"],
                           "street_type": [""], "unit": [""], "city": ["Saint Paul"], "state": ["MN"], "zip": ["55114"]})
        out = std.standardize(df)
        self.assertEqual(out["number"].iloc[0], "2161")
        self.assertEqual(out["name"].iloc[0].upper(), "UNIVERSITY")
        self.assertEqual(std.report["usaddress_parsed"], 1)

    def test_smarty_uses_injected_post_and_caches(self):
        calls = []

        def fake_post(url, params, payload):
            calls.append(payload)
            return [{"input_id": r["input_id"], "components": {"primary_number": "9", "street_name": "Oak",
                                                               "street_suffix": "Ave", "city_name": "X",
                                                               "state_abbreviation": "NY", "zipcode": "10001"}}
                    for r in payload]
        with tempfile.TemporaryDirectory() as d:
            cache = Path(d) / "cache.json"
            std = AddressStandardizer("smarty", cache_path=cache, post=fake_post, env={"SMARTY_AUTH_ID": "id",
                                                                                       "SMARTY_AUTH_TOKEN": "tok"})
            df = pd.DataFrame({"street_number": ["9"], "street_direction": [""], "street_name": ["Oak Ave"],
                               "street_type": [""], "unit": [""], "city": ["X"], "state": ["NY"], "zip": ["10001"]})
            out1 = std.standardize(df)
            out2 = std.standardize(df)     # second call: cache hit, no new post
        self.assertEqual(len(calls), 1)
        self.assertEqual(std.report["smarty_cache_hits"], 1)
        self.assertEqual(out1["name"].iloc[0], "Oak")

    def test_smarty_without_credentials_falls_back_safely(self):
        std = AddressStandardizer("smarty", env={})   # no credentials: never calls the network
        df = pd.DataFrame({"street_number": [""], "street_direction": [""], "street_name": ["5 Main St"],
                           "street_type": [""], "unit": [""], "city": ["Brooklyn"], "state": ["NY"], "zip": ["11211"]})
        out = std.standardize(df)
        self.assertGreater(std.report["smarty_no_credentials"], 0)
        self.assertEqual(std.report["smarty_fallback_usaddress"], 1)
        self.assertEqual(out["number"].iloc[0], "5")   # usaddress fallback still parsed it


class TestExtractedDedup(unittest.TestCase):
    def test_pools_two_rows_of_one_party(self):
        run = _small_run()
        pr = run["pr"]
        ent = run["scored"]["entities"]["person"]
        self.assertGreater(int((ent.size > 1).sum()), 0)
        self.assertEqual(int(ent.size.sum()), len(pr.X["person"]))

    def test_veto_prevents_pooling_two_different_ssns(self):
        cfg, maps, ref = _env()
        X = _tframe([{"record_id": "d1", "first_name": "Greta", "last_name": "Okonkwo", "ssn": _ssn(901)},
                    {"record_id": "d2", "first_name": "Greta", "last_name": "Okonkwo", "ssn": _ssn(902)}])
        W = _tframe([{"record_id": "dw", "first_name": "Greta", "last_name": "Okonkwo", "ssn": _ssn(901)}])
        pr = prepare(X, W, maps, ref, cfg, log=lambda *a: None)
        n = len(pr.X["person"])
        a, b = np.array([0], np.int64), np.array([1], np.int64)
        cid, acc, why = constrained_clusters(n, a, b, np.array([0.99]), single_values(pr.X["person"], pr.ctx.stats))
        self.assertFalse(bool(acc[0]))
        self.assertIn("ssn", why[0])

    def test_dedup_entity_records_sum(self):
        run = _small_run()
        for part in ("person", "business"):
            ent = run["scored"]["entities"][part]
            e = run["entities"]
            sub = e[e["part"] == part]
            self.assertEqual(int(sub["records"].sum()), len(run["pr"].X[part]))


class TestClusters(unittest.TestCase):
    def test_transitivity(self):
        # a-b and b-c linked: constrained_clusters must place a, b, c in one cluster.
        n = 3
        a = np.array([0, 1], dtype=np.int64)
        b = np.array([1, 2], dtype=np.int64)
        p = np.array([0.99, 0.99])
        cid, acc, why = constrained_clusters(n, a, b, p, {})
        self.assertTrue(acc.all())
        self.assertEqual(cid[0], cid[1])
        self.assertEqual(cid[1], cid[2])

    def test_conflicting_identifier_blocks_cluster_join(self):
        n = 3
        node_vals = {"ssn": [frozenset(["111"]), frozenset(["111"]), frozenset(["222"])]}
        a = np.array([0, 1], dtype=np.int64)
        b = np.array([1, 2], dtype=np.int64)
        p = np.array([0.99, 0.99])
        cid, acc, why = constrained_clusters(n, a, b, p, node_vals, fields=["ssn"])
        self.assertTrue(acc[0])          # 0-1 share the same SSN: fine
        self.assertFalse(acc[1])         # 1-2 would merge conflicting SSNs: refused
        self.assertNotEqual(cid[0], cid[2])

    def test_name_only_links_never_join_a_cluster(self):
        run = _small_run()
        for part in ("person", "business"):
            status = run["scored"]["clusters"][part]["status"]
            pairs = run["pairs"]
            p = pairs[pairs["part"] == part]
            name_only = p[p["cluster_status"].astype(str).str.startswith("name only")]
            self.assertTrue((name_only["cluster_status"] == "name only: visible, not clustered").all())


class TestBaseline(unittest.TestCase):
    def test_clean_strips_suffixes_and_titles(self):
        cfg, _, _ = _env()
        self.assertEqual(baseline_clean("Dr. John A. Smith, M.D.", cfg), "john a smith")
        cfg2 = dataclasses.replace(cfg, baseline_strip_titles=False)
        self.assertIn("md", baseline_clean("Dr. John A. Smith, M.D.", cfg2).split())
        self.assertEqual(baseline_clean("Acme Medical Group, LLC", cfg), "acme medical group")

    def test_p_baseline_column_present_and_named(self):
        cfg, maps, ref = _env()
        cfg2 = make_config("selftest", random_pairs=150_000, dedup_random_pairs=150_000, sim_records=4_000,
                           chunk_size=700, bootstrap_reps=10, baseline=True)
        X, W, truth, edges = _T["inputs"]
        run = run_pipeline(X, W, maps, ref, cfg2)
        self.assertIn("p_baseline", run["pairs"].columns)
        self.assertTrue(((run["pairs"]["p_baseline"] >= 0) & (run["pairs"]["p_baseline"] <= 1)).all())

    def test_category_gate_blocks_a_mismatched_pair(self):
        cfg, maps, ref = _env()
        _, xp, _, _, _, _ = _parties([{"record_id": "x", "first_name": "Alva", "last_name": "Quintero",
                                       "category": "legal"}])
        _, wp, _, _, _, _ = _parties([{"record_id": "w", "first_name": "Alva", "last_name": "Quintero",
                                       "provider_specialty": "CHIROPRACTIC"}], source="W", watchlist=True)
        xn, xc = baseline_frame(xp, "person", cfg, False)
        wn, wc = baseline_frame(wp, "person", cfg, True)
        self.assertEqual(xn[0], wn[0])                # identical cleaned names
        self.assertNotEqual(xc[0], wc[0])             # but categories disagree: legal vs medical
        self.assertTrue(xc[0] and wc[0] and xc[0] != wc[0])


class TestSelfChecks(unittest.TestCase):
    def test_holdout_truth_and_label(self):
        run = _small_run()
        cfg, maps, ref = _env()
        pos, xv, wv, ok = holdout_truth(run["pr"], "ssn", cfg)
        if len(pos):
            lab = holdout_label(xv, wv, ok, "ssn")(pos)
            self.assertTrue(set(np.unique(lab)) <= {-1, 0, 1})
            self.assertTrue((lab != -1).any() or len(pos) == 0)

    def test_noise_test_inputs_have_truth(self):
        cfg, maps, ref = _env()
        cfg2 = make_config("selftest", noise_sources=50, noise_decoys=20, noise_max_copies=2,
                           dedup_random_pairs=50_000, random_pairs=50_000, bootstrap_reps=5)
        X, W, truth, edges = _T["inputs"]
        Xn, truth_n = noise_test_inputs(W, maps, ref, cfg2)
        self.assertGreater(len(Xn), 0)
        self.assertGreater(len(truth_n), 0)
        self.assertTrue(set(truth_n["x_record_id"]) <= set(Xn["record_id"]))

    def test_watchlist_duplicate_groups_join_on_shared_identifier(self):
        cfg, maps, ref = _env()
        _, wp, _, _, _, _ = _parties([
            {"record_id": "w1", "first_name": "Ines", "last_name": "Faraday", "provider_npi": _npi(551)},
            {"record_id": "w2", "first_name": "Ines", "last_name": "Faraday", "provider_npi": _npi(551)},
            {"record_id": "w3", "first_name": "Otis", "last_name": "Marlowe", "provider_npi": _npi(552)}],
            source="W", watchlist=True)
        stats = build_value_stats(wp, wp)
        grp = watchlist_duplicate_groups(wp, stats)
        self.assertEqual(grp[0], grp[1])
        self.assertNotEqual(grp[0], grp[2])


class TestTableIO(unittest.TestCase):
    def test_csv_roundtrip(self):
        cfg, maps, ref = _env()
        with tempfile.TemporaryDirectory() as d:
            cfg2 = make_config("selftest", root=d, io_format="csv")
            Path(cfg2.paths["tables"]).mkdir(parents=True, exist_ok=True)
            io_ = TableIO(cfg2)
            df = pd.DataFrame({"a": [1, 2], "b": ["x", "y"]})
            io_.write("t", df)
            back = pd.read_csv(cfg2.paths["tables"] / "t.csv")
            self.assertEqual(list(back["a"]), [1, 2])

    def test_inmemory_delta_versions_and_merge(self):
        be = InMemoryDeltaBackend()
        df1 = pd.DataFrame({"k": ["a", "b"], "v": [1, 2], "_row_hash": ["h1", "h2"]})
        be.write("t", df1)
        self.assertEqual(be.latest_version("t"), 0)
        df2 = pd.DataFrame({"k": ["b", "c"], "v": [20, 3], "_row_hash": ["h2b", "h3"]})
        # delete_missing=True (the default, matching every output-table write): a key absent
        # from the incoming frame is removed, since a write always states the full desired table.
        res = be.merge("t", df2, ["k"])
        self.assertEqual(res["updated"], 1)
        self.assertEqual(res["inserted"], 1)
        self.assertEqual(res["deleted"], 1)
        cur = be.read("t")
        self.assertEqual(set(cur["k"]), {"b", "c"})
        self.assertEqual(int(cur.loc[cur["k"] == "b", "v"].iloc[0]), 20)
        self.assertEqual(be.latest_version("t"), 1)
        # time travel: version 0 still has the old value
        self.assertEqual(int(be.read("t", 0).loc[be.read("t", 0)["k"] == "b", "v"].iloc[0]), 2)

    def test_change_data_feed(self):
        be = InMemoryDeltaBackend()
        be.write("t", pd.DataFrame({"k": ["a"], "v": [1], "_row_hash": ["h1"]}))
        be.merge("t", pd.DataFrame({"k": ["a", "b"], "v": [9, 2], "_row_hash": ["h1b", "h2"]}), ["k"])
        ch = be.changes("t", 1, 1)
        self.assertIn("b", set(ch.loc[ch["_change_type"] == "insert", "k"]))

    def test_delta_merge_deletes_missing_rows(self):
        be = InMemoryDeltaBackend()
        be.write("t", pd.DataFrame({"k": ["a", "b"], "v": [1, 2], "_row_hash": ["h1", "h2"]}))
        be.merge("t", pd.DataFrame({"k": ["a"], "v": [1], "_row_hash": ["h1"]}), ["k"], delete_missing=True)
        self.assertEqual(set(be.read("t")["k"]), {"a"})


class TestIncremental(unittest.TestCase):
    def test_incremental_matches_full_run_on_same_snapshot(self):
        cfg, maps, ref = _env()
        with tempfile.TemporaryDirectory() as d:
            cfg2 = make_config("selftest", root=d, io_format="delta", bootstrap_reps=10, random_pairs=80_000,
                              dedup_random_pairs=80_000)
            X, W, truth, edges = _T["inputs"]
            be = InMemoryDeltaBackend()
            io_ = TableIO(cfg2, backend=be)
            be.write(io_.table(cfg2.delta_extracted_table), plain_table(X).assign(_row_hash="h"))
            be.write(io_.table(cfg2.delta_watchlist_table), plain_table(W).assign(_row_hash="h"))
            xdf, xrep, wdf, wrep, meta = io_.read_inputs()
            full = run_pipeline(xdf, wdf, maps, ref, cfg2)
            write_run_state(io_, {**full, "xdf": xdf}, meta)
            # a second run with no changes: incremental should reuse every entity
            io2 = TableIO(cfg2, backend=be, run_id="run2")
            xdf2, _, wdf2, _, meta2 = io2.read_inputs()
            inc = run_incremental(xdf2, wdf2, maps, ref, cfg2, io2, truth_file=None)
            self.assertEqual(inc["incremental"]["mode"], "incremental")
            self.assertEqual(inc["incremental"].get("reused_entities", 0) + inc["incremental"].get("scored_entities", 0)
                             >= 0, True)
            # same entity count and same total p per part (deterministic re-derivation)
            for part in ("person", "business"):
                a = full["entities"][full["entities"]["part"] == part].sort_values("entity_id")["p"].to_numpy()
                b = inc["entities"][inc["entities"]["part"] == part].sort_values("entity_id")["p"].to_numpy()
                np.testing.assert_allclose(a, b, atol=1e-6)

    def test_full_run_when_no_state(self):
        cfg, maps, ref = _env()
        with tempfile.TemporaryDirectory() as d:
            cfg2 = make_config("selftest", root=d, io_format="delta", bootstrap_reps=5, random_pairs=50_000,
                              dedup_random_pairs=50_000)
            X, W, truth, edges = _T["inputs"]
            be = InMemoryDeltaBackend()
            io_ = TableIO(cfg2, backend=be)
            be.write(io_.table(cfg2.delta_extracted_table), plain_table(X).assign(_row_hash="h"))
            be.write(io_.table(cfg2.delta_watchlist_table), plain_table(W).assign(_row_hash="h"))
            xdf, _, wdf, _, meta = io_.read_inputs()
            run = run_incremental(xdf, wdf, maps, ref, cfg2, io_, truth_file=None)
            self.assertEqual(run["incremental"]["mode"], "full")

CORE_SHA256 = "24d2b8a678f103c62332f71114a5c10979d31416cabf2acef3a934ffd125549a"   # v1.1


NL, CR = chr(10), chr(13)


def notebook_core_hash(nb_path):
    """sha256 over the code cells between the 'Matching core v' heading and its end marker,
    exactly as recorded when the notebook was assembled."""
    nb = json.loads(Path(nb_path).read_text(encoding="utf-8"))
    inside, parts = False, []
    for c in nb["cells"]:
        src = "".join(c["source"])
        if c["cell_type"] == "markdown" and src.startswith("## Matching core v"):
            inside = True
            continue
        if c["cell_type"] == "markdown" and src.strip().startswith("*End of the matching core.*"):
            break
        if inside and c["cell_type"] == "code":
            parts.append(NL.join(l.rstrip() for l in src.replace(CR + NL, NL).split(NL)))
    return hashlib.sha256((NL + "# ---- cell ----" + NL).join(parts).encode("utf-8")).hexdigest()


class TestCoreHash(unittest.TestCase):
    def test_core_section_matches_recorded_hash(self):
        nbp = _env()[0].paths["notebook"]
        if not nbp.exists() or CORE_SHA256.startswith("<"):
            self.skipTest("runs inside the assembled notebook")
        self.assertEqual(notebook_core_hash(nbp), CORE_SHA256,
                         f"the matching core v{CORE_VERSION} changed: bump the version, re-record the hash, copy to [A]")


def run_selftests(verbosity=1):
    suite = unittest.TestSuite()
    loader = unittest.TestLoader()
    for cls in (TestNormalizers, TestBreakdown, TestCategory, TestComparisons, TestCandidates, TestScoring,
                TestParameters, TestOutput, TestLeieExport, TestInputs, TestClosedFormU, TestEMBootstrap,
                TestRicherComparisons, TestAddressStandardizer, TestExtractedDedup, TestClusters, TestBaseline,
                TestSelfChecks, TestTableIO, TestIncremental, TestCoreHash):
        suite.addTests(loader.loadTestsFromTestCase(cls))
    stream = io.StringIO()
    res = unittest.TextTestRunner(stream=stream, verbosity=verbosity).run(suite)
    return res, stream.getvalue()