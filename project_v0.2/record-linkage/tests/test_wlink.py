# %% [markdown]
# ## 17 · Self-tests
#
# `unittest` cases for every part of the package (normalizers, breakdown, category, candidates,
# comparison levels, score invariants, parameters, output, the LEIE exporter), an end-to-end
# run on a small synthetic set with its hand-written edge cases, and the hash check of the
# matching core. They run on their own small inputs, so they are the same for every dataset.
#
# **Where Splink would do better.** Splink has a large test suite across DuckDB, Spark and
# SQLite backends; these tests cover this package only.

# %%
import unittest, tempfile, io
from wlink.core import *
from wlink.core import _undot, _DBA_RE, _VIN_MAP, _VIN_W, _words, _near
from wlink.config import *
from wlink.load import *
from wlink.explode import *
from wlink.category import *
from wlink.simulate import *
from wlink.synth import *
from wlink.leie import *
from wlink.prepare import *
from wlink.candidates import *
from wlink.params import *
from wlink.score import *
from wlink.output import *
from wlink.diagnostics import *
from wlink.edgecheck import *
from wlink.pipeline import *

_T = {}


def _env():
    """Mappings, reference tables and a small configuration, loaded once."""
    if "cfg" not in _T:
        cfg = make_config("selftest", random_pairs=150_000, sim_records=4_000, chunk_size=700)
        _T["cfg"] = cfg
        _T["maps"] = load_mappings(cfg.paths["mappings"])
        _T["ref"] = load_reference(cfg.paths["reference"], cfg.core)
    return _T["cfg"], _T["maps"], _T["ref"]


def _small_run():
    """One end-to-end run on a small synthetic set with every edge case (cached)."""
    if "run" not in _T:
        cfg, maps, ref = _env()
        X, W, truth, edges, _ = make_dataset(400, 100, 1200, ref, maps["simulation_noise"], cfg.seed)
        X, _ = validate_input(X, "X"); W, _ = validate_input(W, "W")
        tf = cfg.paths["out"] / "selftest_truth.csv"
        truth.to_csv(tf, index=False)
        _T["run"] = run_pipeline(X, W, maps, ref, cfg, truth_file=tf)
        _T["inputs"] = (X, W, truth, edges)
    return _T["run"]


def _tframe(rows):
    df = pd.DataFrame([blank_row(r["record_id"], **{k: v for k, v in r.items() if k != "record_id"})
                       for r in rows], columns=SCHEMA)
    return validate_input(df, "t")[0]


def _parties(rows, source="X", watchlist=False):
    cfg, maps, ref = _env()
    nick = Nicknames(ref["nicknames"])
    return rows_to_parties(_tframe(rows), source, maps, nick, cfg.core, watchlist)


class TestNormalizers(unittest.TestCase):
    def test_ssn(self):
        self.assertEqual(norm_ssn("212-45-6781"), ("212456781", True, ""))
        for bad in ("000-12-3456", "666123456", "912345678", "123456789", "111111111", "12345"):
            self.assertFalse(norm_ssn(bad)[1], bad)

    def test_tin(self):
        self.assertTrue(norm_tin("45-1234567")[1])
        self.assertFalse(norm_tin("07-1234567")[1])
        self.assertFalse(norm_tin("000000000")[1])

    def test_npi_luhn(self):
        self.assertTrue(npi_luhn_ok("1234567893"))
        self.assertFalse(npi_luhn_ok("1234567890"))
        self.assertEqual(norm_npi("0000000000"), ("", False, ""))
        self.assertEqual(norm_npi("1234567890")[2], "check_digit")
        self.assertTrue(npi_luhn_ok(luhn_npi("123456789")))

    def test_phone_email(self):
        self.assertEqual(norm_phone("1 (718) 392-4411")[0], "7183924411")
        self.assertFalse(norm_phone("2125550123")[1])
        self.assertFalse(norm_phone("0123456789")[1])
        self.assertEqual(norm_email(" A.B@Mail.Test ")[0], "a.b@mail.test")
        self.assertFalse(norm_email("none@none.com")[1])
        self.assertFalse(norm_email("not an email")[1])

    def test_vin(self):
        self.assertTrue(norm_vin("1M8GDM9AXKP042788")[1])
        self.assertEqual(norm_vin("1M8GDM9AYKP042788")[2], "check_digit")
        self.assertEqual(norm_vin("1M8GDM9AXKP04278I")[2], "format")

    def test_dob(self):
        self.assertEqual(norm_dob("3/4/1961")[0], "1961-03-04")
        self.assertEqual(norm_dob("19610304")[0], "1961-03-04")
        self.assertEqual(norm_dob("1900-01-01")[2], "placeholder")
        self.assertEqual(norm_dob("02/30/1960")[2], "impossible")
        self.assertEqual(norm_dob("2090-01-01")[2], "range")

    def test_qualified_ids(self):
        self.assertEqual(norm_dl("d123-456", "New York")[0], "NY:D123456")
        self.assertEqual(norm_plate("abc 123", "")[0], ":ABC123")
        self.assertFalse(norm_license("00000", "NY")[1])

    def test_address(self):
        self.assertEqual(parse_street_line("2161 UNIVERSITY AVENUE W, STE 5"), ("2161", "", "UNIVERSITY", "AVE", "5"))
        self.assertEqual(parse_street_line("P O BOX 2161")[:3], ("2161", "", "PO BOX"))
        self.assertEqual(parse_street_line("222 N W 45TH AVE")[1], "NW")
        self.assertEqual(norm_street_parts("12", "", "Main Street", "", "Apt 4B"), ("12", "", "MAIN", "ST", "4B"))
        self.assertEqual(norm_state("new york"), "NY")
        self.assertEqual(norm_zip("2134"), "02134")
        self.assertEqual(norm_city("St. Louis"), "SAINT LOUIS")

    def test_person_names(self):
        self.assertEqual(clean_person("Dr. John", "a.", "Smith Jr."), ("JOHN", "A", "SMITH"))
        self.assertEqual(clean_person("", "", "Moy, Marvin"), ("MARVIN", "", "MOY"))
        self.assertEqual(clean_person("William A Weiner", "", ""), ("WILLIAM", "A", "WEINER"))
        self.assertEqual(clean_person("José", "", "O'Brien"), ("JOSE", "", "OBRIEN"))
        self.assertEqual(parse_person("WILLIAM A. WEINER, D.O.")["creds"], ["DO"])

    def test_org_aliases(self):
        self.assertEqual(org_aliases("Rutland Medical P.C. d/b/a Soul Radiology Medical Imaging"),
                         [["RUTLAND", "MEDICAL"], ["SOUL", "RADIOLOGY", "MEDICAL", "IMAGING"]])
        self.assertEqual(org_aliases("AMERICAN TRANSIT INS. CO."), [["AMERICAN", "TRANSIT", "INSURANCE"]])
        self.assertEqual(org_aliases("Smith & Jones LLP"), [["SMITH", "JONES"]])
        self.assertEqual(org_alias_string("Nexray Medical Imaging, P.C."), "NEXRAY MEDICAL IMAGING")

    def test_nicknames(self):
        _, _, ref = _env()
        n = Nicknames(ref["nicknames"])
        self.assertTrue(n.roots("Bill") & n.roots("William"))
        self.assertFalse(n.roots("Robert") & n.roots("William"))
        self.assertEqual(nysiis("Weiner"), nysiis("Wiener"))

    def test_near(self):
        self.assertTrue(_near("123456789", "123456780"))
        self.assertTrue(_near("123456789", "123457689"))
        self.assertFalse(_near("123456789", "123456798x"))
        self.assertFalse(_near("123456789", "123450780"))


class TestBreakdown(unittest.TestCase):
    def test_person_and_business_tied(self):
        rows, parties, details, ties, empty, _ = _parties([
            {"record_id": "r1", "first_name": "Ann", "last_name": "Lee", "business_name": "Lee Chiropractic PC",
             "tin": "45-1234567", "work_phone": "7183924411", "home_phone": "7183924412", "email": "ann@lee.test",
             "street_number": "5", "street_name": "Main", "street_type": "St", "zip": "11211", "state": "NY"}])
        self.assertEqual(sorted(parties["part"]), ["business", "person"])
        self.assertEqual(len(ties), 1)
        P = parties[parties["part"] == "person"].iloc[0]
        B = parties[parties["part"] == "business"].iloc[0]
        self.assertEqual(P["tie"], B["party_id"])
        self.assertEqual((P["id_email"], B["id_email"]), ("ann@lee.test", ""))
        self.assertEqual((P["phone_own"], P["phone_row"], B["phone_row"]), ("7183924412", "7183924411", "7183924411"))
        self.assertEqual(B["id_tin"], "451234567")
        self.assertEqual(P["addr_street"], B["addr_street"])
        own = details.set_index(["party_id", "type", "value"])["ownership"]
        self.assertEqual(own[(P["party_id"], "phone", "7183924411")], "row")
        self.assertEqual(own[(P["party_id"], "phone", "7183924412")], "own")

    def test_tin_only_and_empty(self):
        _, parties, _, _, empty, _ = _parties([{"record_id": "t", "tin": "451234567"},
                                               {"record_id": "e", "city": "Brooklyn"},
                                               {"record_id": "b", "business_name": "Acme Towing", "email": "x@acme.test"}])
        self.assertEqual(list(parties["record_id"]), ["t", "b"])
        self.assertEqual(list(empty["record_id"]), ["e"])
        self.assertEqual(parties.set_index("record_id").loc["b", "id_email"], "x@acme.test")

    def test_invalid_values_kept_visible_not_compared(self):
        _, parties, details, _, _, _ = _parties([{"record_id": "r", "first_name": "A", "last_name": "B",
                                                   "provider_npi": "1234567890", "ssn": "123456789"}])
        self.assertEqual(parties.iloc[0]["id_npi"], "")
        d = details.set_index("type")
        self.assertFalse(bool(d.loc["npi", "valid"]))
        self.assertEqual(d.loc["ssn", "reason"], "placeholder")

    def test_value_index(self):
        _, p1, d1, _, _, _ = _parties([{"record_id": "a", "first_name": "Ann", "last_name": "Lee", "ssn": "212456781"},
                                       {"record_id": "b", "first_name": "A", "last_name": "Lee", "ssn": "212456781"},
                                       {"record_id": "c", "first_name": "Bo", "last_name": "Kim", "ssn": "313456781"},
                                       {"record_id": "d", "first_name": "Cy", "last_name": "Poe", "ssn": "313456781"}])
        vi = mark_junk(value_index_fast(p1, d1), CoreParams(junk_holders=1)).set_index(["type", "value"])
        self.assertTrue(bool(vi.loc[("ssn", "212456781"), "single_holder"]))      # Ann Lee = A. Lee
        self.assertFalse(bool(vi.loc[("ssn", "313456781"), "single_holder"]))
        self.assertTrue(bool(vi.loc[("ssn", "313456781"), "junk"]))


class TestCategory(unittest.TestCase):
    def test_extracted_precedence(self):
        cfg, maps, _ = _env()
        _, p, _, _, _, _ = _parties([
            {"record_id": "npi", "category": "legal", "first_name": "A", "last_name": "B", "provider_npi": luhn_npi("123456789")},
            {"record_id": "giv", "category": "witness", "first_name": "C", "last_name": "D", "business_name": "D Auto Body"},
            {"record_id": "kw", "business_name": "Pemberton & Vasquez LLP"},
            {"record_id": "oth", "category": "other", "first_name": "E", "last_name": "F"},
            {"record_id": "lic", "first_name": "G", "last_name": "H", "professional_license_type": "DC"}])
        e = p.set_index(["record_id", "part"])
        self.assertEqual(e.loc[("npi", "person"), "category"], "medical")
        self.assertTrue(bool(e.loc[("npi", "person"), "category_mismatch"]))
        self.assertEqual(e.loc[("giv", "person"), "category"], "witness")          # given before keyword
        self.assertEqual(e.loc[("giv", "person"), "category_inferred"], "repair shop")
        self.assertEqual(e.loc[("giv", "person"), "prior_group"], "private")
        self.assertEqual(e.loc[("giv", "business"), "prior_group"], "business")
        self.assertEqual(e.loc[("kw", "business"), "category"], "legal")
        self.assertEqual(e.loc[("oth", "person"), "category"], "")                 # other = no information
        self.assertEqual(e.loc[("oth", "person"), "prior_group"], "unknown")
        self.assertEqual(e.loc[("lic", "person"), "cat_strength"], "strong")
        self.assertNotIn("witness", set(p["category_inferred"]))

    def test_watchlist_mapping_and_unmapped(self):
        _, p, _, _, _, unm = _parties([
            {"record_id": "w1", "first_name": "A", "last_name": "B", "provider_specialty": "CHIROPRACTIC"},
            {"record_id": "w2", "first_name": "C", "last_name": "D", "professional_license_type": "LAW PRACTICE"},
            {"record_id": "w3", "first_name": "E", "last_name": "F", "provider_specialty": "BASKET WEAVING"}],
            source="W", watchlist=True)
        self.assertEqual(list(p["category"]), ["medical", "legal", ""])
        self.assertIn("BASKET WEAVING", set(unm["value"]))


def _mini_ctx(xp, wp):
    cfg, maps, ref = _env()
    nick = Nicknames(ref["nicknames"])
    rar = Rarity(ref["surnames"], ref["first_names"], ref["org_tokens"], ref["meta"], cfg.core)
    stats = build_value_stats(xp, wp)
    dba = declared_dba_pairs(pd.concat([xp["org_aliases"], wp["org_aliases"]]))
    return CompareContext(cfg.core, rar, nick, stats, dba,
                          {"first_nick_or_close": 0.01, "first_initial": 0.05, "first_empty": 0.05,
                           "first_differs": 0.9, "last_close": 0.002})


def _pair_levels(xrow, wrow, part="person"):
    _, xp, _, _, _, _ = _parties([xrow])
    _, wp, _, _, _, _ = _parties([wrow], source="W", watchlist=True)
    X, W = split_parts(xp)[part], split_parts(wp)[part]
    ctx = _mini_ctx(xp, wp)
    lv = compare_pairs(X, W, [0], [0], part, ctx)
    return lv.iloc[0], ctx


def _lev(lv, f):
    c = int(lv[f])
    return "empty" if c < 0 else FIELD_LEVELS[f][c]


class TestComparisons(unittest.TestCase):
    def test_name_levels(self):
        cases = [(("John", "Smith"), ("John", "Smith"), "exact"),
                 (("Bill", "Szczepanski"), ("William", "Szczepanski"), "first_nick_or_close"),
                 (("Maria", "Garcia-Villanueva"), ("Maria", "Garcia"), "last_close_first_agrees"),
                 (("R", "Featherstone"), ("Rupert", "Featherstone"), "initial_agrees"),
                 (("Chidi", "Nwachukwu"), ("Nwachukwu", "Chidi"), "swapped"),
                 (("", "Webb"), ("Henry", "Webb"), "first_empty"),
                 (("Alan", "Webb"), ("Henry", "Webb"), "first_differs"),
                 (("Alan", "Webb"), ("Henry", "Moss"), "else")]
        for (f1, l1), (f2, l2), want in cases:
            row = {"record_id": "x", "first_name": f1, "last_name": l1}
            if not l1:
                row["business_name"] = "placeholder biz"    # keep a party with no surname
                row["first_name"] = ""
                row["middle_name"] = ""
            lv, _ = _pair_levels(row, {"record_id": "w", "first_name": f2, "last_name": l2})
            got = _lev(lv, "name") if l1 else "empty"
            self.assertEqual(got, want, (f1, l1, f2, l2))

    def test_dob_levels(self):
        for a, b, want in (("1975-04-09", "1975-04-09", "exact"), ("1975-04-09", "1975-09-04", "swap_or_typo"),
                           ("1975-04-09", "1975-04-08", "swap_or_typo"), ("1975-04-09", "1975-04-21", "year_month"),
                           ("1975-04-09", "1975-11-21", "year"), ("1975-04-09", "1980-11-21", "differs"),
                           ("1975-04-09", "1900-01-01", "empty")):
            lv, _ = _pair_levels({"record_id": "x", "first_name": "A", "last_name": "B", "dob": a},
                                 {"record_id": "w", "first_name": "A", "last_name": "B", "dob": b})
            self.assertEqual(_lev(lv, "dob"), want, (a, b))

    def test_address_levels(self):
        base = {"street_number": "5", "street_name": "Main", "street_type": "St", "unit": "2", "city": "Brooklyn",
                "state": "NY", "zip": "11211"}
        for change, want in (({}, "exact"), ({"unit": "3"}, "street"), ({"street_number": "7"}, "zip"),
                             ({"street_number": "7", "zip": "11222"}, "city_state"),
                             ({"street_number": "7", "zip": "11222", "city": "Albany"}, "state"),
                             ({"street_number": "7", "zip": "90026", "city": "Los Angeles", "state": "CA"}, "differs")):
            lv, _ = _pair_levels({"record_id": "x", "first_name": "A", "last_name": "B", **base},
                                 {"record_id": "w", "first_name": "A", "last_name": "B", **{**base, **change}})
            self.assertEqual(_lev(lv, "address"), want, change)

    def test_identifier_levels_and_veto(self):
        lv, _ = _pair_levels({"record_id": "x", "first_name": "A", "last_name": "B", "ssn": "212456781"},
                             {"record_id": "w", "first_name": "A", "last_name": "B", "ssn": "212456781"})
        self.assertEqual(_lev(lv, "ssn"), "exact")
        lv, _ = _pair_levels({"record_id": "x", "first_name": "A", "last_name": "B", "ssn": "212456781"},
                             {"record_id": "w", "first_name": "A", "last_name": "B", "ssn": "212456718"})
        self.assertEqual(_lev(lv, "ssn"), "near")
        self.assertFalse(bool(lv["veto_ssn"]))
        lv, _ = _pair_levels({"record_id": "x", "first_name": "A", "last_name": "B", "ssn": "212456781"},
                             {"record_id": "w", "first_name": "A", "last_name": "B", "ssn": "313999222"})
        self.assertEqual(_lev(lv, "ssn"), "differs")
        self.assertTrue(bool(lv["veto_ssn"]))
        lv, _ = _pair_levels({"record_id": "x", "first_name": "A", "last_name": "B", "driver_license_number": "D1234567", "driver_license_state": "NY"},
                             {"record_id": "w", "first_name": "A", "last_name": "B", "driver_license_number": "K9876543", "driver_license_state": "NJ"})
        self.assertFalse(bool(lv["veto_dl"]))                                       # two states: no veto
        lv, _ = _pair_levels({"record_id": "x", "business_name": "A Corp", "clinic_npi": luhn_npi("111111112")},
                             {"record_id": "w", "business_name": "A Corp", "clinic_npi": luhn_npi("222222223")}, "business")
        self.assertEqual(_lev(lv, "cnpi"), "differs")
        self.assertFalse(bool(lv["veto_cnpi"]))                                     # clinic NPI never vetoes

    def test_phone_levels(self):
        lv, _ = _pair_levels({"record_id": "x", "first_name": "A", "last_name": "B", "home_phone": "7183924411"},
                             {"record_id": "w", "first_name": "A", "last_name": "B", "home_phone": "7183924411"})
        self.assertEqual(_lev(lv, "phone"), "exact_owned_single")
        lv, _ = _pair_levels({"record_id": "x", "first_name": "A", "last_name": "B", "work_phone": "7183924411"},
                             {"record_id": "w", "first_name": "C", "last_name": "D", "work_phone": "7183924411"})
        self.assertEqual(_lev(lv, "phone"), "exact_shared")

    def test_org_levels(self):
        cases = [("Lakeshore Med Ctr", "Lakeshore Medical Center Inc", "exact"),
                 ("Kestrel Imaging LLC d/b/a Harborview Radiology", "Harborview Radiology", "exact"),
                 ("Pinecrest Orthopedic", "Pinecrest Orthopedic Rehabilitation", "short_form"),
                 ("Allport Indemnity Company", "Allport Casualty Surety Company", "sibling"),
                 ("Quixotic Holdings", "Unrelated Towing", "none")]
        for a, b, want in cases:
            lv, _ = _pair_levels({"record_id": "x", "business_name": a}, {"record_id": "w", "business_name": b}, "business")
            self.assertEqual(_lev(lv, "org"), want, (a, b))

    def test_spec_cat(self):
        lv, _ = _pair_levels({"record_id": "x", "first_name": "A", "last_name": "B", "provider_specialty": "Chiropractor", "provider_npi": luhn_npi("123456789")},
                             {"record_id": "w", "first_name": "A", "last_name": "B", "provider_specialty": "CHIROPRACTIC"})
        self.assertEqual(_lev(lv, "spec_cat"), "specialty")


class TestCandidates(unittest.TestCase):
    def setUp(self):
        run = _small_run()
        self.pr = run["pr"]
        self.cfg = _env()[0]

    def test_chunked_equals_unchunked(self):
        for part in ("person", "business"):
            X, W = self.pr.X[part], self.pr.W[part]
            plan = CandidatePlan(X, W, part, self.pr.ctx.rarity, self.cfg, self.pr.ctx.dba_pairs)
            one = plan.chunk(0, len(X)).sort_values(["l", "r"]).reset_index(drop=True)
            many = pd.concat([plan.chunk(lo, min(len(X), lo + 97)) for lo in range(0, len(X), 97)])
            many = many.sort_values(["l", "r"]).reset_index(drop=True)
            self.assertTrue(one.equals(many), part)

    def test_each_rule_proposes_its_target(self):
        X = _tframe([{"record_id": "x1", "first_name": "Bill", "last_name": "Kowalski"},
                    {"record_id": "x6", "first_name": "Rita", "last_name": "Oldname", "dob": "1970-01-02"},
                    {"record_id": "x2", "first_name": "Ann", "last_name": "Zyx", "ssn": "212456781"},
                    {"record_id": "x3", "first_name": "Mo", "last_name": "Qwe", "home_phone": "7183924411"},
                    {"record_id": "x4", "first_name": "Chidi", "last_name": "Nwachukwu"},
                    {"record_id": "x5", "last_name": "Plover", "dob": "1981-01-01"}])
        W = _tframe([{"record_id": "w1", "first_name": "William", "last_name": "Kowalski"},
                    {"record_id": "w6", "first_name": "Rita", "last_name": "Newname", "dob": "1970-01-02"},
                    {"record_id": "w2", "first_name": "Zed", "last_name": "Other", "ssn": "212456781"},
                    {"record_id": "w3", "first_name": "Zed", "last_name": "Else", "home_phone": "7183924411"},
                    {"record_id": "w4", "first_name": "Nwachukwu", "last_name": "Chidi"},
                    {"record_id": "w5", "first_name": "Ruth", "last_name": "Plover"}])
        cfg, maps, ref = _env()
        pr = prepare(X, W, maps, ref, cfg, log=lambda *a: None)
        plan = CandidatePlan(pr.X["person"], pr.W["person"], "person", pr.ctx.rarity, cfg)
        c = plan.chunk(0, len(pr.X["person"]))
        got = {(pr.X["person"]["record_id"][l], pr.W["person"]["record_id"][r]): set(n.split(","))
               for l, r, n in zip(c["l"], c["r"], plan.rule_names(c["rules"]))}
        self.assertIn("nysiis_canon_initial", got[("x1", "w1")])
        self.assertNotIn("nysiis_initial", got[("x1", "w1")])
        self.assertIn("dob_initial", got[("x6", "w6")])
        self.assertIn("ssn", got[("x2", "w2")])
        self.assertIn("phone", got[("x3", "w3")])
        self.assertIn("swapped", got[("x4", "w4")])
        self.assertIn("nysiis_first_missing", got[("x5", "w5")])

    def test_cross_file_same_part_only(self):
        S = _small_run()["scored"]
        for part in ("person", "business"):
            s = S[part][0]
            self.assertTrue((s["l"] < len(self.pr.X[part])).all() and (s["r"] < len(self.pr.W[part])).all())

    def test_strict_rules_subset_of_blocking(self):
        cfg = self.cfg
        for part in ("person", "business"):
            X, W = self.pr.X[part], self.pr.W[part]
            l, r = strict_pairs(X, W, part, self.pr.ctx)
            keys = self.pr and _small_run()["scored"][part][3]["keys"]
            k = l.astype(np.int64) * np.int64(1 << 32) + r
            self.assertTrue(np.isin(k, keys).all(), part)


class TestScoring(unittest.TestCase):
    def setUp(self):
        self.run = _small_run()

    def test_empty_fields_are_zero_bits(self):
        for part in ("person", "business"):
            S, L = self.run["scored"][part][0], self.run["scored"][part][1]
            for f in COMPARED_FIELDS[part]:
                self.assertTrue((S.loc[L[f].to_numpy() < 0, f"bits_{f}"] == 0).all(), (part, f))

    def test_evidence_sums_to_total(self):
        ev = self.run["evidence"].groupby("pair_id")["bits"].sum()
        pairs = self.run["pairs"].set_index("pair_id")["bits"]
        diff = (pairs.reindex(ev.index) - ev).abs()
        self.assertLess(float(diff.max()), 1e-6)
        zero = pairs[~pairs.index.isin(ev.index)]
        self.assertTrue((zero.abs() < 1e-9).all())

    def test_agreement_never_lowers_p(self):
        part = "person"
        S, L = self.run["scored"][part][0], self.run["scored"][part][1]
        w = self.run["params"]["weights"][part]
        lv = L.copy()
        base = score_pairs(lv, part, w, S["prior_logit"].to_numpy())
        for f in ("ssn", "npi", "email", "dob", "name"):
            lv2 = lv.copy()
            empty = lv2[f].to_numpy() < 0
            lv2.loc[empty, f] = 0
            lv2.loc[empty, f"uv_{f}"] = 1e-6
            more = score_pairs(lv2, part, w, S["prior_logit"].to_numpy())
            ok = base["veto"].to_numpy() == ""
            self.assertTrue((more["p"].to_numpy()[ok] >= base["p"].to_numpy()[ok] - 1e-12).all(), f)

    def test_veto_visible_with_p_zero(self):
        P = self.run["pairs"]
        v = P[P["veto"] != ""]
        self.assertGreater(len(v), 0)
        self.assertTrue((v["p"] == 0).all())

    def test_name_only_never_identifier(self):
        part = "person"
        S, L = self.run["scored"][part][0], self.run["scored"][part][1]
        idn = S["basis"].to_numpy() == "identifier"
        any_id = np.zeros(len(L), bool)
        for f in IDENTIFIER_FIELDS:
            if f in L:
                any_id |= (L[f].to_numpy() == 0) & (S[f"bits_{f}"].to_numpy() > 0)
        self.assertTrue((any_id[idn]).all())
        self.assertTrue((S.loc[~any_id, "basis"] != "identifier").all())

    def test_coparty_needs_anchor_and_name(self):
        S, L = self.run["scored"]["person"][0], self.run["scored"]["person"][1]
        cp = S["bits_co_party"].to_numpy() > 0
        self.assertTrue((S["bits_name"].to_numpy()[cp] > 0).all())
        self.assertTrue((L["co_party"].to_numpy()[cp] == 0).all())
        anchors = business_anchors(self.run["scored"]["business"][0], _env()[0])
        X, W = self.run["pr"].X["person"], self.run["pr"].W["person"]
        tl, tr = X["tie_pos"].to_numpy()[S["l"].to_numpy()[cp]], W["tie_pos"].to_numpy()[S["r"].to_numpy()[cp]]
        self.assertTrue(np.isin(tl.astype(np.int64) * np.int64(1 << 32) + tr, anchors).all())

    def test_deterministic(self):
        cfg, maps, ref = _env()
        X, W, truth, edges = _T["inputs"]
        again = run_pipeline(X, W, maps, ref, cfg)
        a = self.run["pairs"].drop(columns=["_l", "_r"]).reset_index(drop=True)
        b = again["pairs"].drop(columns=["_l", "_r"]).reset_index(drop=True)
        pd.testing.assert_frame_equal(a, b)

    def test_edge_cases(self):
        X, W, truth, edges = _T["inputs"]
        r = check_edge_cases(self.run["pr"], self.run["pairs"], None, self.run["scored"], self.run["entities"], edges)
        self.assertTrue(r["ok"].all(), r[~r["ok"]].to_string())

    def test_every_true_name_only_pair_visible(self):
        t = [x for x in self.run["diag"]["truth"] if x.get("true_pairs")]
        for x in t:
            self.assertEqual(x["proposed"], x["kept"], x["part"])

    def test_entities_one_row_per_party(self):
        e = self.run["entities"]
        self.assertTrue(e["party_id"].is_unique)
        n = sum(len(self.run["pr"].X[p]) for p in ("person", "business"))
        self.assertEqual(len(e), n)


class TestParameters(unittest.TestCase):
    def test_u_closed_form(self):
        lv = pd.DataFrame({"email": np.array([0] * 3 + [1] * 97 + [-1] * 50, dtype=np.int8)})
        u = estimate_u(lv, ["email"], CoreParams(u_pseudo=0.0)).set_index("level")["u"]
        self.assertAlmostEqual(u["exact"], 0.03)
        self.assertAlmostEqual(u["differs"], 0.97)

    def test_anchor_field_excluded(self):
        lv = pd.DataFrame({"ssn": np.array([0, 0, 0, 2], dtype=np.int8), "dob": np.array([0, 0, 4, 4], dtype=np.int8)})
        c = level_counts(lv, ["ssn", "dob"], exclude={"ssn": np.array([True, True, True, False])})
        self.assertEqual(c["ssn"][1], 1.0)
        self.assertEqual(c["dob"][1], 4.0)

    def test_em_recovers_m(self):
        rng = np.random.default_rng(3)
        n, lam = 20000, 0.3
        m = {"a": np.array([0.8, 0.15, 0.05]), "b": np.array([0.7, 0.3]), "c": np.array([0.9, 0.1])}
        u = {"a": np.array([0.01, 0.09, 0.9]), "b": np.array([0.05, 0.95]), "c": np.array([0.1, 0.9])}
        match = rng.random(n) < lam
        data = {}
        for f in m:
            k = len(m[f])
            data[f] = np.where(match, rng.choice(k, n, p=m[f]), rng.choice(k, n, p=u[f])).astype(np.int8)
        lv = pd.DataFrame(data)
        saved = {f: FIELD_LEVELS.get(f) for f in m}
        FIELD_LEVELS.update({"a": ["x", "y", "z"], "b": ["x", "y"], "c": ["x", "y"]})
        try:
            est, lam_hat, _ = em_fixed_u(lv, ["a", "b"], ["c"], {"c": m["c"]}, u, CoreParams())
        finally:
            for f, v in saved.items():
                if v is None:
                    FIELD_LEVELS.pop(f, None)
        self.assertAlmostEqual(lam_hat, lam, delta=0.02)
        self.assertLess(np.abs(est["a"][0] - m["a"]).max(), 0.03)
        self.assertLess(np.abs(est["b"][0] - m["b"]).max(), 0.03)

    def test_combine_chain(self):
        p = CoreParams(alpha=5, n_min=50)
        rows = combine_m("email", [("anchors_strict", np.array([8.0, 2.0]), 10.0),
                                   ("simulation", np.array([900.0, 100.0]), 1000.0)], [0.7, 0.3], p)
        m = rows[0]["m"]
        sim = (900 + 5 * 0.7) / (1000 + 5)
        self.assertAlmostEqual(m, (8 + 5 * sim) / 15)
        rows = combine_m("email", [("anchors_strict", np.array([80.0, 20.0]), 100.0),
                                   ("simulation", np.array([0.0, 1000.0]), 1000.0)], [0.7, 0.3], p)
        self.assertAlmostEqual(rows[0]["m"], (80 + 5 * 0.7) / 105)       # simulation skipped: enough pairs
        self.assertEqual(rows[0]["m_chain"], "anchors_strict(100)")

    def test_prior_formula(self):
        p = CoreParams()
        self.assertAlmostEqual(prior_estimate(10, 1_000_000, 0.5, p)[0], 2e-5)
        v, flagged = prior_estimate(0, 1_000_000, 0.5, p)
        self.assertTrue(flagged)
        self.assertAlmostEqual(v, 1e-6)
        self.assertAlmostEqual(strict_recall({"ssn": 0.5}, lambda f, l: 0.9), 0.45)

    def test_manifest_has_every_parameter(self):
        run = _small_run()
        W = pd.concat([run["params"]["weights"]["person"], run["params"]["weights"]["business"]])
        self.assertTrue((W["m_chain"] != "").all() and W["m_pairs"].notna().all())
        self.assertTrue((W["u_source"].fillna("") != "").all())
        n_levels = sum(len(FIELD_LEVELS[f]) for p in ("person", "business") for f in PART_FIELDS[p])
        self.assertEqual(len(W), n_levels)
        self.assertTrue((run["prior_df"]["prior"] > 0).all())


class TestOutput(unittest.TestCase):
    def test_continuation_sheets(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / "t.xlsx"
            df = pd.DataFrame({"a": np.arange(25), "b": ["x"] * 25, "c": [True] * 25, "d": [np.nan] * 25})
            written = write_workbook(path, {"cand": df, "manifest": df.head(2)}, row_limit=10)
            self.assertEqual(written["cand"], ["cand", "cand (2)", "cand (3)"])
            wb = openpyxl.load_workbook(path, read_only=True)
            self.assertEqual(wb.sheetnames, ["cand", "cand (2)", "cand (3)", "manifest"])
            rows = list(wb["cand (3)"].iter_rows(values_only=True))
            self.assertEqual(rows[0], ("a", "b", "c", "d"))
            self.assertEqual(len(rows), 6)
            wb.close()

    def test_manifest_sections(self):
        cfg, maps, ref = _env()
        run = _small_run()
        sw = Stopwatch(); sw.mark("t")
        m = build_manifest(cfg, {"x": "0"}, ref, ({"source": "x", "rows": 1, "all_empty_columns": [], "added_columns": [],
                                                   "passthrough_columns": []},), run["pr"], run["params"],
                           run["prior_df"], sensitivity(pd.DataFrame({"part": [], "prior_group": [], "bits": []}),
                                                        run["prior_df"], cfg), run["scored"], run["diag"], sw)
        for s in ("config", "config.core", "reference", "m", "u", "prior", "blocking", "data_quality", "timing"):
            self.assertIn(s, set(m["section"]), s)


class TestLeieExport(unittest.TestCase):
    def test_export(self):
        raw = ('LASTNAME,FIRSTNAME,MIDNAME,BUSNAME,GENERAL,SPECIALTY,UPIN,NPI,DOB,ADDRESS,CITY,STATE,ZIP,EXCLTYPE,EXCLDATE,REINDATE,WAIVERDATE,WVRSTATE\n'
               '"DOE","JANE","Q","","PHYSICIAN (MD, DO)","FAMILY PRACTICE","","1234567893","19600102","12 W MAIN ST, APT 4","ALBANY","NY","12207","1128b4","20200101","00000000","00000000",""\n'
               '"","","","ACME DME, INC","DME COMPANY","DME - GENERAL","","1234567893","","P O BOX 7","TROY","NY","12180","1128a1","20190101","00000000","00000000",""\n')
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "raw.csv"
            p.write_text(raw, encoding="utf-8")
            out = export_leie(p)
        self.assertEqual(list(out.columns[:len(SCHEMA)]), SCHEMA)
        a, b = out.iloc[0], out.iloc[1]
        self.assertEqual((a["provider_npi"], a["clinic_npi"], b["provider_npi"], b["clinic_npi"]),
                         ("1234567893", "", "", "1234567893"))
        self.assertEqual(a["dob"], "1960-01-02")
        self.assertEqual((a["street_number"], a["street_direction"], a["street_name"], a["street_type"], a["unit"]),
                         ("12", "W", "MAIN", "ST", "4"))
        self.assertEqual(a["professional_license_type"], "PHYSICIAN (MD, DO)")
        self.assertEqual(a["x_excldate"], "2020-01-01")
        self.assertTrue(out["record_id"].is_unique and out["record_id"].str.startswith("leie:").all())


class TestInputs(unittest.TestCase):
    def test_validation(self):
        with self.assertRaises(InputError):
            validate_input(pd.DataFrame({"record_id": ["a", "a"]}), "t")
        with self.assertRaises(InputError):
            validate_input(pd.DataFrame({"first_name": ["a"]}), "t")
        df, rep = validate_input(pd.DataFrame({"record_id": ["a"], "odd": ["1"]}), "t")
        self.assertIn("x_odd", df.columns)
        self.assertIn("ssn", rep["all_empty_columns"])

    def test_mappings_valid(self):
        cfg, maps, ref = _env()
        self.assertEqual(set(ref["hash_checks"]["status"]), {"ok"})
        self.assertGreater(len(maps["category_map"]), 280)


CORE_SHA256 = "<recorded at assembly>"


NL, CR = chr(10), chr(13)


def notebook_core_hash(nb_path):
    """sha256 over the code cells between the 'Matching core v' heading and its end marker,
    exactly as recorded when the notebook was assembled."""
    nb = json.loads(Path(nb_path).read_text(encoding="utf-8"))
    inside, parts = False, []
    for c in nb["cells"]:
        src = "".join(c["source"])
        if c["cell_type"] == "markdown" and src.startswith("## Matching core v"):
            inside = True
            continue
        if c["cell_type"] == "markdown" and src.strip().startswith("*End of the matching core.*"):
            break
        if inside and c["cell_type"] == "code":
            parts.append(NL.join(l.rstrip() for l in src.replace(CR + NL, NL).split(NL)))
    return hashlib.sha256((NL + "# ---- cell ----" + NL).join(parts).encode("utf-8")).hexdigest()


class TestCoreHash(unittest.TestCase):
    def test_core_section_matches_recorded_hash(self):
        nbp = _env()[0].paths["notebook"]
        if not nbp.exists() or CORE_SHA256.startswith("<"):
            self.skipTest("runs inside the assembled notebook")
        self.assertEqual(notebook_core_hash(nbp), CORE_SHA256,
                         f"the matching core v{CORE_VERSION} changed: bump the version, re-record the hash, copy to [A]")


def run_selftests(verbosity=1):
    suite = unittest.TestSuite()
    loader = unittest.TestLoader()
    for cls in (TestNormalizers, TestBreakdown, TestCategory, TestComparisons, TestCandidates, TestScoring,
                TestParameters, TestOutput, TestLeieExport, TestInputs, TestCoreHash):
        suite.addTests(loader.loadTestsFromTestCase(cls))
    stream = io.StringIO()
    res = unittest.TextTestRunner(stream=stream, verbosity=verbosity).run(suite)
    return res, stream.getvalue()


if __name__ == "__main__":
    res, text = run_selftests(2)
    print(text)
    print(f"{res.testsRun} tests, {len(res.failures)} failures, {len(res.errors)} errors")
