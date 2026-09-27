# %% [markdown]
# ## B2 · Gold labels: review samples, SME packets, evaluation, label-driven improvements
#
# This part of the notebook turns the pipeline's output into **pairs for subject-matter experts
# (SMEs) to label**, reads the labels back, **measures** the pipeline on a sealed test split, and
# tests four **improvements that use labels**. It adopts an improvement only when it beats the
# pipeline on the sealed split. It replaces the skeleton "18 · Review sample".
#
# | Section | What it does |
# |---|---|
# | B2.0 | Configuration, the adapters to the pipeline's tables, the file IO switch |
# | B2.1 | The gold label schema |
# | B2.2 | Features from the evidence table; re-scoring with other parameters |
# | B2.3 | Review-sample generator: stratified sample, splits, inclusion weights, active learning |
# | B2.4 | SME packets: blind workbook, instructions, data dictionary, label template, adjudication sheet |
# | B2.5 | Label ingestion and validation, annotator agreement, adjudication |
# | B2.6 | Evaluation on the sealed split (Horvitz-Thompson weighted) |
# | B2.7 | Improvements: calibration, semi-supervised EM, supervised model, cost threshold; adoption |
# | B2.8 | The improvement loop: next review batch |
# | B2.9 | Synthetic annotators (for tests and the synthetic run only) |
# | B2.10 | Run |
# | B2.11 | Self-tests |
#
# **Rules** (project_v0.2/AGENTS.md 12, 13). Labels never change a runtime parameter by
# themselves: every improvement is behind a switch in `B2Config.apply`, off by default, and the
# adoption table only says whether the evidence supports switching it on. The sealed test split is
# assigned when the sample is drawn, is a pure function of the extracted party and the seed, and is
# reachable only through `B2Gold.sealed_view()`, which the evaluation calls and every fitting
# function refuses (`b2_guard_tuning`). A self-test changes every sealed label and checks that no
# fitted parameter moves. Every metric states its denominator; unsure labels, unlabelled pairs and
# strata with no sealed labels are counted, never silently dropped (rule 11).
#
# **Dependencies on the pipeline** (all through `b2_standardize`; a renamed column is one edit
# in `B2_DEPENDENCIES`):
#
# | Table (notebook variable / file) | Columns B2 reads |
# |---|---|
# | candidates (`CAND_OUT` or `PAIRS`; `out/<ds>/candidates.parquet`) | `pair_id`, `extracted_record_id`, `watchlist_record_id`, `part`, `p`, `bits`, `prior_bits`, `basis`, `veto`; optional `rank`, `rests_on_name`, a baseline probability (`p_baseline` / `baseline_p` / `p_ecm` / `ecm_p`) |
# | evidence (`EVIDENCE`; `evidence.parquet`) | `pair_id`, `field`, `level`, `m`, `u`, `bits`; optional `value_extracted`, `value_watchlist` |
# | entities (`ENTITIES`; `entities.parquet`) | `record_id`, `part`, `p`, `basis`; optional `status` |
# | manifest (`MANIFEST`; `manifest.parquet`) | `section`, `key`, `value`; optional `ci_low`, `ci_high`. Rows used: `m` (`part.field.level`, intervals), `version`/`matching_core`, `input`/`*.sha256`, `run`/`run_id` if present, `diagnostics`/`identifier_anchor_recall.*` and `truth.*` |
# | input rows (`XDF`, `WDF`; `data/<ds>_extracted.csv`, `data/<ds>_watchlist.csv`) | `record_id` and the input schema columns, for the side-by-side display only |
# | core tables (`FIELD_LEVELS`, `AGREEMENT_LEVELS`, `DISAGREEMENT_LEVELS`, `ZERO_LEVELS`) | level order and the bit clamps; a copy of core v1.0's is used when run outside the notebook |
#
# **Where Splink would do better (overall).** Splink has labelling and evaluation built in: a
# labels table (`clerical_match_score`), `truth_space_table_from_labels_table`,
# `accuracy_analysis_from_labels_table` with ROC and precision-recall charts, threshold selection
# charts, and a comparison viewer dashboard for clerical review. B2 reimplements the evaluation
# because it needs things Splink's tools do not do: Horvitz-Thompson weights for a stratified
# sample, a sealed split, several annotators with adjudication, and Excel packets for SMEs who
# never see a score.

# %% [markdown]
# ### B2.0 · Configuration, adapters and IO
#
# `B2Config` holds every B2 number (sample sizes, strata multipliers, split shares, costs,
# bootstrap size, model settings, switches); the B2 manifest section dumps it. The adapters read
# the pipeline's tables from the notebook's variables when B2 runs inside the notebook, or from
# `out/<dataset>/` otherwise, and rename columns to B2's names. `b2_read_table` /
# `b2_write_table` read and write CSV, XLSX, Parquet and Delta by extension; set
# `B2_IO["read"]` / `B2_IO["write"]` to the pipeline's own IO functions to route everything
# through its switch instead.
#
# **Where Splink would do better.** Splink's linker reads and writes through its database
# backend (DuckDB, Spark, Athena, Postgres), so a labels table is just another table in the same
# engine. Here labels travel as files because SMEs work in Excel.

# %% [code]
# ---- B2.0 configuration, adapters, IO --------------------------------------------------------
import os, json, math, hashlib, datetime as _dt, re as _re, textwrap as _tw
from dataclasses import dataclass, field as _field, asdict as _asdict
from pathlib import Path
import numpy as np
import pandas as pd

B2_VERSION = "1.0"
B2_LABELS = ("match", "non-match", "unsure")
B2_CONFIDENCE = ("high", "medium", "low")
B2_SPLITS = ("train", "calibration", "sealed_test")
B2_IMPROVEMENTS = ("calibration", "semisup_em", "supervised", "cost_threshold")
# Accepted spellings, including the gold-annotator's watchlist vocabulary (same / different / cant_tell).
B2_LABEL_SYNONYMS = {
    "match": "match", "same": "match", "yes": "match", "y": "match", "1": "match", "true": "match",
    "non-match": "non-match", "nonmatch": "non-match", "non match": "non-match", "not a match": "non-match",
    "no match": "non-match", "different": "non-match", "no": "non-match", "n": "non-match", "0": "non-match",
    "false": "non-match",
    "unsure": "unsure", "cant_tell": "unsure", "can't tell": "unsure", "cannot tell": "unsure",
    "not sure": "unsure", "?": "unsure",
}


@dataclass
class B2Config:
    seed: int = 20260926
    # review sample (B2.3)
    sample_size: int = 1500
    min_per_stratum: int = 10
    band_edges: tuple = (0.1, 0.5, 0.8, 0.9)        # bands p<0.1, 0.1-0.5, 0.5-0.8, 0.8-0.9, p>=0.9; plus "vetoed"
    uncertain_bands: tuple = ("0.1-0.5", "0.5-0.8", "0.8-0.9")
    uncertain_multiplier: float = 4.0                # allocation ~ sqrt(N_h) x multipliers
    name_bases: tuple = ("name_only", "contextual")
    name_multiplier: float = 2.0
    vetoed_multiplier: float = 0.5
    split_fractions: tuple = (("train", 0.4), ("calibration", 0.2), ("sealed_test", 0.4))
    annotators: tuple = ("A1", "A2", "A3")
    double_code_share: float = 0.25                  # share of the other pairs labelled by two annotators
    double_code_splits: tuple = ("sealed_test", "calibration")   # every pair of these splits by two (+ adjudication)
    masked_fields: tuple = ()                        # display fields shown as last four only, e.g. ("SSN",)
    # evaluation (B2.6)
    b1_threshold: float = 0.5                        # the pipeline's operating point ("p >= 0.5")
    thresholds: tuple = (0.05, 0.1, 0.2, 0.3, 0.5, 0.7, 0.8, 0.9, 0.95, 0.99)
    reliability_edges: tuple = (0.0, 0.001, 0.01, 0.1, 0.3, 0.5, 0.7, 0.9, 0.99, 0.999, 1.0)
    cost_false_alarm: float = 1.0                    # (d): cost of flagging a non-match
    cost_miss: float = 5.0                           # (d): cost of missing a true watchlist match
    bootstrap_reps: int = 1000
    adopt_alpha: float = 0.05                        # family-wise; Bonferroni over the improvements tested
    min_labels_to_fit: int = 30
    # improvements (B2.7)
    calibration_methods: tuple = ("platt", "isotonic")
    calibration_folds: int = 5
    em_alpha: float = 5.0                            # pseudo-pairs toward the pipeline's m (as the core's alpha)
    em_max_iter: int = 200
    em_tol: float = 1e-6
    gb_params: dict = _field(default_factory=lambda: dict(n_estimators=200, max_depth=3, learning_rate=0.05,
                                                          subsample=0.8, min_samples_leaf=10))
    # improvement loop (B2.8)
    al_batch_size: int = 300
    al_per_entity: int = 1
    al_disagreement_weight: float = 1.0
    prior_shift_bits: float = 3.3219                 # parameter-source spread: prior x10 and /10
    # switches: labels change nothing until a person turns one on (AGENTS rule 13)
    apply: dict = _field(default_factory=lambda: {k: False for k in B2_IMPROVEMENTS})
    apply_order: tuple = ("supervised", "semisup_em", "calibration")


# ---- adapters ---------------------------------------------------------------------------------
B2_DEPENDENCIES = {
    "candidates": {
        "required": {"pair_id": ("pair_id",), "extracted_record_id": ("extracted_record_id", "x_record_id"),
                     "watchlist_record_id": ("watchlist_record_id", "w_record_id"), "part": ("part",),
                     "p": ("p",), "bits": ("bits",), "prior_bits": ("prior_bits", "prior_logit"),
                     "basis": ("basis",), "veto": ("veto",)},
        "optional": {"rank": ("rank",), "rests_on_name": ("rests_on_name",),
                     "p_baseline": ("p_baseline", "baseline_p", "p_ecm", "ecm_p")}},
    "evidence": {
        "required": {"pair_id": ("pair_id",), "field": ("field",), "level": ("level",), "m": ("m",),
                     "u": ("u", "u_used"), "bits": ("bits",)},
        "optional": {"value_extracted": ("value_extracted",), "value_watchlist": ("value_watchlist",)}},
    "entities": {
        "required": {"record_id": ("record_id",), "part": ("part",), "p": ("p",), "basis": ("basis",)},
        "optional": {"status": ("status",)}},
    "manifest": {
        "required": {"section": ("section",), "key": ("key",), "value": ("value",)},
        "optional": {"source": ("source",), "ci_low": ("ci_low",), "ci_high": ("ci_high",), "note": ("note",)}},
    "extracted_rows": {"required": {"record_id": ("record_id",)}, "optional": {}},
    "watchlist_rows": {"required": {"record_id": ("record_id",)}, "optional": {}},
}


def b2_standardize(df, table):
    """Rename the pipeline's columns to B2's names; raise naming every missing column."""
    spec = B2_DEPENDENCIES[table]
    ren, missing = {}, []
    for canon, names in spec["required"].items():
        hit = next((n for n in names if n in df.columns), None)
        if hit is None:
            missing.append(f"{canon} (looked for {', '.join(names)})")
        elif hit != canon:
            ren[hit] = canon
    for canon, names in spec.get("optional", {}).items():
        hit = next((n for n in names if n in df.columns), None)
        if hit is not None and hit != canon and canon not in df.columns:
            ren[hit] = canon
    if missing:
        raise KeyError(f"B2 adapter: table '{table}' lacks {'; '.join(missing)}")
    return df.rename(columns=ren)


# ---- IO switch --------------------------------------------------------------------------------
B2_IO = {"read": None, "write": None}      # set to the pipeline's IO functions to route through its switch


def _b2_is_delta(path):
    return Path(path).is_dir() and (Path(path) / "_delta_log").exists()


def b2_read_table(path, sheet=None):
    """CSV / XLSX / Parquet / Delta by extension. Text columns stay text (no NA guessing)."""
    if B2_IO["read"] is not None:
        return B2_IO["read"](path) if sheet is None else B2_IO["read"](path, sheet=sheet)
    path = Path(path)
    ext = path.suffix.lower()
    if _b2_is_delta(path) or ext == ".delta":
        try:
            import deltalake
        except ImportError as ex:
            raise ImportError("reading Delta needs the 'deltalake' package (pip install deltalake), "
                              "or set B2_IO['read'] to the pipeline's reader") from ex
        return deltalake.DeltaTable(str(path)).to_pandas()
    if ext == ".csv":
        return pd.read_csv(path, dtype=str, keep_default_na=False, encoding="utf-8-sig")
    if ext in (".xlsx", ".xlsm"):
        return pd.read_excel(path, sheet_name=sheet if sheet is not None else 0, dtype=object)
    if ext == ".parquet":
        return pd.read_parquet(path)
    raise ValueError(f"unsupported table file: {path}")


def b2_write_table(df, path):
    if B2_IO["write"] is not None:
        return B2_IO["write"](df, path)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    ext = path.suffix.lower()
    if ext == ".csv":
        df.to_csv(path, index=False, encoding="utf-8")
    elif ext == ".parquet":
        d = df.copy()
        for c in d.columns:
            if d[c].dtype == object:
                d[c] = d[c].map(lambda v: "" if v is None else (v if isinstance(v, str) else str(v)))
        d.to_parquet(path, index=False)
    elif ext == ".delta":
        import deltalake
        deltalake.write_deltalake(str(path), df, mode="overwrite")
    elif ext == ".xlsx":
        df.to_excel(path, index=False)
    else:
        raise ValueError(f"unsupported table file: {path}")
    return path


def b2_find_table(folder, name):
    """The first existing of name.parquet, name (Delta folder), name.delta, name.csv."""
    folder = Path(folder)
    for cand in (folder / f"{name}.parquet", folder / name, folder / f"{name}.delta", folder / f"{name}.csv"):
        if cand.exists() and (cand.is_file() or _b2_is_delta(cand)):
            return cand
    raise FileNotFoundError(f"no {name} table in {folder} (looked for .parquet, Delta, .csv)")


def b2_root():
    g = globals()
    cfg = g.get("CFG")
    if cfg is not None and getattr(cfg, "root", ""):
        return Path(cfg.root)
    if "__file__" in g:
        return Path(g["__file__"]).resolve().parent
    here = Path.cwd()
    for p in [here, *here.parents]:
        if (p / "mappings" / "columns.csv").exists():
            return p
    return here


@dataclass
class B2Run:
    dataset: str
    root: Path
    candidates: pd.DataFrame
    evidence: pd.DataFrame
    entities: pd.DataFrame
    manifest: pd.DataFrame
    xrows: pd.DataFrame
    wrows: pd.DataFrame
    run_id: str
    core_version: str
    truth_path: object
    source: str


def _b2_manifest_value(man, section, key, default=""):
    hit = man[(man["section"] == section) & (man["key"] == key)]
    return hit["value"].iloc[0] if len(hit) else default


def b2_load_run(dataset=None, root=None, use_notebook=True):
    """The pipeline's output tables: from the notebook's variables when present, else from files."""
    g = globals()
    root = Path(root) if root else b2_root()
    in_nb = use_notebook and g.get("EVIDENCE") is not None and (g.get("CAND_OUT") is not None or g.get("PAIRS") is not None)
    if in_nb:
        ds = getattr(g.get("CFG"), "dataset", dataset or "synthetic")
        cand = g["CAND_OUT"] if g.get("CAND_OUT") is not None else g["PAIRS"]
        ev, ent, man = g["EVIDENCE"], g.get("ENTITIES"), g.get("MANIFEST")
        xr, wr, truth = g.get("XDF"), g.get("WDF"), g.get("TF")
        source = "notebook variables"
    else:
        ds = dataset or os.environ.get("RL_DATASET", "synthetic")
        out = root / "out" / ds
        cand = b2_read_table(b2_find_table(out, "candidates"))
        ev = b2_read_table(b2_find_table(out, "evidence"))
        ent = b2_read_table(b2_find_table(out, "entities"))
        man = b2_read_table(b2_find_table(out, "manifest"))
        data = root / "data"
        if ds == "files":
            xf, wf, tf = Path(os.environ["RL_EXTRACTED"]), Path(os.environ["RL_WATCHLIST"]), None
        else:
            xf, wf, tf = data / f"{ds}_extracted.csv", data / f"{ds}_watchlist.csv", data / f"{ds}_truth.csv"
        xr, wr = b2_read_table(xf), b2_read_table(wf)
        truth = tf if tf is not None and Path(tf).exists() else None
        source = f"files in {out}"
    cand = b2_standardize(cand.drop(columns=[c for c in ("_l", "_r") if c in cand.columns]), "candidates")
    ev = b2_standardize(ev, "evidence")
    ent = b2_standardize(ent, "entities") if ent is not None else pd.DataFrame(columns=["record_id", "part", "p", "basis"])
    man = b2_standardize(man, "manifest") if man is not None else pd.DataFrame(columns=["section", "key", "value"])
    xr = b2_standardize(xr.fillna("").astype(str), "extracted_rows")
    wr = b2_standardize(wr.fillna("").astype(str), "watchlist_rows")
    cand = cand.copy()
    cand["veto"] = cand["veto"].fillna("").astype(str)
    for c in ("p", "bits", "prior_bits"):
        cand[c] = pd.to_numeric(cand[c], errors="coerce").astype(float)
    if cand["pair_id"].duplicated().any():
        raise ValueError("B2: candidates.pair_id is not unique")
    core = str(_b2_manifest_value(man, "version", "matching_core", g.get("CORE_VERSION", "")))
    run_id = str(_b2_manifest_value(man, "run", "run_id", ""))
    if not run_id:
        hashes = man[(man["section"] == "input") & man["key"].astype(str).str.endswith(".sha256")]["value"].astype(str)
        basis = "|".join(sorted(hashes)) + "|" + core + "|" + ds
        run_id = "run-" + hashlib.sha256(basis.encode()).hexdigest()[:12]
    return B2Run(ds, root, cand.reset_index(drop=True), ev, ent, man, xr, wr, run_id, core,
                 Path(truth) if truth is not None and Path(truth).exists() else None, source)


def b2_core_defs():
    """Level order and clamps: the notebook's core tables, else a copy of core v1.0's."""
    g = globals()
    names = ("FIELD_LEVELS", "AGREEMENT_LEVELS", "DISAGREEMENT_LEVELS", "ZERO_LEVELS")
    if all(n in g for n in names):
        return {n: g[n] for n in names}, "notebook matching core"
    return B2_CORE_FALLBACK, "B2's copy of matching core v1.0 level tables"


B2_CORE_FALLBACK = {
    "FIELD_LEVELS": {
        "name": ["exact", "first_nick_or_close", "last_close_first_agrees", "initial_agrees", "swapped",
                 "first_empty", "first_differs", "else"],
        "middle": ["exact", "initial", "differs"],
        "org": ["exact", "dba", "short_form", "rare_shared", "common_shared", "sibling", "none"],
        "dob": ["exact", "swap_or_typo", "year_month", "year", "differs"],
        "address": ["exact", "street", "zip", "city_state", "state", "differs"],
        "ssn": ["exact", "near", "differs"], "npi": ["exact", "near", "differs"], "dl": ["exact", "near", "differs"],
        "tin": ["exact", "near", "differs"], "license": ["exact", "differs"], "email": ["exact", "differs"],
        "vin": ["exact", "differs"], "plate": ["exact", "differs"], "cnpi": ["exact", "differs"],
        "phone": ["exact_owned_single", "exact_shared", "differs"],
        "spec_cat": ["specialty", "category_id", "category_weak", "differs"],
        "co_party": ["anchored", "not_anchored"]},
    "AGREEMENT_LEVELS": {
        "name": {"exact", "first_nick_or_close", "last_close_first_agrees", "initial_agrees", "swapped"},
        "middle": {"exact", "initial"}, "org": {"exact", "dba", "short_form", "rare_shared"},
        "dob": {"exact", "swap_or_typo"}, "address": {"exact", "street"},
        "ssn": {"exact", "near"}, "npi": {"exact", "near"}, "dl": {"exact", "near"},
        "tin": {"exact", "near"}, "license": {"exact"}, "email": {"exact"}, "vin": {"exact"},
        "plate": {"exact"}, "cnpi": {"exact"}, "phone": {"exact_owned_single", "exact_shared"},
        "spec_cat": {"specialty", "category_id", "category_weak"}, "co_party": {"anchored"}},
    "DISAGREEMENT_LEVELS": {
        "name": {"else"}, "middle": {"differs"}, "org": {"sibling", "none"}, "dob": {"differs"},
        "address": {"differs"}, "ssn": {"differs"}, "npi": {"differs"}, "dl": {"differs"},
        "tin": {"differs"}, "license": {"differs"}, "email": {"differs"}, "vin": {"differs"},
        "plate": {"differs"}, "cnpi": {"differs"}, "phone": {"differs"}, "spec_cat": {"differs"}},
    "ZERO_LEVELS": {"co_party": {"not_anchored"}},
}


def b2_hash_unit(*parts):
    """A deterministic number in [0, 1) from strings (split and assignment draws)."""
    h = hashlib.sha256("|".join(str(p) for p in parts).encode()).digest()
    return int.from_bytes(h[:8], "big") / 2.0 ** 64


def b2_config_rows(cfg):
    return [{"section": "b2.config", "key": k, "value": json.dumps(v, default=str), "source": "B2Config"}
            for k, v in _asdict(cfg).items()]


# %% [markdown]
# ### B2.1 · The gold label schema
#
# Three tables. **Sample key** (the coordinator's copy, one row per pair put up for review):
# the pair's identity, the run and core version it came from, the stratum, the pipeline's `p`
# and basis at sampling time, the population and sample counts of its stratum, the inclusion
# probability and weight, the selection route (`stratified` or `active`), the split, the round
# and the seed. **Labels** (one row per annotator and pair): what the SME returned plus where it
# came from. **Final labels** (one row per pair): the resolved label after agreement or
# adjudication, with the number of annotators, their labels and the adjudicator. The pair is
# identified by extracted `record_id` + part and watchlist `record_id` + part; the SME sees only
# an opaque `review_id`.
#
# **Where Splink would do better.** Splink's labels table is two record ids and a
# `clerical_match_score` in [0, 1]; it can also infer labels from a column that already
# identifies clusters (`cluster_col`). It has no fields for annotators, confidence, splits or
# weights: those are B2's.

# %% [code]
# ---- B2.1 gold label schema -------------------------------------------------------------------
B2_KEY_SCHEMA = [
    ("review_id", "str", "opaque id shown to the SME; unique per pair and round"),
    ("pair_id", "str", "the pipeline's pair id"),
    ("extracted_record_id", "str", "extracted row"), ("extracted_part", "str", "person | business"),
    ("watchlist_record_id", "str", "watchlist row"), ("watchlist_part", "str", "person | business"),
    ("run_id", "str", "pipeline run the pair came from"), ("core_version", "str", "matching core version"),
    ("b2_version", "str", "B2 version"), ("round", "int", "review round (0 = first stratified sample)"),
    ("selection", "str", "stratified | active"), ("seed", "int", "sampling seed"),
    ("stratum", "str", "band|basis at sampling time"), ("band", "str", "p band at sampling time or 'vetoed'"),
    ("basis", "str", "pipeline basis at sampling time"), ("p_at_sampling", "float", "pipeline p at sampling time"),
    ("rank", "float", "rank of the pair within its extracted party"),
    ("N_h", "int", "candidate pairs in the stratum (population)"), ("n_h", "int", "pairs drawn from the stratum"),
    ("inclusion_prob", "float", "n_h / N_h (stratified); empty for active selection"),
    ("weight", "float", "1 / inclusion_prob; empty for active selection"),
    ("split", "str", "train | calibration | sealed_test (from the extracted party and the seed)"),
    ("group", "str", "extracted record_id|part: the unit the split is assigned to"),
    ("annotators", "str", "annotator ids the pair was assigned to, ';'-joined"),
    ("double_coded", "bool", "assigned to two annotators"),
]
B2_LABEL_SCHEMA = [
    ("review_id", "str", "as in the sample key"), ("annotator_id", "str", "who labelled"),
    ("label", "str", "match | non-match | unsure"), ("confidence", "str", "high | medium | low (optional)"),
    ("reason", "str", "free text; required for unsure"),
    ("labelled_at", "str", "ISO 8601; from the file's modified time when the cell is empty"),
    ("timestamp_source", "str", "cell | file"), ("source_file", "str", "file the label was read from"),
    ("source_row", "int", "row in that file"),
]
B2_FINAL_SCHEMA = [
    ("review_id", "str", ""), ("pair_id", "str", ""), ("final_label", "str", "match | non-match | unsure"),
    ("y", "float", "1 match, 0 non-match, empty unsure"), ("n_annotators", "int", ""),
    ("labels", "str", "annotator:label pairs"), ("unanimous", "bool", ""),
    ("adjudicated", "bool", ""), ("adjudicator_id", "str", ""), ("adjudication_reason", "str", ""),
    ("label_source", "str", "single | unanimous | adjudicated | unresolved"),
]


def b2_schema_frame():
    rows = []
    for table, sch in (("sample_key", B2_KEY_SCHEMA), ("labels", B2_LABEL_SCHEMA), ("final_labels", B2_FINAL_SCHEMA)):
        rows += [{"table": table, "column": c, "type": t, "meaning": m} for c, t, m in sch]
    return pd.DataFrame(rows)


# %% [markdown]
# ### B2.2 · Features from the evidence table, and re-scoring
#
# The evidence table has one row per pair and non-empty field with the level, the m and the u
# the pipeline used (value-specific where the level has one) and the bits. B2 pivots it into one
# level, u, m and bits array per field, and can **re-score** every candidate with other m
# values or a prior shift using the core's own rules (log2(m/u); agreement never below 0,
# disagreement never above 0, empty = 0; vetoed pairs p = 0). With the pipeline's own m the
# re-score reproduces the pipeline's bits (a self-test). B2's models change p, never the basis:
# a name-only pair stays labelled name-only whatever its p.
#
# **Where Splink would do better.** Splink's `predict()` re-scores from a saved settings JSON
# in SQL for any parameter set, and its waterfall chart shows each comparison's contribution;
# here the re-score is a numpy pass over the evidence table.

# %% [code]
# ---- B2.2 features and re-scoring -------------------------------------------------------------
@dataclass
class B2Features:
    pair_id: np.ndarray
    part: np.ndarray
    fields: list
    level: dict
    u: dict
    m: dict
    bits: dict
    vetoed: np.ndarray
    prior_bits: np.ndarray
    p: np.ndarray


def b2_features(run):
    c, ev = run.candidates, run.evidence
    n = len(c)
    pos = pd.Index(c["pair_id"]).get_indexer(ev["pair_id"])
    ok = pos >= 0
    fields = sorted(ev.loc[ok, "field"].astype(str).unique())
    fcol = ev["field"].astype(str).to_numpy()
    lev_all = ev["level"].astype(str).to_numpy(dtype=object)
    u_all = pd.to_numeric(ev["u"], errors="coerce").to_numpy(float)
    m_all = pd.to_numeric(ev["m"], errors="coerce").to_numpy(float)
    b_all = pd.to_numeric(ev["bits"], errors="coerce").to_numpy(float)
    level, u, m, bits = {}, {}, {}, {}
    for f in fields:
        sel = ok & (fcol == f)
        p_ = pos[sel]
        level[f] = np.full(n, "", dtype=object); level[f][p_] = lev_all[sel]
        u[f] = np.full(n, np.nan); u[f][p_] = u_all[sel]
        m[f] = np.full(n, np.nan); m[f][p_] = m_all[sel]
        bits[f] = np.zeros(n); bits[f][p_] = b_all[sel]
    return B2Features(c["pair_id"].to_numpy(dtype=object), c["part"].astype(str).to_numpy(dtype=object), fields,
                      level, u, m, bits, (c["veto"] != "").to_numpy(), c["prior_bits"].to_numpy(float),
                      c["p"].to_numpy(float))


def b2_m_table(run, feats):
    """m per part, field and level: the manifest's rows (with intervals) where present, the
    evidence table's values otherwise."""
    rows = []
    for f in feats.fields:
        has = feats.level[f] != ""
        if has.any():
            d = pd.DataFrame({"part": feats.part[has], "field": f, "level": feats.level[f][has], "m": feats.m[f][has]})
            rows.append(d.drop_duplicates(["part", "field", "level"]))
    ev_m = pd.concat(rows, ignore_index=True) if rows else pd.DataFrame(columns=["part", "field", "level", "m"])
    man = run.manifest
    mm = man[man["section"] == "m"].copy()
    if len(mm):
        k = mm["key"].astype(str).str.split(".", n=2, expand=True)
        mm = pd.DataFrame({"part": k[0], "field": k[1], "level": k[2],
                           "m": pd.to_numeric(mm["value"], errors="coerce").to_numpy(),
                           "m_lo": pd.to_numeric(mm.get("ci_low", pd.Series(np.nan, index=mm.index)), errors="coerce").to_numpy(),
                           "m_hi": pd.to_numeric(mm.get("ci_high", pd.Series(np.nan, index=mm.index)), errors="coerce").to_numpy()})
        out = mm.merge(ev_m[["part", "field", "level"]].assign(seen=True), on=["part", "field", "level"], how="outer")
        miss = out["m"].isna()
        if miss.any():
            fill = ev_m.set_index(["part", "field", "level"])["m"]
            out.loc[miss, "m"] = [fill.get((a, b, c_), np.nan) for a, b, c_ in
                                  zip(out.loc[miss, "part"], out.loc[miss, "field"], out.loc[miss, "level"])]
        out["seen"] = out["seen"].fillna(False).astype(bool)
        return out
    ev_m["m_lo"], ev_m["m_hi"], ev_m["seen"] = np.nan, np.nan, True
    return ev_m


def _b2_lookup(feats, f, table):
    """m for each pair's (part, level) of field f from {(part, field, level): m}; NaN if absent."""
    keys = list(zip(feats.part, feats.level[f]))
    return np.array([table.get((pt, f, lv), np.nan) if lv != "" else np.nan for pt, lv in keys], dtype=float)


def b2_rescore(feats, m_lookup=None, prior_shift_bits=0.0):
    """Re-score every candidate: per field log2(m/u) with the core's clamps. `m_lookup`:
    {(part, field, level): m}; missing entries keep the pipeline's m. `prior_shift_bits`: scalar
    or per-pair array added to the prior. Returns (p, total bits, {field: bits})."""
    defs, _ = b2_core_defs()
    n = len(feats.pair_id)
    total = np.zeros(n)
    per = {}
    for f in feats.fields:
        lev = feats.level[f]
        has = lev != ""
        m = feats.m[f] if m_lookup is None else np.where(np.isnan(mv := _b2_lookup(feats, f, m_lookup)), feats.m[f], mv)
        b = np.log2(np.clip(np.nan_to_num(m, nan=1.0), 1e-15, 1.0) / np.clip(np.nan_to_num(feats.u[f], nan=1.0), 1e-15, 1.0))
        agree = np.isin(lev, list(defs["AGREEMENT_LEVELS"].get(f, ())))
        disagree = np.isin(lev, list(defs["DISAGREEMENT_LEVELS"].get(f, ())))
        zero = np.isin(lev, list(defs["ZERO_LEVELS"].get(f, ())))
        b = np.where(agree, np.maximum(b, 0.0), b)
        b = np.where(disagree, np.minimum(b, 0.0), b)
        b = np.where(zero | ~has, 0.0, b)
        per[f] = b
        total += b
    logit = feats.prior_bits + prior_shift_bits + total
    p = 1.0 / (1.0 + np.exp2(-np.clip(logit, -1000, 1000)))
    p = np.where(feats.vetoed, 0.0, p)
    return p, total, per


def b2_parameter_spread(run, feats, cfg):
    """Probabilities under other defensible parameter sets, for active learning: m at the
    manifest's 95% bounds (pessimistic: agreement at the low bound, disagreement at the high
    bound; optimistic: the reverse) and the prior x10 and /10."""
    defs, _ = b2_core_defs()
    mt = b2_m_table(run, feats)
    out = {"B1": feats.p.copy()}
    if mt["m_lo"].notna().any():
        agree = np.array([lv in defs["AGREEMENT_LEVELS"].get(f, set()) for f, lv in zip(mt["field"], mt["level"])])
        lo = mt["m_lo"].fillna(mt["m"]).to_numpy(); hi = mt["m_hi"].fillna(mt["m"]).to_numpy()
        pess = dict(zip(zip(mt["part"], mt["field"], mt["level"]), np.where(agree, lo, hi)))
        opt = dict(zip(zip(mt["part"], mt["field"], mt["level"]), np.where(agree, hi, lo)))
        out["m_pessimistic"] = b2_rescore(feats, pess)[0]
        out["m_optimistic"] = b2_rescore(feats, opt)[0]
    out["prior_x10"] = b2_rescore(feats, None, cfg.prior_shift_bits)[0]
    out["prior_div10"] = b2_rescore(feats, None, -cfg.prior_shift_bits)[0]
    return out


# %% [markdown]
# ### B2.3 · Review-sample generator
#
# **Strata** = p band × basis, from the pipeline's candidates (every kept pair; vetoed pairs form
# their own band, so a wrong veto can be found). Each stratum gets a share of the sample
# proportional to sqrt(N_h), multiplied by `uncertain_multiplier` for the bands 0.1-0.9 and by
# `name_multiplier` for name-only and contextual pairs; at least `min_per_stratum`, at most N_h.
# Pairs are drawn at random within the stratum with the seed; the **inclusion probability**
# n_h / N_h and **weight** N_h / n_h are recorded.
#
# **Splits** are assigned when the sample is drawn, per *extracted party* (record_id + part), by a
# hash of the party and the seed: `train` 40%, `calibration` 20%, `sealed_test` 40%. All pairs of
# one party share a split (no leakage through a party seen in training), and a party keeps its
# split in every later round. The split is in the coordinator's key only.
#
# **Active learning** (B2.8) picks the most informative unlabelled pairs: highest entropy of the
# mean p across the available scorers plus their spread (the pipeline's p under the m intervals
# and the prior sensitivity; after labels, the calibrated, EM and supervised p). Active pairs have
# no inclusion probability, so they go to `train` only and never enter a weighted metric; they are
# drawn only from parties whose split is `train`.
#
# **Where Splink would do better.** Splink has no sampler for clerical review; its comparison
# viewer dashboard shows a sample of pairs per comparison-vector pattern (gamma pattern), which is
# a useful stratum B2 could add. It has no notion of a sealed split.

# %% [code]
# ---- B2.3 review sample --------------------------------------------------------------------------
def b2_band(p, vetoed, cfg):
    e = cfg.band_edges
    labels = [f"p<{e[0]}"] + [f"{e[i]}-{e[i + 1]}" for i in range(len(e) - 1)] + [f"p>={e[-1]}"]
    idx = np.searchsorted(np.asarray(e), p, side="right")
    return np.where(vetoed, "vetoed", np.asarray(labels, dtype=object)[idx])


def b2_split_of(groups, seed, fractions):
    cum = np.cumsum([f for _, f in fractions])
    names = [s for s, _ in fractions]
    u = np.array([b2_hash_unit("split", seed, g) for g in groups])
    return np.asarray(names, dtype=object)[np.minimum(np.searchsorted(cum, u, side="right"), len(names) - 1)]


def b2_frame(run, cfg):
    """The sampling frame: every candidate pair with its band, stratum and group."""
    c = run.candidates
    vet = (c["veto"] != "").to_numpy()
    f = pd.DataFrame({"pair_id": c["pair_id"], "extracted_record_id": c["extracted_record_id"].astype(str),
                      "part": c["part"].astype(str), "watchlist_record_id": c["watchlist_record_id"].astype(str),
                      "p": c["p"].to_numpy(float), "basis": c["basis"].astype(str),
                      "rank": pd.to_numeric(c["rank"], errors="coerce") if "rank" in c else np.nan})
    f["band"] = b2_band(f["p"].to_numpy(), vet, cfg)
    f["stratum"] = f["band"] + "|" + f["basis"]
    f["group"] = f["extracted_record_id"] + "|" + f["part"]
    return f


def b2_allocate(N, cfg, total):
    """Sample size per stratum. N: Series stratum -> population count."""
    band = N.index.str.split("|").str[0]
    basis = N.index.str.split("|").str[1]
    mult = np.ones(len(N))
    mult *= np.where(np.isin(band, cfg.uncertain_bands), cfg.uncertain_multiplier, 1.0)
    mult *= np.where(np.isin(basis, cfg.name_bases), cfg.name_multiplier, 1.0)
    mult *= np.where(band == "vetoed", cfg.vetoed_multiplier, 1.0)
    score = np.sqrt(N.to_numpy(float)) * mult
    n = np.minimum(N.to_numpy(), np.maximum(cfg.min_per_stratum, np.round(total * score / score.sum()))).astype(int)
    for _ in range(50):                          # spread what the caps freed over the strata still open
        gap = total - n.sum()
        open_ = n < N.to_numpy()
        if gap <= 0 or not open_.any():
            break
        add = np.floor(gap * score * open_ / (score * open_).sum()).astype(int)
        if add.sum() == 0:
            add[np.flatnonzero(open_)[np.argmax(score[open_])]] = 1
        n = np.minimum(N.to_numpy(), n + add)
    return pd.Series(n, index=N.index)


def _b2_review_id(round_id, seed, pair_id):
    return f"R{round_id:02d}-" + hashlib.sha256(f"{seed}|{pair_id}".encode()).hexdigest()[:8].upper()


def b2_draw_sample(run, cfg, round_id=0, exclude=(), size=None):
    """Stratified sample (see above). `exclude`: pair ids already put up for review."""
    fr = b2_frame(run, cfg)
    fr = fr[~fr["pair_id"].isin(set(exclude))].sort_values("pair_id", kind="stable").reset_index(drop=True)
    N = fr.groupby("stratum").size().sort_index()
    n = b2_allocate(N, cfg, size or cfg.sample_size)
    rng = np.random.default_rng([cfg.seed, round_id])
    picks = []
    for s in N.index:
        rows = np.flatnonzero(fr["stratum"].to_numpy() == s)
        k = int(n[s])
        if k > 0:
            picks.append(np.sort(rng.choice(rows, size=k, replace=False)))
    take = fr.iloc[np.concatenate(picks)].copy() if picks else fr.iloc[:0].copy()
    take["N_h"] = take["stratum"].map(N).astype(int)
    take["n_h"] = take["stratum"].map(n).astype(int)
    take["inclusion_prob"] = take["n_h"] / take["N_h"]
    take["weight"] = 1.0 / take["inclusion_prob"]
    take["selection"] = "stratified"
    return _b2_key_rows(take, run, cfg, round_id)


def _b2_key_rows(take, run, cfg, round_id):
    k = pd.DataFrame({
        "review_id": [_b2_review_id(round_id, cfg.seed, p) for p in take["pair_id"]],
        "pair_id": take["pair_id"].to_numpy(), "extracted_record_id": take["extracted_record_id"].to_numpy(),
        "extracted_part": take["part"].to_numpy(), "watchlist_record_id": take["watchlist_record_id"].to_numpy(),
        "watchlist_part": take["part"].to_numpy(), "run_id": run.run_id, "core_version": run.core_version,
        "b2_version": B2_VERSION, "round": round_id, "selection": take["selection"].to_numpy(), "seed": cfg.seed,
        "stratum": take["stratum"].to_numpy(), "band": take["band"].to_numpy(), "basis": take["basis"].to_numpy(),
        "p_at_sampling": take["p"].to_numpy(), "rank": take["rank"].to_numpy(),
        "N_h": take["N_h"].to_numpy(), "n_h": take["n_h"].to_numpy(),
        "inclusion_prob": take["inclusion_prob"].to_numpy(), "weight": take["weight"].to_numpy(),
        "group": take["group"].to_numpy()})
    k["split"] = b2_split_of(k["group"], cfg.seed, cfg.split_fractions)
    if (k["selection"] == "active").any():
        assert (k.loc[k["selection"] == "active", "split"] == "train").all(), "active pairs must be train"
    if k["review_id"].duplicated().any():
        raise ValueError("B2: review_id collision; change the seed")
    return k.reset_index(drop=True)


def b2_assign(key, cfg, annotators=None):
    """Annotator assignment: every pair of the `double_code_splits` and a
    `double_code_share` of the others (by hash) to two annotators, the rest to one, balanced
    round-robin in hash order. Returns (key with annotators, assignments)."""
    ann = list(annotators or cfg.annotators)
    k = key.copy()
    order = np.argsort([b2_hash_unit("order", cfg.seed, r) for r in k["review_id"]], kind="stable")
    dbl = np.array([b2_hash_unit("double", cfg.seed, r) < cfg.double_code_share for r in k["review_id"]])
    if cfg.double_code_splits and "split" in k.columns:
        dbl |= k["split"].isin(cfg.double_code_splits).to_numpy()
    if len(ann) < 2:
        dbl[:] = False
    rows, j = [], 0
    names = [""] * len(k)
    for i in order:
        who = [ann[j % len(ann)]]
        j += 1
        if dbl[i]:
            who.append(ann[j % len(ann)])
            j += 1
        names[i] = ";".join(who)
        rows += [{"review_id": k["review_id"].iloc[i], "annotator_id": a} for a in who]
    k["annotators"] = names
    k["double_coded"] = dbl
    asg = pd.DataFrame(rows, columns=["review_id", "annotator_id"])
    asg["_h"] = [b2_hash_unit("pos", cfg.seed, a, r) for a, r in zip(asg["annotator_id"], asg["review_id"])]
    asg["position"] = asg.groupby("annotator_id")["_h"].rank(method="first").astype(int)
    return k, asg.drop(columns="_h").sort_values(["annotator_id", "position"]).reset_index(drop=True)


# %% [markdown]
# ### B2.4 · SME packets
#
# One folder per annotator with: the **blind review workbook** (sheets *Instructions*, *Persons*,
# *Businesses*, *Data dictionary*; one row per pair with the extracted and the watchlist values
# side by side, and `label` / `confidence` / `reason` / `labelled_on` cells with drop-down
# validation; the data cells are locked, the label cells are not), the **instructions** as a
# Markdown file, the **data dictionary** as CSV, and a **label template** CSV for anyone who
# returns labels without Excel. The reviewer sees only an opaque `review_id` and the input values:
# no p, basis, bits, stratum, split, weight or record id. Double-coded pairs appear in both
# annotators' files in different orders. The **coordinator's copy** holds the sample key (strata,
# p, weights, split, seed), the assignments, a blind flat file for an app's review queue, and a
# packet manifest with a SHA-256 per file. The **adjudication sheet** lists every pair the
# annotators disagreed on (or marked unsure) with their labels and reasons, anonymised as
# Reviewer 1/2, and cells for the final label.
#
# **Where Splink would do better.** Splink's comparison viewer is an interactive HTML page for
# one analyst; it shows the scores and waterfall, so it is not blind, and it has no label entry.
# For SMEs in Excel there is no Splink equivalent: packets, blinding and adjudication are B2's.

# %% [code]
# ---- B2.4 SME packets -----------------------------------------------------------------------------
def _b2_join(row, cols, sep=" "):
    return sep.join(str(row.get(c, "") or "").strip() for c in cols if str(row.get(c, "") or "").strip())


def _b2_address(r):
    street = _b2_join(r, ["street_number", "street_direction", "street_name", "street_type"])
    unit = str(r.get("unit", "") or "").strip()
    if unit:
        street = f"{street}, unit {unit}" if street else f"unit {unit}"
    place = _b2_join(r, ["city", "state"], ", ")
    z = str(r.get("zip", "") or "").strip()
    place = f"{place} {z}".strip()
    return "; ".join(x for x in (street, place) if x)


# (label shown to the SME, how to build it from an input row). "Same row also lists" is the other
# party of the same input row: context, not a field of this party.
B2_DISPLAY = {
    "person": [
        ("Name", lambda r: _b2_join(r, ["first_name", "middle_name", "last_name"])),
        ("Date of birth", lambda r: r.get("dob", "")),
        ("SSN", lambda r: r.get("ssn", "")),
        ("Driver licence", lambda r: _b2_join(r, ["driver_license_state", "driver_license_number"])),
        ("Provider NPI", lambda r: r.get("provider_npi", "")),
        ("Professional licence", lambda r: _b2_join(r, ["professional_license_state", "professional_license_type",
                                                          "professional_license_number"])),
        ("Specialty", lambda r: r.get("provider_specialty", "")),
        ("Category", lambda r: r.get("category", "")),
        ("Address", _b2_address),
        ("Home phone", lambda r: r.get("home_phone", "")),
        ("Work phone", lambda r: r.get("work_phone", "")),
        ("Email", lambda r: r.get("email", "")),
        ("Vehicle", lambda r: _b2_join(r, ["vin", "plate_state", "plate_number"])),
        ("Same row also lists", lambda r: r.get("business_name", "")),
    ],
    "business": [
        ("Business name", lambda r: r.get("business_name", "")),
        ("TIN", lambda r: r.get("tin", "")),
        ("Clinic NPI", lambda r: r.get("clinic_npi", "")),
        ("Specialty", lambda r: r.get("provider_specialty", "")),
        ("Category", lambda r: r.get("category", "")),
        ("Address", _b2_address),
        ("Work phone", lambda r: r.get("work_phone", "")),
        ("Email", lambda r: r.get("email", "")),
        ("Same row also lists", lambda r: _b2_join(r, ["first_name", "middle_name", "last_name"])),
    ],
}
B2_LABEL_COLUMNS = ["review_id", "label", "confidence", "reason", "annotator_id", "labelled_on"]
B2_SHEETS = {"person": "Persons", "business": "Businesses"}


def _b2_mask(v):
    s = str(v or "")
    digits = _re.sub(r"\D", "", s)
    return ("*" * max(0, len(digits) - 4) + digits[-4:]) if digits else s


def b2_display(run, key, cfg):
    """Side-by-side display rows per part: review_id, then '<field> (extracted)' / '(watchlist)'."""
    xr = run.xrows.set_index("record_id", drop=False)
    wr = run.wrows.set_index("record_id", drop=False)
    out = {}
    for part, spec in B2_DISPLAY.items():
        k = key[key["extracted_part"] == part]
        rows = []
        for _, r in k.iterrows():
            xrow = xr.loc[r["extracted_record_id"]].to_dict() if r["extracted_record_id"] in xr.index else {}
            wrow = wr.loc[r["watchlist_record_id"]].to_dict() if r["watchlist_record_id"] in wr.index else {}
            d = {"review_id": r["review_id"]}
            for name, fn in spec:
                a, b = str(fn(xrow) or "").strip(), str(fn(wrow) or "").strip()
                if name in cfg.masked_fields:
                    a, b = _b2_mask(a), _b2_mask(b)
                d[f"{name} (extracted)"] = a
                d[f"{name} (watchlist)"] = b
            rows.append(d)
        cols = ["review_id"] + [f"{n} ({s})" for n, _ in spec for s in ("extracted", "watchlist")]
        out[part] = pd.DataFrame(rows, columns=cols)
    return out


B2_SME_INSTRUCTIONS = """\
# Reviewing possible watchlist matches: instructions for reviewers

## What you are doing

Each row of the workbook is one **pair**: on the left, a person or business as it was recorded
in claim data ("extracted"); on the right, a person or business from the watchlist. Your job is
to say whether the two describe **the same real-world party**. You label pairs; you do not
search for matches, and nothing you do changes any live system.

No score or computer judgement is shown, on purpose: we use your labels to measure how well
the matching works, so they must be yours alone.

## The three labels

- **match**: the same person, or the same business.
- **non-match**: different parties, even if they are related (family members, a doctor and
  the clinic she works at, two separately registered branches of one brand, a parent company
  and its subsidiary).
- **unsure**: the values shown do not settle it either way. Use it when you cannot lean either
  way, not as "probably". Always say why (for example "common name, nothing else to compare").

## What counts as the same party

**Persons: the same human being.** These do *not* make a non-match on their own:
nicknames and short forms (Bill / William, Peggy / Margaret), a first initial instead of a first
name, small spelling mistakes, first and last name swapped, a missing middle name, a surname
change (marriage), a different address or phone (people move), empty cells.

These point strongly to a **non-match**: a different date of birth that is not a typo or a
day/month swap; a different SSN, driver licence or NPI that is not a one-digit slip; Jr / Sr or
other generations with different dates of birth.

**Businesses: the same business entity.** These do *not* make a non-match on their own:
"d/b/a" (doing business as) names, a missing or different suffix (LLC, Inc, PC, PLLC),
abbreviations (Ctr / Center, & / And), a move.

These point strongly to a **non-match**: a different TIN; the same brand name in a different
location with a different TIN or clinic NPI ("sibling" businesses); different owners named
alongside.

## Weighing what you see

- **Identifiers** (SSN, TIN, NPI, licence, driver licence, VIN) are the strongest evidence.
  The same identifier is close to decisive; one digit different may be a typing error.
- **Common names alone** (for example "John Smith" with nothing else to compare) are not
  enough for a match: label **unsure** unless something else agrees or conflicts.
- **An empty cell is no information**, not a disagreement.
- **"Same row also lists"** shows the other party recorded on the same row (for example the
  clinic a doctor was recorded with). Use it as context only.
- Values may be written differently (upper/lower case, date formats, spaces in phone
  numbers). Judge the value, not the formatting.

## Confidence

- **high**: you would be surprised to be wrong.
- **medium**: you lean one way but the evidence is not conclusive.
- **low**: a weak lean. If you have no lean at all, choose **unsure** instead.

## Reason

One short line naming the fields that decided it, for example "same SSN and date of birth;
Bill/William". Required for **unsure** and for anything not **high**; welcome for the rest.

## Rules

1. Work alone and do not discuss pairs with other reviewers: some pairs are given to two
   reviewers to measure agreement.
2. Use only the values in the workbook. Do not look the parties up in other systems unless the
   study coordinator tells you otherwise.
3. Do not edit the data cells, add or delete rows, or rename sheets. Only the label, confidence,
   reason and date cells are editable.
4. Label every row, then return the file with its original name to the study coordinator.

## Examples (fictional)

| Extracted | Watchlist | Label | Why |
|---|---|---|---|
| William R. Hart, DOB 1961-03-14, SSN ending 4471 | Bill Hart, DOB 1961-03-14, SSN ending 4471 | match, high | same SSN and DOB; Bill is short for William |
| Maria Lopez, DOB 1979-07-02, Queens NY | Maria Lopez, DOB 1983-11-19, Queens NY | non-match, high | same common name and area, different date of birth |
| John Smith, no other details | John Smith, DOB 1955-01-09, Phoenix AZ | unsure, low | common name, nothing to compare |
| Kaur Anita, DOB 1970-05-06 | Anita Kaur, DOB 1970-06-05 | match, medium | names swapped; DOB day/month swapped |
| Pine Street Family Dental, TIN 12-3456789 | Harbor Dental Group LLC d/b/a Pine Street Family Dental, TIN 12-3456789 | match, high | same TIN; the d/b/a name matches |
| Quick Fix Collision - Newark, TIN 22-1111111 | Quick Fix Collision - Paterson, TIN 22-2222222 | non-match, high | same brand, different TIN and town: sibling businesses |
| Dr. Ana Ruiz, NPI 1234567893 | Ana Ruiz, NPI 1234567839 | match, medium | NPI digits transposed; name agrees; nothing conflicts |
| Robert Nguyen Jr, DOB 1988-02-10 | Robert Nguyen, DOB 1959-08-30, same address | non-match, high | same name and address, a generation apart |

## How to fill in the workbook

1. Open the *Persons* and *Businesses* sheets. Columns A-F are yours: `review_id` (do not
   change), `label`, `confidence`, `reason`, `annotator_id` (already filled in) and
   `labelled_on` (today's date).
2. Pick `label` and `confidence` from the drop-down lists.
3. Save the workbook with its original name and return it.
"""


def b2_data_dictionary():
    rows = [
        ("review_id", "both", "Opaque identifier of the pair. Do not change it.", "fixed"),
        ("label", "both", "Your decision.", "match | non-match | unsure"),
        ("confidence", "both", "How sure you are.", "high | medium | low"),
        ("reason", "both", "One short line: which fields decided it. Required for unsure and for anything not high.", "free text"),
        ("annotator_id", "both", "Your reviewer id, filled in for you.", "fixed"),
        ("labelled_on", "both", "Date you labelled the pair.", "date (YYYY-MM-DD)"),
    ]
    src = {
        "Name": "first_name middle_name last_name", "Date of birth": "dob", "SSN": "ssn",
        "Driver licence": "driver_license_state driver_license_number", "Provider NPI": "provider_npi",
        "Professional licence": "professional_license_state professional_license_type professional_license_number",
        "Specialty": "provider_specialty", "Category": "category (medical, legal, repair shop, witness, claimant, other)",
        "Address": "street_number street_direction street_name street_type, unit; city, state zip",
        "Home phone": "home_phone", "Work phone": "work_phone", "Email": "email",
        "Vehicle": "vin plate_state plate_number", "Business name": "business_name", "TIN": "tin",
        "Clinic NPI": "clinic_npi",
        "Same row also lists": "the other party on the same input row (a person's business, a business's person)",
    }
    for part, spec in B2_DISPLAY.items():
        for name, _ in spec:
            for side in ("extracted", "watchlist"):
                who = "the claim-data record" if side == "extracted" else "the watchlist record"
                rows.append((f"{name} ({side})", B2_SHEETS[part], f"{name} of {who}, as recorded.", src.get(name, "")))
    return pd.DataFrame(rows, columns=["column", "sheet", "meaning", "values / source columns"])


def _b2_strip_md(line):
    return line.replace("**", "").replace("`", "").lstrip("#").strip()


def b2_write_review_workbook(path, display, annotator, labels=None, fixed_time=None, title="Watchlist pair review"):
    """The blind workbook. `labels`: optional frame (review_id, label, confidence, reason,
    labelled_at) to pre-fill (used for the filled examples and the synthetic run)."""
    import xlsxwriter
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = xlsxwriter.Workbook(str(path), {"strings_to_urls": False, "strings_to_formulas": False,
                                         "strings_to_numbers": False})
    created = fixed_time or _dt.datetime(2026, 9, 26, 9, 0, 0)
    wb.set_properties({"title": title, "author": "B2 packet generator", "created": created,
                       "comments": f"annotator {annotator}"})
    bold = wb.add_format({"bold": True})
    wrap = wb.add_format({"text_wrap": True, "valign": "top"})
    h_lab = wb.add_format({"bold": True, "bg_color": "#F4E3B5", "border": 1, "text_wrap": True, "valign": "top", "locked": True})
    h_x = wb.add_format({"bold": True, "bg_color": "#DCE8F5", "border": 1, "text_wrap": True, "valign": "top", "locked": True})
    h_w = wb.add_format({"bold": True, "bg_color": "#FBEFD9", "border": 1, "text_wrap": True, "valign": "top", "locked": True})
    c_x = wb.add_format({"bg_color": "#F1F6FB", "locked": True, "text_wrap": True, "valign": "top"})
    c_w = wb.add_format({"bg_color": "#FFFaF1", "locked": True, "text_wrap": True, "valign": "top"})
    c_id = wb.add_format({"locked": True, "font_name": "Consolas", "valign": "top"})
    c_in = wb.add_format({"locked": False, "bg_color": "#FFFFFF", "border": 1, "valign": "top", "text_wrap": True})
    c_date = wb.add_format({"locked": False, "num_format": "yyyy-mm-dd", "border": 1, "valign": "top"})
    ws = wb.add_worksheet("Instructions")
    ws.set_column(0, 0, 110, wrap)
    r = 0
    for line in B2_SME_INSTRUCTIONS.splitlines():
        if line.startswith("|---"):
            continue
        text = _b2_strip_md(line)
        if line.startswith("|"):
            text = "   " + "  |  ".join(x.strip() for x in line.strip("|").split("|"))
        ws.write_string(r, 0, text, bold if line.startswith("#") else wrap)
        r += 1
    lab = labels.set_index("review_id") if labels is not None and len(labels) else None
    for part in ("person", "business"):
        d = display.get(part)
        ws = wb.add_worksheet(B2_SHEETS[part])
        cols = B2_LABEL_COLUMNS + [c for c in (d.columns if d is not None else []) if c != "review_id"]
        for j, c in enumerate(cols):
            fmt = h_lab if j < len(B2_LABEL_COLUMNS) else (h_x if c.endswith("(extracted)") else h_w)
            ws.write_string(0, j, c, fmt)
        ws.set_row(0, 32)
        ws.set_column(0, 0, 13); ws.set_column(1, 1, 11); ws.set_column(2, 2, 11); ws.set_column(3, 3, 40)
        ws.set_column(4, 4, 11); ws.set_column(5, 5, 12)
        ws.set_column(len(B2_LABEL_COLUMNS), max(len(B2_LABEL_COLUMNS), len(cols) - 1), 22)
        n = 0 if d is None else len(d)
        for i in range(n):
            row = d.iloc[i]
            rid = row["review_id"]
            ws.write_string(i + 1, 0, rid, c_id)
            L = lab.loc[rid] if lab is not None and rid in lab.index else None
            ws.write_string(i + 1, 1, "" if L is None else str(L.get("label", "") or ""), c_in)
            ws.write_string(i + 1, 2, "" if L is None else str(L.get("confidence", "") or ""), c_in)
            ws.write_string(i + 1, 3, "" if L is None else str(L.get("reason", "") or ""), c_in)
            ws.write_string(i + 1, 4, annotator, c_id)
            when = None if L is None else pd.to_datetime(L.get("labelled_at", None), errors="coerce")
            if when is not None and not pd.isna(when):
                ws.write_datetime(i + 1, 5, when.to_pydatetime().replace(tzinfo=None), c_date)
            else:
                ws.write_blank(i + 1, 5, None, c_date)
            for j, c in enumerate(cols[len(B2_LABEL_COLUMNS):], start=len(B2_LABEL_COLUMNS)):
                ws.write_string(i + 1, j, str(row[c] or ""), c_x if c.endswith("(extracted)") else c_w)
        last = max(1, n)
        ws.data_validation(1, 1, last, 1, {"validate": "list", "source": list(B2_LABELS),
                                           "input_title": "Label", "input_message": "match, non-match or unsure",
                                           "error_title": "Label", "error_message": "Choose match, non-match or unsure."})
        ws.data_validation(1, 2, last, 2, {"validate": "list", "source": list(B2_CONFIDENCE),
                                           "error_message": "Choose high, medium or low."})
        ws.data_validation(1, 5, last, 5, {"validate": "date", "criteria": "between",
                                           "minimum": _dt.date(2020, 1, 1), "maximum": _dt.date(2100, 12, 31),
                                           "error_message": "Enter a date (YYYY-MM-DD)."})
        ws.freeze_panes(1, 6)
        ws.autofilter(0, 0, last, len(cols) - 1)
        ws.protect("", {"autofilter": True, "format_columns": True, "format_rows": True, "select_locked_cells": True,
                        "select_unlocked_cells": True})
    ws = wb.add_worksheet("Data dictionary")
    dd = b2_data_dictionary()
    for j, c in enumerate(dd.columns):
        ws.write_string(0, j, c, bold)
    for i, row in enumerate(dd.itertuples(index=False), start=1):
        for j, v in enumerate(row):
            ws.write_string(i, j, str(v), wrap)
    ws.set_column(0, 0, 32); ws.set_column(1, 1, 12); ws.set_column(2, 2, 70); ws.set_column(3, 3, 50)
    wb.close()
    return path


def b2_label_template(assignments, annotator):
    rids = assignments.loc[assignments["annotator_id"] == annotator, "review_id"]
    return pd.DataFrame({"review_id": rids.to_numpy(), "annotator_id": annotator, "label": "", "confidence": "",
                         "reason": "", "labelled_at": ""})


def _b2_sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def b2_write_packet(run, key, assignments, out_dir, cfg, labels_by_annotator=None, fixed_time=None, round_id=0):
    """Write one review round: to-annotators/<id>/..., coordinator/..., and (when labels are
    given) the filled workbooks. Returns {relative path: sha256}."""
    out_dir = Path(out_dir)
    disp = b2_display(run, key, cfg)
    files = []
    for a in sorted(assignments["annotator_id"].unique()):
        rids = assignments.loc[assignments["annotator_id"] == a].sort_values("position")["review_id"]
        mine = {p: d.set_index("review_id").reindex([r for r in rids if r in set(d["review_id"])]).reset_index()
                for p, d in disp.items()}
        folder = out_dir / "to-annotators" / a
        lab = None if labels_by_annotator is None else labels_by_annotator.get(a)
        name = f"review-workbook_{a}_round{round_id:02d}.xlsx"
        files.append(b2_write_review_workbook(folder / name, mine, a, labels=lab, fixed_time=fixed_time))
        (folder / "SME-INSTRUCTIONS.md").write_text(B2_SME_INSTRUCTIONS, encoding="utf-8")
        files.append(folder / "SME-INSTRUCTIONS.md")
        files.append(b2_write_table(b2_data_dictionary(), folder / "data-dictionary.csv"))
        files.append(b2_write_table(b2_label_template(assignments, a), folder / f"labels-template_{a}_round{round_id:02d}.csv"))
    co = out_dir / "coordinator"
    files.append(b2_write_table(key, co / f"sample-key_round{round_id:02d}.csv"))
    files.append(b2_write_table(assignments, co / f"assignments_round{round_id:02d}.csv"))
    flat = pd.concat([d.assign(party_type=p) for p, d in disp.items()], ignore_index=True)
    files.append(b2_write_table(flat, co / f"pair-queue-import_round{round_id:02d}.csv"))
    files.append(b2_write_table(b2_schema_frame(), co / "label-schema.csv"))
    manifest = {
        "b2_version": B2_VERSION, "run_id": run.run_id, "core_version": run.core_version, "dataset": run.dataset,
        "round": round_id, "seed": cfg.seed, "created": (fixed_time or _dt.datetime.now()).isoformat(timespec="seconds"),
        "pairs": int(len(key)), "double_coded": int(key.get("double_coded", pd.Series(dtype=bool)).sum()),
        "by_split": key["split"].value_counts().to_dict(), "by_selection": key["selection"].value_counts().to_dict(),
        "by_stratum": key.groupby("stratum").agg(N_h=("N_h", "first"), n_h=("n_h", "first")).reset_index().to_dict("records"),
        "annotators": assignments["annotator_id"].value_counts().sort_index().to_dict(),
        "files": {str(Path(f).relative_to(out_dir)).replace("\\", "/"): _b2_sha(f) for f in files},
    }
    (co / f"packet-manifest_round{round_id:02d}.json").write_text(json.dumps(manifest, indent=1, default=str), encoding="utf-8")
    return manifest["files"]


def b2_write_adjudication(path, run, key, queue, labels, cfg, filled=None, fixed_time=None):
    """Adjudication sheet: the pair side by side, each reviewer's label, confidence and reason
    (as Reviewer 1, 2, ...), and cells for the final decision."""
    import xlsxwriter
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    k = key[key["review_id"].isin(queue["review_id"])]
    disp = b2_display(run, k, cfg)
    wb = xlsxwriter.Workbook(str(path), {"strings_to_urls": False, "strings_to_formulas": False})
    wb.set_properties({"title": "Adjudication", "author": "B2 packet generator",
                       "created": fixed_time or _dt.datetime(2026, 9, 26, 9, 0, 0)})
    head = wb.add_format({"bold": True, "bg_color": "#DDE3EA", "border": 1, "text_wrap": True, "valign": "top"})
    lockw = wb.add_format({"locked": True, "text_wrap": True, "valign": "top"})
    free = wb.add_format({"locked": False, "border": 1, "text_wrap": True, "valign": "top"})
    fdate = wb.add_format({"locked": False, "border": 1, "num_format": "yyyy-mm-dd", "valign": "top"})
    ws = wb.add_worksheet("How to adjudicate")
    ws.set_column(0, 0, 110, lockw)
    for i, t in enumerate([
            "Adjudication: decide the final label for each pair below.",
            "These pairs were labelled differently by two reviewers, or marked unsure. Reviewers are shown as Reviewer 1, 2 in random order.",
            "Read the values, the reviewers' labels and reasons, then fill final_label (match / non-match / unsure), adjudicator_id, adjudication_reason and adjudicated_on.",
            "Keep 'unsure' when the values truly do not settle it: unsure pairs are counted and reported, never forced.",
            "The same definitions apply as in the reviewers' instructions (SME-INSTRUCTIONS.md)."]):
        ws.write_string(i, 0, t)
    fl = filled.set_index("review_id") if filled is not None and len(filled) else None
    max_r = int(labels[labels["review_id"].isin(queue["review_id"])].groupby("review_id").size().max() or 1) if len(labels) else 1
    for part in ("person", "business"):
        d = disp.get(part)
        if d is None or not len(d):
            continue
        ws = wb.add_worksheet(B2_SHEETS[part])
        rcols = []
        for j in range(1, max_r + 1):
            rcols += [f"Reviewer {j} label", f"Reviewer {j} confidence", f"Reviewer {j} reason"]
        cols = ["review_id", "why here", "final_label", "adjudicator_id", "adjudication_reason", "adjudicated_on"] + rcols + \
               [c for c in d.columns if c != "review_id"]
        for j, c in enumerate(cols):
            ws.write_string(0, j, c, head)
        ws.set_row(0, 32)
        for i, row in enumerate(d.itertuples(index=False), start=1):
            rid = row[0]
            q = queue.set_index("review_id").loc[rid]
            L = labels[labels["review_id"] == rid].copy()
            L["o"] = [b2_hash_unit("adj", cfg.seed, rid, a) for a in L["annotator_id"]]
            L = L.sort_values("o")
            ws.write_string(i, 0, rid, lockw)
            ws.write_string(i, 1, str(q["why"]), lockw)
            F = fl.loc[rid] if fl is not None and rid in fl.index else None
            ws.write_string(i, 2, "" if F is None else str(F["final_label"]), free)
            ws.write_string(i, 3, "" if F is None else str(F["adjudicator_id"]), free)
            ws.write_string(i, 4, "" if F is None else str(F["adjudication_reason"]), free)
            when = None if F is None else pd.to_datetime(F.get("adjudicated_at"), errors="coerce")
            if when is not None and not pd.isna(when):
                ws.write_datetime(i, 5, when.to_pydatetime().replace(tzinfo=None), fdate)
            else:
                ws.write_blank(i, 5, None, fdate)
            j = 6
            for _, lr in L.iterrows():
                for v in (lr["label"], lr.get("confidence", ""), lr.get("reason", "")):
                    ws.write_string(i, j, str(v or ""), lockw)
                    j += 1
            j = 6 + 3 * max_r
            for v in row[1:]:
                ws.write_string(i, j, str(v or ""), lockw)
                j += 1
        n = len(d)
        ws.data_validation(1, 2, n, 2, {"validate": "list", "source": list(B2_LABELS)})
        ws.data_validation(1, 5, n, 5, {"validate": "date", "criteria": "between", "minimum": _dt.date(2020, 1, 1),
                                        "maximum": _dt.date(2100, 12, 31)})
        ws.set_column(0, 0, 13); ws.set_column(1, 1, 22); ws.set_column(2, 3, 13); ws.set_column(4, 4, 40)
        ws.set_column(5, 5, 12); ws.set_column(6, len(cols) - 1, 22)
        ws.freeze_panes(1, 6)
        ws.protect("", {"autofilter": True, "format_columns": True, "format_rows": True})
    wb.close()
    return path


# %% [markdown]
# ### B2.5 · Label ingestion, validation, agreement, adjudication
#
# Labels come back as the review workbooks (sheets *Persons* / *Businesses*), CSV label
# templates, or any table with the label columns (Parquet or Delta through the IO switch).
# Checks, each with a reason in the issues table:
#
# | Check | Outcome |
# |---|---|
# | `review_id` not in the sample key | rejected (`unknown_review_id`) |
# | pair no longer in the run's candidates | rejected (`stale_pair`) |
# | annotator not assigned the pair | rejected (`not_assigned`) |
# | label not one of match / non-match / unsure (after synonyms) | rejected (`invalid_label`) |
# | the same annotator, pair and label twice | one kept (`duplicate`) |
# | the same annotator and pair with different labels | all rejected, pair needs relabelling (`contradictory`) |
# | empty label | pending, not an error (`unlabelled`) |
# | unsure, or confidence not high, without a reason | kept, warned (`reason_missing`) |
# | confidence not high / medium / low | kept with confidence emptied (`invalid_confidence`) |
#
# **Agreement** on pairs with two or more labels: raw agreement, Cohen's kappa per annotator
# pair (three categories, and match vs non-match only), Fleiss' kappa over all multi-labelled
# pairs. **Resolution**: one label → that label; unanimous → that label; otherwise the pair goes
# to adjudication, and so does any single `unsure` (a second look). The adjudicator's final label
# wins; a pair still unresolved stays `unresolved` and is counted, never forced.
#
# **Where Splink would do better.** Splink takes labels as a table and has no notion of several
# annotators, agreement or adjudication; B2 has no Splink counterpart to lean on here.

# %% [code]
# ---- B2.5 ingestion and validation -----------------------------------------------------------------
def _b2_norm_label(v):
    s = str(v if v is not None else "").strip().lower()
    if s in ("", "nan", "none"):
        return ""
    return B2_LABEL_SYNONYMS.get(s, "?" + s)


def b2_read_labels(paths):
    """Raw label rows from workbooks, CSV templates or tables. Adds source_file, source_row and
    the file's modified time for empty timestamps."""
    out = []
    for path in [Path(p) for p in (paths if isinstance(paths, (list, tuple)) else [paths])]:
        mtime = _dt.datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds") if path.exists() else ""
        frames = []
        if path.suffix.lower() in (".xlsx", ".xlsm"):
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            for sname in wb.sheetnames:
                if sname not in B2_SHEETS.values():
                    continue
                rows = list(wb[sname].iter_rows(values_only=True))
                if not rows:
                    continue
                head = [str(h) if h is not None else "" for h in rows[0]]
                if "review_id" not in head:
                    continue
                idx = {h: head.index(h) for h in B2_LABEL_COLUMNS if h in head}
                recs = []
                for i, r in enumerate(rows[1:], start=2):
                    if r is None or all(v is None for v in r):
                        continue
                    rec = {h: r[j] if j < len(r) else None for h, j in idx.items()}
                    rec["source_row"] = i
                    rec["source_sheet"] = sname
                    recs.append(rec)
                frames.append(pd.DataFrame(recs))
            wb.close()
        else:
            d = b2_read_table(path)
            d["source_row"] = np.arange(2, len(d) + 2)
            d["source_sheet"] = ""
            frames.append(d)
        for d in frames:
            if not len(d):
                continue
            d = d.copy()
            if "labelled_on" in d.columns and "labelled_at" not in d.columns:
                d = d.rename(columns={"labelled_on": "labelled_at"})
            for c in ("review_id", "annotator_id", "label", "confidence", "reason", "labelled_at"):
                if c not in d.columns:
                    d[c] = ""
            ts = d["labelled_at"].map(lambda v: "" if v is None or (isinstance(v, float) and np.isnan(v)) else
                                      (v.isoformat() if hasattr(v, "isoformat") else str(v)))
            d["timestamp_source"] = np.where(ts.str.strip() == "", "file", "cell")
            d["labelled_at"] = np.where(ts.str.strip() == "", mtime, ts)
            d["source_file"] = path.name
            out.append(d[["review_id", "annotator_id", "label", "confidence", "reason", "labelled_at",
                          "timestamp_source", "source_file", "source_sheet", "source_row"]])
    if not out:
        return pd.DataFrame(columns=["review_id", "annotator_id", "label", "confidence", "reason", "labelled_at",
                                     "timestamp_source", "source_file", "source_sheet", "source_row"])
    raw = pd.concat(out, ignore_index=True)
    for c in ("review_id", "annotator_id", "confidence", "reason"):
        raw[c] = raw[c].map(lambda v: "" if v is None or (isinstance(v, float) and np.isnan(v)) else str(v).strip())
    return raw


def b2_validate_labels(raw, key, assignments=None, run=None):
    """Returns (valid labels, issues, pending). See the table above."""
    issues = []
    d = raw.copy()
    d["label_raw"] = d["label"]
    d["label"] = d["label"].map(_b2_norm_label)

    def flag(mask, issue, detail=""):
        for _, r in d[mask].iterrows():
            issues.append({"review_id": r["review_id"], "annotator_id": r["annotator_id"], "issue": issue,
                           "detail": detail or str(r.get("label_raw", "")), "source_file": r["source_file"],
                           "source_row": r["source_row"]})

    pending = d[d["label"] == ""]
    flag(d["label"] == "", "unlabelled", "no label yet")
    d = d[d["label"] != ""]
    bad = ~d["review_id"].isin(set(key["review_id"]))
    flag(bad, "unknown_review_id"); d = d[~bad]
    if run is not None:
        live = set(run.candidates["pair_id"])
        stale_ids = set(key.loc[~key["pair_id"].isin(live), "review_id"])
        bad = d["review_id"].isin(stale_ids)
        flag(bad, "stale_pair", "pair not in this run's candidates"); d = d[~bad]
    if assignments is not None and len(assignments):
        ok_pairs = set(zip(assignments["review_id"], assignments["annotator_id"]))
        bad = np.array([(r, a) not in ok_pairs for r, a in zip(d["review_id"], d["annotator_id"])], dtype=bool)
        flag(bad, "not_assigned"); d = d[~bad]
    bad = d["label"].str.startswith("?")
    flag(bad, "invalid_label"); d = d[~bad]
    dup = d.duplicated(["review_id", "annotator_id", "label"], keep="first")
    flag(dup, "duplicate", "same annotator, pair and label repeated; one kept"); d = d[~dup]
    multi = d.groupby(["review_id", "annotator_id"])["label"].transform("nunique") > 1
    flag(multi, "contradictory", "same annotator gave different labels; all rejected"); d = d[~multi]
    conf = d["confidence"].str.lower()
    badc = (conf != "") & ~conf.isin(B2_CONFIDENCE)
    flag(badc, "invalid_confidence")
    d["confidence"] = np.where(badc, "", conf)
    nor = (d["reason"] == "") & ((d["label"] == "unsure") | (d["confidence"].isin(["medium", "low"])))
    flag(nor, "reason_missing", "reason expected for unsure or confidence below high")
    iss = pd.DataFrame(issues, columns=["review_id", "annotator_id", "issue", "detail", "source_file", "source_row"])
    return d.drop(columns=["label_raw"]).reset_index(drop=True), iss, pending.reset_index(drop=True)


def b2_cohen_kappa(a, b, cats=None):
    a, b = np.asarray(a, dtype=object), np.asarray(b, dtype=object)
    cats = list(cats or sorted(set(a) | set(b)))
    n = len(a)
    if n == 0:
        return float("nan")
    ia = np.array([cats.index(x) for x in a]); ib = np.array([cats.index(x) for x in b])
    cm = np.zeros((len(cats), len(cats)))
    np.add.at(cm, (ia, ib), 1)
    po = np.trace(cm) / n
    pe = float((cm.sum(1) / n) @ (cm.sum(0) / n))
    return float("nan") if pe >= 1 else float((po - pe) / (1 - pe))


def b2_fleiss_kappa(counts):
    """counts: items x categories, each row summing to that item's number of ratings (>= 2)."""
    c = np.asarray(counts, dtype=float)
    c = c[c.sum(1) >= 2]
    if not len(c):
        return float("nan")
    n_i = c.sum(1)
    P_i = ((c ** 2).sum(1) - n_i) / (n_i * (n_i - 1))
    p_j = c.sum(0) / n_i.sum()
    P_e = float((p_j ** 2).sum())
    return float("nan") if P_e >= 1 else float((P_i.mean() - P_e) / (1 - P_e))


def b2_agreement(valid):
    """Agreement statistics on pairs labelled by two or more annotators."""
    rows = []
    multi = valid[valid.groupby("review_id")["annotator_id"].transform("nunique") >= 2]
    ann = sorted(multi["annotator_id"].unique())
    for i in range(len(ann)):
        for j in range(i + 1, len(ann)):
            a = multi[multi["annotator_id"] == ann[i]].set_index("review_id")["label"]
            b = multi[multi["annotator_id"] == ann[j]].set_index("review_id")["label"]
            common = a.index.intersection(b.index)
            if not len(common):
                continue
            x, y = a[common].to_numpy(), b[common].to_numpy()
            both = (x != "unsure") & (y != "unsure")
            rows.append({"annotators": f"{ann[i]} x {ann[j]}", "pairs": len(common),
                         "raw_agreement": float((x == y).mean()),
                         "cohen_kappa_3cat": b2_cohen_kappa(x, y, list(B2_LABELS)),
                         "pairs_both_definite": int(both.sum()),
                         "cohen_kappa_match_vs_non": b2_cohen_kappa(x[both], y[both], ["match", "non-match"]) if both.any() else float("nan")})
    if len(multi):
        ct = pd.crosstab(multi["review_id"], multi["label"]).reindex(columns=list(B2_LABELS), fill_value=0)
        same = (ct.max(axis=1) == ct.sum(axis=1)).mean()
        rows.append({"annotators": "all (Fleiss)", "pairs": len(ct), "raw_agreement": float(same),
                     "fleiss_kappa_3cat": b2_fleiss_kappa(ct.to_numpy())})
    return pd.DataFrame(rows)


def b2_adjudication_queue(valid):
    """Pairs needing an adjudicator: annotators disagree, or a single label is unsure."""
    g = valid.groupby("review_id")
    s = pd.DataFrame({"n": g["annotator_id"].nunique(), "distinct": g["label"].nunique(),
                      "any_unsure": g["label"].agg(lambda x: (x == "unsure").any())})
    why = np.select([s["distinct"] > 1, s["any_unsure"]], ["reviewers disagree", "marked unsure"], "")
    q = s.assign(why=why)
    return q[q["why"] != ""].reset_index()


def b2_read_adjudication(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    recs = []
    for sname in wb.sheetnames:
        if sname not in B2_SHEETS.values():
            continue
        rows = list(wb[sname].iter_rows(values_only=True))
        head = [str(h) if h is not None else "" for h in rows[0]]
        for r in rows[1:]:
            if r is None or r[0] is None:
                continue
            rec = dict(zip(head, r))
            recs.append({"review_id": str(rec.get("review_id", "")), "final_label": _b2_norm_label(rec.get("final_label")),
                         "adjudicator_id": str(rec.get("adjudicator_id") or ""),
                         "adjudication_reason": str(rec.get("adjudication_reason") or ""),
                         "adjudicated_at": (rec["adjudicated_on"].isoformat() if hasattr(rec.get("adjudicated_on"), "isoformat")
                                            else str(rec.get("adjudicated_on") or ""))})
    wb.close()
    return pd.DataFrame(recs, columns=["review_id", "final_label", "adjudicator_id", "adjudication_reason", "adjudicated_at"])


def b2_final_labels(valid, adjudication, key):
    """One row per labelled pair, with the key's stratum, weight and split."""
    adj = adjudication[adjudication["final_label"].isin(B2_LABELS)] if adjudication is not None and len(adjudication) \
        else pd.DataFrame(columns=["review_id", "final_label", "adjudicator_id", "adjudication_reason"])
    adj = adj.drop_duplicates("review_id", keep="last").set_index("review_id")
    rows = []
    for rid, g in valid.groupby("review_id", sort=True):
        labs = g["label"].tolist()
        uni = len(set(labs)) == 1
        if rid in adj.index:
            fl, src = adj.loc[rid, "final_label"], "adjudicated"
        elif uni and not (len(labs) == 1 and labs[0] == "unsure"):
            fl, src = labs[0], ("unanimous" if len(labs) > 1 else "single")
        elif uni and labs[0] == "unsure":
            fl, src = "unsure", "single"
        else:
            fl, src = "", "unresolved"
        rows.append({"review_id": rid, "final_label": fl, "n_annotators": len(labs),
                     "labels": "; ".join(f"{a}:{l}" for a, l in zip(g["annotator_id"], labs)),
                     "unanimous": uni, "adjudicated": rid in adj.index,
                     "adjudicator_id": adj.loc[rid, "adjudicator_id"] if rid in adj.index else "",
                     "adjudication_reason": adj.loc[rid, "adjudication_reason"] if rid in adj.index else "",
                     "label_source": src})
    f = pd.DataFrame(rows, columns=["review_id", "final_label", "n_annotators", "labels", "unanimous", "adjudicated",
                                    "adjudicator_id", "adjudication_reason", "label_source"])
    f["y"] = f["final_label"].map({"match": 1.0, "non-match": 0.0})
    return key.merge(f, on="review_id", how="inner")


class B2SealedAccessError(RuntimeError):
    pass


class B2Gold:
    """Holds the final labels. Tuning code gets `tuning_view()` (train and calibration only);
    `sealed_view()` is for evaluation and logs every access."""

    def __init__(self, final):
        self._all = final.copy()
        self.access_log = []

    def tuning_view(self):
        v = self._all[self._all["split"] != "sealed_test"].copy()
        v.attrs["b2_view"] = "tuning"
        return v

    def sealed_view(self, purpose):
        self.access_log.append({"purpose": purpose, "at": _dt.datetime.now().isoformat(timespec="seconds"),
                                "rows": int((self._all["split"] == "sealed_test").sum())})
        v = self._all[self._all["split"] == "sealed_test"].copy()
        v.attrs["b2_view"] = "sealed"
        return v

    def counts(self):
        a = self._all
        return a.groupby(["split", "final_label"]).size().unstack(fill_value=0)


def b2_guard_tuning(df):
    """Every fitting function starts here: only a tuning view, never a sealed row."""
    if df.attrs.get("b2_view") != "tuning":
        raise B2SealedAccessError("fitting code must be given B2Gold.tuning_view()")
    if "split" in df.columns and (df["split"] == "sealed_test").any():
        raise B2SealedAccessError("sealed_test rows reached fitting code")


# %% [markdown]
# ### B2.6 · Evaluation on the sealed split
#
# Weights are **post-stratified Horvitz-Thompson**: within stratum h, each sealed pair with a
# definite label counts N_h / m_h, where m_h is the number of such pairs in the stratum. Unsure
# and unresolved pairs are excluded from the weighted metrics (treated as missing at random within
# their stratum) and reported per stratum; a stratum with no sealed label is reported with its
# N_h as uncovered. Reported per model (the pipeline "B1", each improvement, the baseline if its
# column exists):
#
# * precision, recall and F1 at each threshold, and by p band and by basis;
# * **recall inside the candidate space** (the estimated share of all matches among the kept
#   candidates that reach the threshold), and the **recall ceiling**: the share of true pairs
#   the pipeline proposes and keeps at all, from the pipeline's own diagnostics (identifier-anchor
#   recall of the name rules; the truth file on synthetic data), so overall recall ~ candidate-space
#   recall x ceiling;
# * reliability table (mean p vs observed share per p bin), Brier score, log loss, weighted
#   average precision, confusion tables.
#
# Every comparison with B1 uses a **paired stratified bootstrap** on the sealed pairs.
#
# **Where Splink would do better.** `accuracy_analysis_from_labels_table` gives precision,
# recall, F1 and ROC / precision-recall charts over all thresholds, interactively, and
# `truth_space_table` gives the confusion counts; for an unweighted labels table these are
# complete. They do not weight a stratified sample, estimate a recall ceiling, or bootstrap.

# %% [code]
# ---- B2.6 evaluation --------------------------------------------------------------------------------
def b2_poststrat_weights(df):
    """N_h / m_h over rows with a definite label (y not NaN) and stratified selection."""
    ok = df["y"].notna() & (df["selection"] == "stratified")
    m = df[ok].groupby("stratum")["stratum"].transform("size")
    w = pd.Series(np.nan, index=df.index)
    w[ok] = df.loc[ok, "N_h"].astype(float) / m
    return w


def b2_prf(y, p, w, t):
    s = p >= t
    tp = float((w * y * s).sum()); fp = float((w * (1 - y) * s).sum()); fn = float((w * y * ~s).sum())
    prec = tp / (tp + fp) if tp + fp > 0 else float("nan")
    rec = tp / (tp + fn) if tp + fn > 0 else float("nan")
    f1 = 2 * prec * rec / (prec + rec) if prec == prec and rec == rec and prec + rec > 0 else float("nan")
    return {"precision": prec, "recall": rec, "f1": f1, "est_tp": tp, "est_fp": fp, "est_fn": fn,
            "est_tn": float((w * (1 - y) * ~s).sum()), "n_flagged": int(s.sum())}


def b2_brier(y, p, w):
    return float((w * (p - y) ** 2).sum() / w.sum()) if w.sum() > 0 else float("nan")


def b2_logloss(y, p, w, eps=1e-6):
    q = np.clip(p, eps, 1 - eps)
    return float(-(w * (y * np.log(q) + (1 - y) * np.log(1 - q))).sum() / w.sum()) if w.sum() > 0 else float("nan")


def b2_avg_precision(y, p, w):
    """Weighted average precision (area under the precision-recall step curve), ties grouped."""
    pos = float((w * y).sum())
    if pos <= 0:
        return float("nan")
    d = pd.DataFrame({"p": p, "tp": w * y, "fp": w * (1 - y)}).groupby("p", sort=True).sum().iloc[::-1]
    ctp, cfp = d["tp"].cumsum().to_numpy(), d["fp"].cumsum().to_numpy()
    prec = ctp / np.maximum(ctp + cfp, 1e-300)
    return float((d["tp"].to_numpy() / pos * prec).sum())


def b2_cost(y, p, w, t, cfg):
    s = p >= t
    return float((w * ((1 - y) * s * cfg.cost_false_alarm + y * ~s * cfg.cost_miss)).sum() / w.sum())


def b2_reliability(y, p, w, edges):
    b = np.clip(np.searchsorted(np.asarray(edges), p, side="right") - 1, 0, len(edges) - 2)
    rows = []
    for i in range(len(edges) - 1):
        s = b == i
        if not s.any():
            continue
        W = float(w[s].sum())
        rows.append({"p_from": edges[i], "p_to": edges[i + 1], "pairs": int(s.sum()), "est_pairs": W,
                     "mean_p": float((w[s] * p[s]).sum() / W), "observed": float((w[s] * y[s]).sum() / W)})
    return pd.DataFrame(rows)


def b2_eval_rows(sealed, scores, thresholds_of, cfg):
    """sealed: sealed rows (y, weight_ps, stratum, band, basis). scores: {model: p array aligned
    to sealed}. thresholds_of: {model: operating threshold}. Returns tables."""
    ok = sealed["y"].notna().to_numpy() & sealed["weight_ps"].notna().to_numpy()
    y = sealed["y"].to_numpy(float)[ok]; w = sealed["weight_ps"].to_numpy(float)[ok]
    band = sealed["band"].to_numpy(object)[ok]; basis = sealed["basis"].to_numpy(object)[ok]
    summary, by_t, by_band, by_basis, rel, conf = [], [], [], [], [], []
    for name, p_all in scores.items():
        p = np.asarray(p_all, dtype=float)[ok]
        t0 = thresholds_of.get(name, cfg.b1_threshold)
        r = b2_prf(y, p, w, t0)
        summary.append({"model": name, "threshold": t0, "pairs": int(ok.sum()), "est_pairs": float(w.sum()),
                        "est_matches": float((w * y).sum()), "brier": b2_brier(y, p, w), "log_loss": b2_logloss(y, p, w),
                        "avg_precision": b2_avg_precision(y, p, w),
                        "expected_cost_per_pair": b2_cost(y, p, w, t0, cfg), **{k: r[k] for k in ("precision", "recall", "f1")}})
        for t in sorted(set(cfg.thresholds) | {t0}):
            by_t.append({"model": name, "threshold": t, **b2_prf(y, p, w, t)})
        for key_, arr, out in (("band", band, by_band), ("basis", basis, by_basis)):
            for g in sorted(set(arr)):
                s = arr == g
                rr = b2_prf(y[s], p[s], w[s], t0)
                out.append({"model": name, key_: g, "pairs": int(s.sum()), "est_pairs": float(w[s].sum()),
                            "observed_match_share": float((w[s] * y[s]).sum() / w[s].sum()),
                            "mean_p": float((w[s] * p[s]).sum() / w[s].sum()), "threshold": t0,
                            **{k: rr[k] for k in ("precision", "recall", "f1")}})
        rl = b2_reliability(y, p, w, cfg.reliability_edges)
        rl.insert(0, "model", name)
        rel.append(rl)
        c = b2_prf(y, p, w, t0)
        raw = b2_prf(y, p, np.ones_like(w), t0)
        conf.append({"model": name, "threshold": t0, "est_TP": c["est_tp"], "est_FP": c["est_fp"], "est_FN": c["est_fn"],
                     "est_TN": c["est_tn"], "raw_TP": raw["est_tp"], "raw_FP": raw["est_fp"], "raw_FN": raw["est_fn"],
                     "raw_TN": raw["est_tn"]})
    return {"summary": pd.DataFrame(summary), "by_threshold": pd.DataFrame(by_t), "by_band": pd.DataFrame(by_band),
            "by_basis": pd.DataFrame(by_basis), "reliability": pd.concat(rel, ignore_index=True) if rel else pd.DataFrame(),
            "confusion": pd.DataFrame(conf)}


def b2_denominators(sealed, frame_counts):
    """Per stratum: population, sealed pairs drawn, labelled definite, unsure, unresolved,
    pending; strata of the population with no sealed label are listed as uncovered."""
    s = sealed.copy()
    s["state"] = np.select([s["y"].notna(), s["final_label"] == "unsure", s["label_source"] == "unresolved"],
                           ["definite", "unsure", "unresolved"], "pending")
    t = pd.crosstab(s["stratum"], s["state"]).reindex(columns=["definite", "unsure", "unresolved", "pending"], fill_value=0)
    t = frame_counts.rename("N_h").to_frame().join(t, how="left").fillna(0)
    t["covered"] = t["definite"] > 0
    t["weight"] = np.where(t["covered"], t["N_h"] / t["definite"].clip(lower=1), np.nan)
    return t.reset_index().rename(columns={"index": "stratum"})


def b2_bootstrap_diff(y, w, strata, pa, pb, metric, reps, seed):
    """Paired stratified bootstrap of metric(pa) - metric(pb). Returns the draws."""
    rng = np.random.default_rng(seed)
    groups = [np.flatnonzero(strata == s) for s in np.unique(strata)]
    out = np.empty(reps)
    for r in range(reps):
        idx = np.concatenate([rng.choice(g, size=len(g), replace=True) for g in groups])
        out[r] = metric(y[idx], pa[idx], w[idx]) - metric(y[idx], pb[idx], w[idx])
    return out


def b2_recall_ceiling(run, eval_summary, cfg):
    """Candidate-space recall (sealed, weighted) and what the pipeline's own diagnostics say about
    true pairs never proposed or not kept."""
    man = run.manifest
    rows = []
    for _, r in eval_summary.iterrows():
        rows.append({"measure": f"recall inside the candidate space at p >= {r['threshold']:.3g} [{r['model']}]",
                     "value": r["recall"], "source": "sealed split, HT-weighted", "note": ""})
    diag = man[man["section"] == "diagnostics"]
    anch = diag[diag["key"].astype(str).str.startswith("identifier_anchor_recall")]
    for _, r in anch.iterrows():
        rows.append({"measure": f"blocking recall of the name rules on identifier-anchored pairs [{str(r['key']).split('.')[-1]}]",
                     "value": pd.to_numeric(r["value"], errors="coerce"), "source": "pipeline manifest (diagnostics)",
                     "note": "lower bound for pairs that have no shared identifier"})
    ceiling = []
    for _, r in diag[diag["key"].astype(str).str.startswith("truth.")].iterrows():
        try:
            t = json.loads(r["value"])
        except Exception:
            continue
        if t.get("true_pairs"):
            share = t["kept"] / t["true_pairs"]
            ceiling.append((t["true_pairs"], share))
            rows.append({"measure": f"true pairs proposed and kept / all true pairs [{t['part']}]", "value": share,
                         "source": "truth file (synthetic / test data only)",
                         "note": f"{t['kept']} of {t['true_pairs']}"})
    ent = run.entities
    if "status" in ent.columns and len(ent):
        rows.append({"measure": "extracted parties with no candidate proposed", "value": float((ent["status"] == "no candidate proposed").mean()),
                     "source": "entities", "note": f"{int((ent['status'] == 'no candidate proposed').sum())} of {len(ent)}"})
    if ceiling:
        c = sum(n * s for n, s in ceiling) / sum(n for n, _ in ceiling)
        src = "truth file"
    else:
        vals = [pd.to_numeric(r["value"], errors="coerce") for _, r in anch.iterrows()]
        vals = [v for v in vals if v == v]
        c = float(np.mean(vals)) if vals else float("nan")
        src = "identifier-anchor recall (proxy)"
    rows.append({"measure": "estimated recall ceiling (share of true pairs inside the candidate space)", "value": c,
                 "source": src, "note": "overall recall ~ candidate-space recall x ceiling"})
    b1 = eval_summary[eval_summary["model"] == "B1"]
    if len(b1) and c == c:
        rows.append({"measure": f"estimated overall recall at p >= {b1['threshold'].iloc[0]:.3g} [B1]",
                     "value": float(b1["recall"].iloc[0]) * c, "source": "product of the two above", "note": ""})
    return pd.DataFrame(rows)


# %% [markdown]
# ### B2.7 · Improvements that use labels, and adoption
#
# Each is fitted on the **tuning view** only, is switchable, and is compared with the pipeline on
# the sealed split:
#
# * **(a) Calibration**: isotonic regression or Platt scaling (logistic on logit p) of the
#   pipeline's p, fitted on the `calibration` split with its HT weights; the method is chosen by
#   weighted Brier in k-fold cross-validation inside that split. Vetoed pairs stay at 0.
# * **(b) Semi-supervised EM**: EM over all non-vetoed candidates with u fixed at the pipeline's
#   (value-specific) u, starting from and shrunk toward the pipeline's m (`em_alpha`
#   pseudo-pairs); `train` labels fix their pairs' match weights (1 or 0) instead of the E-step's
#   guess, and anchor a per-part prior shift (a logistic fit of the labels with the current
#   log-likelihood ratio as offset: selection on the features does not bias it). The prior over
#   all pairs is not re-estimated from candidates, which would be biased by blocking.
# * **(c) Supervised model**: gradient boosting (scikit-learn) on the per-field bits and the
#   prior, trained on `train`. Its explanation is exact per-feature tree contributions (each
#   split's change in the tree's value is credited to the split's field; they sum with a bias to
#   the model's log-odds), shown in bits beside the pipeline's bits on the decision card.
# * **(d) Cost threshold**: the threshold on the pipeline's p that minimizes the weighted
#   expected cost `cost_false_alarm` x false alarms + `cost_miss` x misses on the `calibration`
#   split; compared with `b1_threshold` on the sealed split by the same cost.
#
# **Adoption rule** (stated before the sealed split is read): (a)-(c) are adopted when the
# pipeline's sealed Brier score minus the improvement's has a paired stratified bootstrap
# confidence interval entirely above 0; (d) the same with expected cost. The interval's level is
# 1 - `adopt_alpha` / k for k improvements tested (Bonferroni). "Adopted" means the evidence
# supports switching it on; `B2Config.apply` switches it on, by hand.
#
# **Where Splink would do better.** Splink trains m from labels
# (`estimate_m_from_pairwise_labels`) and from a labelled cluster column
# (`estimate_m_from_label_column`), and its threshold-selection chart shows precision and recall
# per match weight. It has no calibration step, no supervised model, and no sealed-split
# adoption test; its explanation (the waterfall) exists only for Fellegi-Sunter weights.

# %% [code]
# ---- B2.7 improvements ------------------------------------------------------------------------------
def _b2_logit(p):
    q = np.clip(np.asarray(p, dtype=float), 1e-12, 1 - 1e-12)
    return np.log(q / (1 - q))


class B2Calibrator:
    def __init__(self, method, model, cv_brier):
        self.method, self.model, self.cv_brier = method, model, cv_brier

    def predict(self, p, vetoed=None):
        p = np.asarray(p, dtype=float)
        q = self.model.predict(p) if self.method == "isotonic" else self.model.predict_proba(_b2_logit(p).reshape(-1, 1))[:, 1]
        q = np.clip(q, 0.0, 1.0)
        return np.where(vetoed, 0.0, q) if vetoed is not None else q


def _b2_fit_cal(method, p, y, w, seed):
    if method == "isotonic":
        from sklearn.isotonic import IsotonicRegression
        return IsotonicRegression(y_min=0.0, y_max=1.0, out_of_bounds="clip").fit(p, y, sample_weight=w)
    from sklearn.linear_model import LogisticRegression
    return LogisticRegression(C=1e6, max_iter=1000).fit(_b2_logit(p).reshape(-1, 1), y.astype(int), sample_weight=w)


def _b2_rows_for(tuning, split, feats):
    """Labelled definite rows of one split with their candidate positions (vetoed pairs dropped)."""
    r = tuning[(tuning["split"] == split) & tuning["y"].notna()].copy()
    pos = pd.Index(feats.pair_id).get_indexer(r["pair_id"])
    r = r[pos >= 0].assign(_pos=pos[pos >= 0])
    return r[~feats.vetoed[r["_pos"].to_numpy()]]


def b2_fit_calibration(feats, tuning, cfg):
    b2_guard_tuning(tuning)
    r = _b2_rows_for(tuning, "calibration", feats)
    r = r[r["selection"] == "stratified"]
    if len(r) < cfg.min_labels_to_fit or r["y"].nunique() < 2:
        return None, f"too few calibration labels ({len(r)})"
    w = b2_poststrat_weights(r).to_numpy(float)
    p = feats.p[r["_pos"].to_numpy()]
    y = r["y"].to_numpy(float)
    from sklearn.model_selection import StratifiedKFold
    k = max(2, min(cfg.calibration_folds, int(min((y == 1).sum(), (y == 0).sum()))))
    cv = {}
    for mth in cfg.calibration_methods:
        errs, ws = 0.0, 0.0
        for tr, te in StratifiedKFold(n_splits=k, shuffle=True, random_state=cfg.seed).split(p, y):
            if len(np.unique(y[tr])) < 2:
                continue
            c = B2Calibrator(mth, _b2_fit_cal(mth, p[tr], y[tr], w[tr], cfg.seed), None)
            q = c.predict(p[te])
            errs += float((w[te] * (q - y[te]) ** 2).sum()); ws += float(w[te].sum())
        cv[mth] = errs / ws if ws else float("inf")
    best = min(cv, key=cv.get)
    return B2Calibrator(best, _b2_fit_cal(best, p, y, w, cfg.seed), cv), f"{best} (CV Brier {cv})"


@dataclass
class B2EM:
    m_lookup: dict
    delta: dict
    table: pd.DataFrame
    iterations: dict
    labelled: dict


def b2_fit_semisup_em(run, feats, tuning, cfg):
    b2_guard_tuning(tuning)
    defs, _ = b2_core_defs()
    FL = defs["FIELD_LEVELS"]
    mt = b2_m_table(run, feats)
    m0_map = {(a, b, c): v for a, b, c, v in zip(mt["part"], mt["field"], mt["level"], mt["m"])}
    lab = tuning[(tuning["split"] == "train") & tuning["y"].notna()]
    ymap = pd.Series(lab["y"].to_numpy(float), index=lab["pair_id"].to_numpy())
    lookup, delta, iters, nlab, rows = {}, {}, {}, {}, []
    for part in sorted(set(feats.part)):
        ii = np.flatnonzero((feats.part == part) & ~feats.vetoed)
        if not len(ii):
            continue
        fields = [f for f in feats.fields if f in FL and (feats.level[f][ii] != "").any()]
        code, lnu, m0 = {}, {}, {}
        for f in fields:
            levs = FL[f]
            li = {lv: j for j, lv in enumerate(levs)}
            code[f] = np.array([li.get(lv, -1) for lv in feats.level[f][ii]])
            lnu[f] = np.log(np.clip(np.nan_to_num(feats.u[f][ii], nan=1.0), 1e-15, 1.0))
            v = np.array([m0_map.get((part, f, lv), np.nan) for lv in levs], dtype=float)
            v = np.where(np.isnan(v) | (v <= 0), 1e-3, v)
            m0[f] = v / v.sum()
        y = ymap.reindex(feats.pair_id[ii]).to_numpy(float)
        L = ~np.isnan(y)
        z0 = math.log(2) * feats.prior_bits[ii]
        m = {f: m0[f].copy() for f in fields}
        d, it = 0.0, 0
        for it in range(1, cfg.em_max_iter + 1):
            llr = np.zeros(len(ii))
            for f in fields:
                ok = code[f] >= 0
                llr[ok] += np.log(np.clip(m[f][code[f][ok]], 1e-15, 1.0)) - lnu[f][ok]
            z = z0 + llr
            if L.sum() >= 2 and len(np.unique(y[L])) == 2:        # prior shift anchored by labels
                for _ in range(25):
                    s = 1 / (1 + np.exp(-np.clip(z[L] + d, -700, 700)))
                    g = float((y[L] - s).sum()); h = float((s * (1 - s)).sum()) + 1e-9
                    step = float(np.clip(g / h, -2, 2)); d += step
                    if abs(step) < 1e-10:
                        break
            w = 1 / (1 + np.exp(-np.clip(z + d, -700, 700)))
            w[L] = y[L]
            change = 0.0
            for f in fields:
                ok = code[f] >= 0
                k = len(FL[f])
                c = np.bincount(code[f][ok], weights=w[ok], minlength=k)[:k]
                new = (c + cfg.em_alpha * m0[f]) / (w[ok].sum() + cfg.em_alpha)
                change = max(change, float(np.abs(new - m[f]).max()))
                m[f] = new
            if change < cfg.em_tol:
                break
        for f in fields:
            for j, lv in enumerate(FL[f]):
                lookup[(part, f, lv)] = float(m[f][j])
                rows.append({"part": part, "field": f, "level": lv, "m_pipeline": float(m0[f][j]), "m_em": float(m[f][j])})
        delta[part], iters[part], nlab[part] = d, it, int(L.sum())
    return B2EM(lookup, delta, pd.DataFrame(rows), iters, nlab)


def b2_predict_em(feats, em):
    shift = np.array([em.delta.get(p, 0.0) for p in feats.part]) / math.log(2)
    return b2_rescore(feats, em.m_lookup, shift)[0]


def b2_gb_matrix(feats):
    cols = ["prior_bits", "is_person"] + [f"bits_{f}" for f in feats.fields]
    X = np.column_stack([feats.prior_bits, (feats.part == "person").astype(float)] + [feats.bits[f] for f in feats.fields])
    return X.astype(np.float32), cols


@dataclass
class B2Supervised:
    model: object
    columns: list
    trained_on: int


def b2_fit_supervised(feats, tuning, cfg):
    b2_guard_tuning(tuning)
    from sklearn.ensemble import GradientBoostingClassifier
    r = _b2_rows_for(tuning, "train", feats)
    if len(r) < cfg.min_labels_to_fit or r["y"].nunique() < 2:
        return None, f"too few train labels ({len(r)})"
    X, cols = b2_gb_matrix(feats)
    model = GradientBoostingClassifier(random_state=cfg.seed, **cfg.gb_params)
    model.fit(X[r["_pos"].to_numpy()], r["y"].to_numpy(int))
    return B2Supervised(model, cols, len(r)), f"{len(r)} train labels"


def b2_predict_supervised(feats, sup):
    X, _ = b2_gb_matrix(feats)
    return np.where(feats.vetoed, 0.0, sup.model.predict_proba(X)[:, 1])


def b2_gb_contributions(sup, X):
    """Exact per-feature contributions (natural log-odds) of a binary GradientBoostingClassifier:
    for each tree, the change in node value along the decision path, credited to the split
    feature, times the learning rate. Returns (bias per row, contributions n x features); bias +
    row sum = decision_function."""
    model = sup.model
    X = np.asarray(X, dtype=np.float32)
    n, k = X.shape
    contrib = np.zeros((n, k))
    lr = model.learning_rate
    rows = np.arange(n)
    for est in model.estimators_[:, 0]:
        t = est.tree_
        vals = t.value[:, 0, 0]
        node = np.zeros(n, dtype=np.int64)
        while True:
            left = t.children_left[node]
            live = left != -1
            if not live.any():
                break
            i = rows[live]; nd = node[live]
            f = t.feature[nd]
            go_left = X[i, f] <= t.threshold[nd]
            child = np.where(go_left, t.children_left[nd], t.children_right[nd])
            np.add.at(contrib, (i, f), lr * (vals[child] - vals[nd]))
            node[live] = child
    bias = model.decision_function(X) - contrib.sum(1)
    return bias, contrib


def b2_decision_card(feats, sup, rows=None):
    """Long table per pair and field: the pipeline's bits and the supervised model's contribution
    in bits (log-odds / ln 2). `prior` holds the prior and the model's bias."""
    X, cols = b2_gb_matrix(feats)
    idx = np.arange(len(feats.pair_id)) if rows is None else np.asarray(rows)
    bias, con = b2_gb_contributions(sup, X[idx])
    out = []
    for j, c in enumerate(cols):
        fld = c[5:] if c.startswith("bits_") else ("prior" if c == "prior_bits" else "part")
        b1 = feats.bits[fld][idx] if fld in feats.bits else (feats.prior_bits[idx] if fld == "prior" else np.zeros(len(idx)))
        out.append(pd.DataFrame({"pair_id": feats.pair_id[idx], "field": fld, "level": feats.level[fld][idx] if fld in feats.level else "",
                                 "pipeline_bits": b1, "supervised_bits": con[:, j] / math.log(2)}))
    out.append(pd.DataFrame({"pair_id": feats.pair_id[idx], "field": "model bias", "level": "", "pipeline_bits": 0.0,
                             "supervised_bits": bias / math.log(2)}))
    card = pd.concat(out, ignore_index=True)
    card = card[(card["pipeline_bits"] != 0) | (card["supervised_bits"].abs() > 1e-9) | (card["field"] == "model bias")]
    return card.sort_values(["pair_id", "field"], kind="stable").reset_index(drop=True)


def b2_fit_cost_threshold(feats, tuning, cfg):
    b2_guard_tuning(tuning)
    r = _b2_rows_for(tuning, "calibration", feats)
    r = r[r["selection"] == "stratified"]
    if len(r) < cfg.min_labels_to_fit or r["y"].nunique() < 2:
        return None, f"too few calibration labels ({len(r)})"
    w = b2_poststrat_weights(r).to_numpy(float)
    p = feats.p[r["_pos"].to_numpy()]
    y = r["y"].to_numpy(float)
    u = np.unique(p)
    cands = np.unique(np.concatenate([(u[1:] + u[:-1]) / 2, [cfg.b1_threshold], np.asarray(cfg.thresholds)]))
    costs = np.array([b2_cost(y, p, w, t, cfg) for t in cands])
    best = cands[costs <= costs.min() + 1e-12]
    bayes = cfg.cost_false_alarm / (cfg.cost_false_alarm + cfg.cost_miss)
    t = float(best[np.argmin(np.abs(best - bayes))])
    return t, f"threshold {t:.4g} (calibration cost {costs.min():.4g}; B1 threshold cost {b2_cost(y, p, w, cfg.b1_threshold, cfg):.4g})"


def b2_fit_all(run, feats, tuning, cfg):
    """Every improvement, fitted on the tuning view only."""
    b2_guard_tuning(tuning)
    fits, notes = {}, {}
    fits["calibration"], notes["calibration"] = b2_fit_calibration(feats, tuning, cfg)
    fits["semisup_em"] = b2_fit_semisup_em(run, feats, tuning, cfg)
    notes["semisup_em"] = f"labelled per part {fits['semisup_em'].labelled}, prior shift (nats) " \
                          f"{ {k: round(v, 3) for k, v in fits['semisup_em'].delta.items()} }, iterations {fits['semisup_em'].iterations}"
    fits["supervised"], notes["supervised"] = b2_fit_supervised(feats, tuning, cfg)
    fits["cost_threshold"], notes["cost_threshold"] = b2_fit_cost_threshold(feats, tuning, cfg)
    return fits, notes


def b2_scores(run, feats, fits):
    """p of every candidate under each model; the threshold each is operated at."""
    s = {"B1": feats.p.copy()}
    th = {"B1": None}
    if fits.get("calibration") is not None:
        s["calibration"] = fits["calibration"].predict(feats.p, feats.vetoed)
    if fits.get("semisup_em") is not None:
        s["semisup_em"] = b2_predict_em(feats, fits["semisup_em"])
    if fits.get("supervised") is not None:
        s["supervised"] = b2_predict_supervised(feats, fits["supervised"])
    if fits.get("cost_threshold") is not None:
        s["cost_threshold"] = feats.p.copy()
        th["cost_threshold"] = fits["cost_threshold"]
    if "p_baseline" in run.candidates.columns:
        s["baseline"] = pd.to_numeric(run.candidates["p_baseline"], errors="coerce").fillna(0).to_numpy(float)
    return s, th


def b2_evaluate_sealed(run, feats, gold, scores, th, cfg, purpose="evaluation"):
    """The only reader of the sealed split. Returns evaluation tables and the adoption table."""
    sealed = gold.sealed_view(purpose)
    assert (sealed["selection"] == "stratified").all(), "sealed pairs must come from the stratified sample"
    sealed = sealed.copy()
    sealed["weight_ps"] = b2_poststrat_weights(sealed)
    pos = pd.Index(feats.pair_id).get_indexer(sealed["pair_id"])
    if (pos < 0).any():
        raise ValueError("sealed pairs missing from the run's candidates (stale key)")
    thr = {m: (t if t is not None else cfg.b1_threshold) for m, t in th.items()}
    for m in scores:
        thr.setdefault(m, cfg.b1_threshold)
    aligned = {m: np.asarray(v)[pos] for m, v in scores.items()}
    tables = b2_eval_rows(sealed, aligned, thr, cfg)
    frame = b2_frame(run, cfg)
    tables["denominators"] = b2_denominators(sealed, frame.groupby("stratum").size())
    ok = sealed["y"].notna().to_numpy() & sealed["weight_ps"].notna().to_numpy()
    y = sealed["y"].to_numpy(float)[ok]; w = sealed["weight_ps"].to_numpy(float)[ok]
    st = sealed["stratum"].to_numpy(object)[ok]
    tested = [m for m in B2_IMPROVEMENTS if m in scores]
    k = max(1, len(tested))
    q = cfg.adopt_alpha / k
    rows = []
    for i, m in enumerate(B2_IMPROVEMENTS):
        if m not in scores:
            rows.append({"improvement": m, "adopted": False, "metric": "", "b1": np.nan, "candidate": np.nan,
                         "improvement_estimate": np.nan, "ci_low": np.nan, "ci_high": np.nan, "ci_level": np.nan,
                         "sealed_pairs": int(ok.sum()), "reason": "not fitted"})
            continue
        pa, pb = aligned["B1"][ok], aligned[m][ok]
        if m == "cost_threshold":
            ta, tb = cfg.b1_threshold, thr[m]
            met = "expected cost per pair"
            fa = lambda yy, pp, ww, t=ta: b2_cost(yy, pp, ww, t, cfg)
            fb = lambda yy, pp, ww, t=tb: b2_cost(yy, pp, ww, t, cfg)
            point = fa(y, pa, w) - fb(y, pb, w)
            rng = np.random.default_rng([cfg.seed, 99, i])
            groups = [np.flatnonzero(st == s) for s in np.unique(st)]
            draws = np.empty(cfg.bootstrap_reps)
            for r in range(cfg.bootstrap_reps):
                idx = np.concatenate([rng.choice(g, size=len(g), replace=True) for g in groups])
                draws[r] = fa(y[idx], pa[idx], w[idx]) - fb(y[idx], pb[idx], w[idx])
            base, cand = fa(y, pa, w), fb(y, pb, w)
        else:
            met = "Brier score"
            base, cand = b2_brier(y, pa, w), b2_brier(y, pb, w)
            point = base - cand
            draws = b2_bootstrap_diff(y, w, st, pa, pb, b2_brier, cfg.bootstrap_reps, [cfg.seed, 99, i])
        lo, hi = np.quantile(draws, [q / 2, 1 - q / 2])
        adopted = bool(lo > 0)
        rows.append({"improvement": m, "adopted": adopted, "metric": met, "b1": base, "candidate": cand,
                     "improvement_estimate": point, "ci_low": float(lo), "ci_high": float(hi), "ci_level": 1 - q,
                     "sealed_pairs": int(ok.sum()),
                     "reason": ("interval above 0: better than B1 on the sealed split" if adopted else
                                "interval includes or is below 0: not shown to beat B1")})
    tables["adoption"] = pd.DataFrame(rows)
    return tables


def b2_truth_check(run, feats, gold, scores, th, cfg, truth):
    """Synthetic / test data only: how often the resolved gold labels agree with the truth file,
    and every model's sealed metrics against the truth instead of the gold labels (so label
    noise can be told apart from model error)."""
    s = gold.sealed_view("truth diagnostic (synthetic only)").copy()
    s["truth_y"] = [float((a, b, c) in truth) for a, b, c in
                    zip(s["extracted_record_id"], s["extracted_part"], s["watchlist_record_id"])]
    rows = []
    for src, g in s.groupby("label_source"):
        d = g[g["y"].notna()]
        rows.append({"table": "gold vs truth", "label_source": src, "pairs": len(g), "definite": len(d),
                     "agree_with_truth": float((d["y"] == d["truth_y"]).mean()) if len(d) else np.nan})
    s["y"] = s["truth_y"]
    s["weight_ps"] = s["N_h"] / s.groupby("stratum")["stratum"].transform("size")
    pos = pd.Index(feats.pair_id).get_indexer(s["pair_id"])
    thr = {m: (th.get(m) if th.get(m) is not None else cfg.b1_threshold) for m in scores}
    summ = b2_eval_rows(s, {m: np.asarray(v)[pos] for m, v in scores.items()}, thr, cfg)["summary"]
    summ.insert(0, "table", "models vs truth (all sealed pairs)")
    return pd.concat([pd.DataFrame(rows), summ], ignore_index=True)


def b2_apply(run, scores, th, adoption, cfg):
    """Probability used downstream: the pipeline's unless a switch is on AND the improvement was
    adopted. Returns the candidates' B2 columns and a note."""
    adopted = set(adoption.loc[adoption["adopted"], "improvement"])
    out = pd.DataFrame({"pair_id": run.candidates["pair_id"], "p_b1": scores["B1"]})
    for m in ("calibration", "semisup_em", "supervised", "baseline"):
        if m in scores:
            out[f"p_{m}"] = scores[m]
    used = "B1"
    for m in cfg.apply_order:
        if cfg.apply.get(m) and m in adopted and m in scores:
            used = m
            break
    out["p_b2"] = scores[used]
    out["p_b2_source"] = used
    t = th.get("cost_threshold") if (cfg.apply.get("cost_threshold") and "cost_threshold" in adopted) else cfg.b1_threshold
    out["flag_b2"] = out["p_b2"] >= t
    note = f"p_b2 = {used}; threshold {t:.4g}" + ("" if any(cfg.apply.values()) else " (all switches off: B2 changes nothing)")
    return out, note


def b2_manifest_rows(run, cfg, key, issues, agreement, final, tables, notes, gold):
    rows = b2_config_rows(cfg)
    rows.append({"section": "b2.version", "key": "b2", "value": B2_VERSION, "source": "b2_sections"})
    rows.append({"section": "b2.run", "key": "run_id", "value": run.run_id, "source": run.source})
    for sp, n in key["split"].value_counts().items():
        rows.append({"section": "b2.sample", "key": f"split.{sp}", "value": int(n), "source": f"seed {cfg.seed}"})
    for (s, sel), n in key.groupby(["stratum", "selection"]).size().items():
        rows.append({"section": "b2.sample", "key": f"{sel}.{s}", "value": int(n), "source": "sample key",
                     "pairs": int(key.loc[key["stratum"] == s, "N_h"].iloc[0]) if sel == "stratified" else ""})
    for lab, n in final["final_label"].replace("", "unresolved").value_counts().items():
        rows.append({"section": "b2.labels", "key": f"final.{lab}", "value": int(n), "source": "final labels"})
    for iss, n in (issues["issue"].value_counts().items() if len(issues) else []):
        rows.append({"section": "b2.labels", "key": f"issue.{iss}", "value": int(n), "source": "validation"})
    for _, r in agreement.iterrows():
        k = "cohen_kappa_3cat" if "cohen_kappa_3cat" in r and r.get("cohen_kappa_3cat") == r.get("cohen_kappa_3cat") else "fleiss_kappa_3cat"
        rows.append({"section": "b2.agreement", "key": r["annotators"], "value": r.get(k), "pairs": r["pairs"],
                     "source": k, "note": f"raw agreement {r['raw_agreement']:.3f}"})
    for _, r in tables["adoption"].iterrows():
        rows.append({"section": "b2.adoption", "key": r["improvement"], "value": "adopted" if r["adopted"] else "not adopted",
                     "source": f"sealed split, {r['sealed_pairs']} pairs, {r['metric']}", "pairs": r["sealed_pairs"],
                     "ci_low": r["ci_low"], "ci_high": r["ci_high"],
                     "note": f"B1 {r['b1']:.5g} vs {r['candidate']:.5g}; {r['reason']}; fit: {notes.get(r['improvement'], '')}"})
    for _, r in tables["summary"].iterrows():
        rows.append({"section": "b2.eval", "key": f"{r['model']}.brier", "value": r["brier"], "source": "sealed, HT"})
        rows.append({"section": "b2.eval", "key": f"{r['model']}.f1@{r['threshold']:.3g}", "value": r["f1"], "source": "sealed, HT"})
    for a in gold.access_log:
        rows.append({"section": "b2.sealed_access", "key": a["purpose"], "value": a["rows"], "source": a["at"]})
    m = pd.DataFrame(rows)
    for c in ("section", "key", "value", "source", "pairs", "ci_low", "ci_high", "note"):
        if c not in m:
            m[c] = ""
    m = m[["section", "key", "value", "source", "pairs", "ci_low", "ci_high", "note"]].fillna("")
    m["value"] = m["value"].map(lambda v: v if isinstance(v, (int, float, np.integer, np.floating)) else str(v))
    return m


# %% [markdown]
# ### B2.8 · The improvement loop: the next review batch
#
# After a round is labelled and evaluated, the next batch is the most informative unlabelled
# pairs: highest binary entropy of the mean p across every available scorer, plus
# `al_disagreement_weight` x their spread (max - min). At most `al_per_entity` pairs per
# extracted party, only from parties in the `train` split, never a pair already put up for
# review, never a vetoed pair. The batch is written as a new packet (round k) with selection
# `active`; its labels train (b) and (c) in the next round. A fresh stratified top-up for the
# calibration and sealed splits is drawn with `b2_draw_sample(..., exclude=...)` when a round's
# sealed metrics need more pairs.
#
# **Where Splink would do better.** Splink has no active learning. Its comparison viewer can
# show pairs near a chosen match weight, which is uncertainty sampling by hand.

# %% [code]
# ---- B2.8 improvement loop ---------------------------------------------------------------------------
def b2_select_active(run, feats, key_all, sources, cfg, round_id):
    fr = b2_frame(run, cfg)
    P = np.column_stack([np.asarray(v, dtype=float) for v in sources.values()])
    pm = P.mean(1)
    q = np.clip(pm, 1e-12, 1 - 1e-12)
    ent = -(q * np.log2(q) + (1 - q) * np.log2(1 - q))
    spread = P.max(1) - P.min(1)
    fr["al_score"] = ent + cfg.al_disagreement_weight * spread
    fr["al_entropy"], fr["al_spread"] = ent, spread
    fr["split_of_group"] = b2_split_of(fr["group"], cfg.seed, cfg.split_fractions)
    pool = fr[(fr["split_of_group"] == "train") & ~feats.vetoed & ~fr["pair_id"].isin(set(key_all["pair_id"]))]
    pool = pool.sort_values(["al_score", "pair_id"], ascending=[False, True], kind="stable")
    pool = pool[pool.groupby("group").cumcount() < cfg.al_per_entity].head(cfg.al_batch_size).copy()
    N = fr.groupby("stratum").size()
    pool["N_h"] = pool["stratum"].map(N).astype(int)
    pool["n_h"] = 0
    pool["inclusion_prob"] = np.nan
    pool["weight"] = np.nan
    pool["selection"] = "active"
    k = _b2_key_rows(pool, run, cfg, round_id)
    k["al_score"] = pool["al_score"].to_numpy()
    k["al_spread"] = pool["al_spread"].to_numpy()
    return k


def b2_next_batch(run, feats, key_all, sources, cfg, round_id, out_dir, fixed_time=None):
    k = b2_select_active(run, feats, key_all, sources, cfg, round_id)
    k, asg = b2_assign(k, cfg)
    files = b2_write_packet(run, k, asg, out_dir, cfg, fixed_time=fixed_time, round_id=round_id)
    return k, asg, files


# %% [markdown]
# ### B2.9 · Synthetic annotators
#
# For the self-tests and the synthetic run only: labels from the synthetic truth file, with
# per-annotator noise (label flips, and "unsure" much more often on pairs where only the name
# can be compared), confidences and reasons written from the pair's field levels, and a
# simulated adjudicator. Real labels never pass through here.
#
# **Where Splink would do better.** Splink's demo datasets come with ground-truth cluster ids, so
# labels can be taken straight from them; there is no annotator-noise simulator in either.

# %% [code]
# ---- B2.9 synthetic annotators -----------------------------------------------------------------------
B2_REASON = {
    "name": {"exact": "same name", "first_nick_or_close": "first name is a nickname or spelling variant",
             "last_close_first_agrees": "surname spelled slightly differently", "initial_agrees": "only the first initial agrees",
             "swapped": "first and last names swapped", "first_empty": "surname agrees, first name missing",
             "first_differs": "same surname, different first name", "else": "names differ"},
    "org": {"exact": "same business name", "dba": "one name is the other's d/b/a", "short_form": "short form of the same name",
            "rare_shared": "names share a distinctive word", "common_shared": "names share only a common word",
            "sibling": "looks like a sibling business", "none": "business names differ"},
    "dob": {"exact": "same date of birth", "swap_or_typo": "date of birth day/month swapped or one-digit typo",
            "year_month": "same birth year and month only", "year": "same birth year only", "differs": "different date of birth"},
    "address": {"exact": "same address", "street": "same street address", "zip": "same ZIP only", "city_state": "same city only",
                "state": "same state only", "differs": "different address"},
    "middle": {"exact": "same middle name", "initial": "middle initial agrees", "differs": "different middle name"},
    "phone": {"exact_owned_single": "same phone", "exact_shared": "same phone", "differs": "different phone"},
    "spec_cat": {"specialty": "same specialty", "category_id": "same category", "category_weak": "same category",
                 "differs": "different specialty or category"},
    "co_party": {"anchored": "the businesses on the same rows are the same"},
}
_B2_ID_NAMES = {"ssn": "SSN", "npi": "NPI", "dl": "driver licence", "tin": "TIN", "license": "licence", "email": "email",
                "vin": "VIN", "plate": "plate", "cnpi": "clinic NPI"}


def b2_reason_text(levels, label):
    defs, _ = b2_core_defs()
    agree, disagree = [], []
    for f, lv in levels.items():
        if not lv or lv in defs["ZERO_LEVELS"].get(f, set()):
            continue
        if f in _B2_ID_NAMES:
            txt = {"exact": f"same {_B2_ID_NAMES[f]}", "near": f"{_B2_ID_NAMES[f]} differs by one digit",
                   "differs": f"different {_B2_ID_NAMES[f]}"}.get(lv, "")
        else:
            txt = B2_REASON.get(f, {}).get(lv, "")
        if not txt:
            continue
        (agree if lv in defs["AGREEMENT_LEVELS"].get(f, set()) else disagree).append(txt)
    if label == "match":
        return "; ".join(agree[:3]) or "values agree"
    if label == "non-match":
        return "; ".join(disagree[:3]) or "nothing agrees beyond a common name"
    only_name = all(f in ("name", "org", "middle") for f in levels if levels[f]) or not agree
    return ("only the name can be compared" if only_name else "evidence points both ways: " + "; ".join((agree[:1] + disagree[:1])))


def b2_truth_keys(truth_path):
    t = b2_read_table(truth_path)
    return set(zip(t["x_record_id"].astype(str), t["part"].astype(str), t["w_record_id"].astype(str)))


B2_SIM_PROFILES = {"A1": {"flip": 0.02, "unsure_hard": 0.30, "unsure_easy": 0.01},
                   "A2": {"flip": 0.04, "unsure_hard": 0.20, "unsure_easy": 0.02},
                   "A3": {"flip": 0.06, "unsure_hard": 0.40, "unsure_easy": 0.02}}


def b2_simulate_labels(key, assignments, truth, feats, seed, profiles=None, start=None):
    profiles = profiles or B2_SIM_PROFILES
    defs, _ = b2_core_defs()
    kk = key.set_index("review_id")
    pos = pd.Index(feats.pair_id).get_indexer(kk["pair_id"])
    posmap = dict(zip(kk.index, pos))
    start = start or _dt.datetime(2026, 10, 5, 9, 0, 0)
    rows = []
    for a, g in assignments.groupby("annotator_id", sort=True):
        pr = profiles.get(a, {"flip": 0.04, "unsure_hard": 0.25, "unsure_easy": 0.02})
        rng = np.random.default_rng([seed, sum(map(ord, a))])
        for i, rid in enumerate(g.sort_values("position")["review_id"]):
            r = kk.loc[rid]
            y = (r["extracted_record_id"], r["extracted_part"], r["watchlist_record_id"]) in truth
            j = posmap[rid]
            levels = {f: feats.level[f][j] for f in feats.fields} if j >= 0 else {}
            strong = [f for f, lv in levels.items() if lv and f not in ("name", "org", "middle")
                      and lv in defs["AGREEMENT_LEVELS"].get(f, set()) | defs["DISAGREEMENT_LEVELS"].get(f, set())]
            hard = len(strong) == 0
            u = rng.random()
            if u < (pr["unsure_hard"] if hard else pr["unsure_easy"]):
                lab, conf = "unsure", "low"
            elif rng.random() < pr["flip"]:
                lab, conf = ("non-match" if y else "match"), rng.choice(["medium", "low"])
            else:
                lab, conf = ("match" if y else "non-match"), ("medium" if hard else "high")
            rows.append({"review_id": rid, "annotator_id": a, "label": lab, "confidence": str(conf),
                         "reason": b2_reason_text(levels, lab),
                         "labelled_at": (start + _dt.timedelta(minutes=3 * i)).isoformat(timespec="seconds")})
    return pd.DataFrame(rows)


def b2_simulate_adjudication(queue, key, truth, seed, error=0.01, adjudicator="ADJ1", start=None):
    rng = np.random.default_rng([seed, 7])
    kk = key.set_index("review_id")
    start = start or _dt.datetime(2026, 10, 12, 9, 0, 0)
    rows = []
    for i, rid in enumerate(queue["review_id"]):
        r = kk.loc[rid]
        y = (r["extracted_record_id"], r["extracted_part"], r["watchlist_record_id"]) in truth
        if rng.random() < error:
            y = not y
        rows.append({"review_id": rid, "final_label": "match" if y else "non-match", "adjudicator_id": adjudicator,
                     "adjudication_reason": "decided on the identifiers and date of birth shown" if y else "conflicting details outweigh the name",
                     "adjudicated_at": (start + _dt.timedelta(minutes=5 * i)).isoformat(timespec="seconds")})
    return pd.DataFrame(rows)


# %% [markdown]
# ### B2.10 · Run
#
# 1. Draw round 0 (stratified), assign annotators, write the packet to
#    `review/<dataset>/<run_id>/round_00/` (gitignored).
# 2. Read returned labels from `.../round_00/returned/` (workbooks, CSV, Parquet). On the
#    synthetic set, when none are there, the synthetic annotators fill in workbooks there first,
#    so the same file path is exercised. On other datasets with no labels the section stops
#    here and says it is waiting for labels.
# 3. Validate, measure agreement, write the adjudication sheet; read the adjudicator's file
#    (simulated on synthetic); resolve final labels.
# 4. Fit every improvement on the tuning view, evaluate everything on the sealed split, decide
#    adoption, build the B2 manifest section, apply the switches (all off by default).
# 5. Write the next active-learning packet (round 1) and the B2 report
#    `out/<dataset>/b2/b2_report.xlsx` with tables as Parquet beside it.
#
# `RL_B2_WRITE_STUBS=1` also regenerates the committed SME stub packets in `review/stubs/`.
#
# **Where Splink would do better.** A Splink session would do steps 3-4 in a few calls
# (`linker.evaluation.accuracy_analysis_from_labels_table`, `prediction_errors_from_labels_table`)
# and chart them; the rest (packets, adjudication, the sealed split) has no Splink equivalent.

# %% [code]
# ---- B2.10 run ----------------------------------------------------------------------------------------
def b2_write_report(path, tables):
    import xlsxwriter
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = xlsxwriter.Workbook(str(path), {"nan_inf_to_errors": True, "strings_to_urls": False, "strings_to_formulas": False})
    head = wb.add_format({"bold": True, "bg_color": "#DDE3EA"})
    for name, df in tables.items():
        ws = wb.add_worksheet(name[:31])
        cols = [str(c) for c in df.columns]
        ws.write_row(0, 0, cols, head)
        for i, row in enumerate(df.itertuples(index=False), start=1):
            ws.write_row(i, 0, [(None if (isinstance(v, float) and not np.isfinite(v)) else
                                 (v if isinstance(v, (int, float, str, bool, np.integer, np.floating)) else str(v))) for v in row])
        ws.autofilter(0, 0, max(1, len(df)), max(0, len(cols) - 1))
        ws.freeze_panes(1, 0)
        ws.set_column(0, max(0, len(cols) - 1), 16)
    wb.close()
    return path


def b2_run(run=None, cfg=None, simulate=None, log=print, review_dir=None, fixed_time=None):
    """Steps 1-5 above. Returns a dict of every table and fitted object."""
    import time
    t0 = time.time()
    cfg = cfg or B2Config()
    run = run or b2_load_run()
    simulate = (run.dataset == "synthetic" and run.truth_path is not None) if simulate is None else simulate
    feats = b2_features(run)
    rd = Path(review_dir) if review_dir else run.root / "review" / run.dataset / run.run_id
    r0 = rd / "round_00"
    key, asg = b2_assign(b2_draw_sample(run, cfg, 0), cfg)
    b2_write_packet(run, key, asg, r0, cfg, fixed_time=fixed_time, round_id=0)
    log(f"B2 round 0: {len(key)} pairs ({int(key['double_coded'].sum())} double-coded) to {asg['annotator_id'].nunique()} "
        f"annotators; splits {key['split'].value_counts().to_dict()}; packet in {r0}")
    ret = r0 / "returned"
    files = sorted([p for p in ret.glob("*") if p.suffix.lower() in (".xlsx", ".xlsm", ".csv", ".parquet")]) if ret.exists() else []
    labels_files = [p for p in files if not p.name.startswith("adjudication")]
    truth = b2_truth_keys(run.truth_path) if run.truth_path is not None else None
    if not labels_files and simulate:
        sim = b2_simulate_labels(key, asg, truth, feats, cfg.seed)
        disp = b2_display(run, key, cfg)
        for a, g in sim.groupby("annotator_id"):
            mine_ids = asg.loc[asg["annotator_id"] == a].sort_values("position")["review_id"]
            mine = {p: d.set_index("review_id").reindex([r for r in mine_ids if r in set(d["review_id"])]).reset_index()
                    for p, d in disp.items()}
            labels_files.append(b2_write_review_workbook(ret / f"review-workbook_{a}_round00.xlsx", mine, a, labels=g,
                                                          fixed_time=fixed_time))
        log(f"  synthetic annotators filled {len(sim)} labels into {ret}")
    if not labels_files:
        log(f"  waiting for labels: none in {ret} (the packet is ready; evaluation not run)")
        return {"key": key, "assignments": asg, "status": "waiting for labels", "packet": r0}
    raw = b2_read_labels(labels_files)
    valid, issues, pending = b2_validate_labels(raw, key, asg, run)
    agreement = b2_agreement(valid)
    queue = b2_adjudication_queue(valid)
    b2_write_adjudication(rd / "round_00" / "adjudication" / "adjudication_round00.xlsx", run, key, queue, valid, cfg,
                          fixed_time=fixed_time)
    adj_files = sorted(ret.glob("adjudication*.xlsx")) if ret.exists() else []
    if not adj_files and simulate and len(queue):
        filled = b2_simulate_adjudication(queue, key, truth, cfg.seed)
        adj_files = [b2_write_adjudication(ret / "adjudication_round00_filled.xlsx", run, key, queue, valid, cfg,
                                           filled=filled, fixed_time=fixed_time)]
    adjud = pd.concat([b2_read_adjudication(p) for p in adj_files], ignore_index=True) if adj_files else None
    final = b2_final_labels(valid, adjud, key)
    log(f"  labels: {len(raw)} read, {len(valid)} valid, {len(pending)} pending, issues "
        f"{issues['issue'].value_counts().to_dict() if len(issues) else {}}; adjudication queue {len(queue)}; "
        f"final {final['final_label'].replace('', 'unresolved').value_counts().to_dict()}")
    gold = B2Gold(final)
    tuning = gold.tuning_view()
    fits, notes = b2_fit_all(run, feats, tuning, cfg)
    scores, th = b2_scores(run, feats, fits)
    tables = b2_evaluate_sealed(run, feats, gold, scores, th, cfg)
    tables["recall_ceiling"] = b2_recall_ceiling(run, tables["summary"], cfg)
    if truth is not None:
        tables["vs_truth"] = b2_truth_check(run, feats, gold, scores, th, cfg, truth)
    applied, apply_note = b2_apply(run, scores, th, tables["adoption"], cfg)
    man = b2_manifest_rows(run, cfg, key, issues, agreement, final, tables, notes, gold)
    # step 5: the next batch, from the current uncertainty and disagreement of every scorer
    sources = b2_parameter_spread(run, feats, cfg)
    sources.update({m: v for m, v in scores.items() if m in ("calibration", "semisup_em", "supervised")})
    k1, a1, _ = b2_next_batch(run, feats, key, sources, cfg, 1, rd / "round_01", fixed_time=fixed_time)
    card = b2_decision_card(feats, fits["supervised"]) if fits.get("supervised") is not None else pd.DataFrame()
    out = run.root / "out" / run.dataset / "b2"
    report = {"adoption": tables["adoption"], "summary": tables["summary"], "by_threshold": tables["by_threshold"],
              "by_band": tables["by_band"], "by_basis": tables["by_basis"], "reliability": tables["reliability"],
              "confusion": tables["confusion"], "recall_ceiling": tables["recall_ceiling"],
              "denominators": tables["denominators"], "agreement": agreement, "label_issues": issues,
              "em_m": fits["semisup_em"].table, "next_batch": k1, "b2_manifest": man, "schema": b2_schema_frame()}
    if "vs_truth" in tables:
        report["vs_truth"] = tables["vs_truth"]
    b2_write_report(out / "b2_report.xlsx", report)
    b2_write_table(applied, out / "b2_candidates.parquet")
    if len(card):
        b2_write_table(card, out / "b2_decision_card.parquet")
    b2_write_table(man, out / "b2_manifest.parquet")
    log(f"  fitted: {notes}")
    log(f"  sealed split: {tables['summary']['pairs'].iloc[0]} definite labels; sealed accesses {len(gold.access_log)}")
    log(f"  {apply_note}; next batch: {len(k1)} active pairs in {rd / 'round_01'}; report {out / 'b2_report.xlsx'} "
        f"({time.time() - t0:.1f}s)")
    return {"run": run, "feats": feats, "key": key, "assignments": asg, "raw": raw, "valid": valid, "issues": issues,
            "pending": pending, "agreement": agreement, "queue": queue, "final": final, "gold": gold, "fits": fits,
            "notes": notes, "scores": scores, "thresholds": th, "tables": tables, "applied": applied, "manifest": man,
            "next_key": k1, "next_assignments": a1, "card": card, "status": "evaluated"}


def b2_build_stubs(run, stub_dir, cfg=None):
    """The committed SME stub packets: a small fictional round from the synthetic run, exactly as
    handed to annotators, with a filled example set and an adjudication sheet."""
    base = cfg or B2Config()
    cfg = B2Config(**{**_asdict(base), "sample_size": 40, "min_per_stratum": 2, "annotators": ("A1", "A2"),
                      "double_code_share": 0.35})
    fixed = _dt.datetime(2026, 9, 28, 9, 0, 0)
    stub_dir = Path(stub_dir)
    feats = b2_features(run)
    key = b2_draw_sample(run, cfg, 0)
    # keep the stub readable: at most 3 pairs per stratum
    key = key[key.groupby("stratum").cumcount() < 3].reset_index(drop=True)
    key, asg = b2_assign(key, cfg)
    files = b2_write_packet(run, key, asg, stub_dir, cfg, fixed_time=fixed, round_id=0)
    truth = b2_truth_keys(run.truth_path)
    prof = {"A1": {"flip": 0.0, "unsure_hard": 0.35, "unsure_easy": 0.0}, "A2": {"flip": 0.08, "unsure_hard": 0.2, "unsure_easy": 0.03}}
    sim = b2_simulate_labels(key, asg, truth, feats, cfg.seed, profiles=prof, start=_dt.datetime(2026, 10, 5, 9, 0, 0))
    ex = stub_dir / "examples"
    disp = b2_display(run, key, cfg)
    for a, g in sim.groupby("annotator_id"):
        mine_ids = asg.loc[asg["annotator_id"] == a].sort_values("position")["review_id"]
        mine = {p: d.set_index("review_id").reindex([r for r in mine_ids if r in set(d["review_id"])]).reset_index()
                for p, d in disp.items()}
        b2_write_review_workbook(ex / f"review-workbook_{a}_round00_FILLED.xlsx", mine, a, labels=g, fixed_time=fixed)
        b2_write_table(g[["review_id", "annotator_id", "label", "confidence", "reason", "labelled_at"]],
                       ex / f"labels_{a}_round00_FILLED.csv")
    raw = b2_read_labels(sorted(ex.glob("labels_*_FILLED.csv")))
    valid, issues, _ = b2_validate_labels(raw, key, asg, run)
    queue = b2_adjudication_queue(valid)
    b2_write_adjudication(stub_dir / "to-adjudicator" / "adjudication_round00.xlsx", run, key, queue, valid, cfg, fixed_time=fixed)
    filled = b2_simulate_adjudication(queue, key, truth, cfg.seed, error=0.0, start=_dt.datetime(2026, 10, 12, 9, 0, 0))
    b2_write_adjudication(ex / "adjudication_round00_FILLED.xlsx", run, key, queue, valid, cfg, filled=filled, fixed_time=fixed)
    final = b2_final_labels(valid, b2_read_adjudication(ex / "adjudication_round00_FILLED.xlsx"), key)
    b2_write_table(final.drop(columns=["p_at_sampling"]), ex / "final-labels_round00_EXAMPLE.csv")
    b2_write_table(b2_agreement(valid), ex / "agreement_round00_EXAMPLE.csv")
    return {"key": key, "assignments": asg, "labels": sim, "queue": queue, "final": final, "files": files}


# %% [code]
if __name__ == "__main__" or "CAND_OUT" in globals():
    B2CFG = B2Config()
    B2RUN = b2_load_run()
    print(f"B2 on {B2RUN.dataset} ({B2RUN.source}); run {B2RUN.run_id}, core v{B2RUN.core_version}; "
          f"{len(B2RUN.candidates):,} candidate pairs")
    B2 = b2_run(B2RUN, B2CFG)
    if B2["status"] == "evaluated":
        with pd.option_context("display.width", 200, "display.max_columns", 30):
            print(B2["tables"]["adoption"][["improvement", "adopted", "metric", "b1", "candidate", "ci_low", "ci_high", "reason"]].to_string(index=False))
            print(B2["tables"]["summary"].to_string(index=False, float_format=lambda v: f"{v:.4g}"))
            print(B2["agreement"].to_string(index=False, float_format=lambda v: f"{v:.3f}"))
    if os.environ.get("RL_B2_WRITE_STUBS") == "1" and B2RUN.truth_path is not None:
        B2STUB = b2_build_stubs(B2RUN, B2RUN.root / "review" / "stubs")
        print(f"stub packets written: {len(B2STUB['key'])} pairs, {len(B2STUB['queue'])} to adjudicate")


# %% [markdown]
# ### B2.11 · Self-tests
#
# Each component on small constructed data or on the run's own tables, with labels made from the
# synthetic truth file with injected annotator noise. The sealed split is checked two ways: every
# fitting function refuses a sealed view, and changing every sealed label leaves every fitted
# parameter unchanged. The end-to-end test runs the whole loop on the synthetic run and checks
# that each improvement is adopted only when its sealed interval is above zero; a deliberately
# miscalibrated B1 must lead calibration to be adopted.
#
# **Where Splink would do better.** Splink's own test suite covers its evaluation functions;
# nothing here depends on it.

# %% [code]
import unittest as _b2_unittest


class B2TestUnits(_b2_unittest.TestCase):
    def test_standardize_renames_and_names_missing(self):
        d = pd.DataFrame({"pair_id": ["a"], "x_record_id": ["x"], "watchlist_record_id": ["w"], "part": ["person"],
                          "p": [0.5], "bits": [1.0], "prior_logit": [-3.0], "basis": ["name_only"], "veto": [""]})
        s = b2_standardize(d, "candidates")
        self.assertIn("extracted_record_id", s.columns); self.assertIn("prior_bits", s.columns)
        with self.assertRaises(KeyError) as cm:
            b2_standardize(d.drop(columns=["basis"]), "candidates")
        self.assertIn("basis", str(cm.exception))

    def test_cohen_matches_sklearn(self):
        from sklearn.metrics import cohen_kappa_score
        rng = np.random.default_rng(1)
        a = rng.choice(list(B2_LABELS), 200); b = np.where(rng.random(200) < 0.7, a, rng.choice(list(B2_LABELS), 200))
        self.assertAlmostEqual(b2_cohen_kappa(a, b, list(B2_LABELS)), cohen_kappa_score(a, b, labels=list(B2_LABELS)), places=10)

    def test_fleiss_textbook(self):
        # Fleiss (1971) worked example as tabulated on Wikipedia: kappa = 0.210
        t = [[0, 0, 0, 0, 14], [0, 2, 6, 4, 2], [0, 0, 3, 5, 6], [0, 3, 9, 2, 0], [2, 2, 8, 1, 1],
             [7, 7, 0, 0, 0], [3, 2, 6, 3, 0], [2, 5, 3, 2, 2], [6, 5, 2, 1, 0], [0, 2, 2, 3, 7]]
        self.assertAlmostEqual(b2_fleiss_kappa(t), 0.210, places=3)

    def test_ht_metrics_by_hand(self):
        y = np.array([1, 0, 1, 0.]); p = np.array([0.9, 0.8, 0.2, 0.1]); w = np.array([1, 3, 2, 4.])
        r = b2_prf(y, p, w, 0.5)
        self.assertAlmostEqual(r["precision"], 1 / 4); self.assertAlmostEqual(r["recall"], 1 / 3)
        self.assertAlmostEqual(b2_brier(y, p, w), (0.01 + 3 * 0.64 + 2 * 0.64 + 4 * 0.01) / 10)
        self.assertAlmostEqual(b2_avg_precision(np.array([1, 0, 1.]), np.array([.9, .8, .7]), np.ones(3)), (1 + 2 / 3) / 2)

    def test_poststrat_weights_sum_to_population(self):
        d = pd.DataFrame({"stratum": ["a", "a", "a", "b"], "N_h": [30, 30, 30, 5], "y": [1, 0, np.nan, 1.],
                          "selection": "stratified"})
        w = b2_poststrat_weights(d)
        self.assertAlmostEqual(w.sum(), 35); self.assertTrue(np.isnan(w[2]))

    def test_allocation_oversamples_uncertain_and_respects_caps(self):
        N = pd.Series({"p>=0.9|identifier": 10000, "0.1-0.5|name_only": 400, "p<0.1|none": 10000, "vetoed|none": 3})
        n = b2_allocate(N, B2Config(), 600)
        rate = n / N
        self.assertGreater(rate["0.1-0.5|name_only"], rate["p>=0.9|identifier"])
        self.assertLessEqual(n["vetoed|none"], 3)
        self.assertTrue((n <= N).all())

    def test_split_is_deterministic_per_group(self):
        g = [f"X{i}|person" for i in range(3000)]
        a = b2_split_of(g, 5, B2Config().split_fractions); b = b2_split_of(g, 5, B2Config().split_fractions)
        self.assertTrue((a == b).all())
        share = pd.Series(a).value_counts(normalize=True)
        self.assertAlmostEqual(share["sealed_test"], 0.4, delta=0.03)

    def test_label_validation(self):
        key = pd.DataFrame({"review_id": ["R1", "R2", "R3"], "pair_id": ["a", "b", "c"]})
        asg = pd.DataFrame({"review_id": ["R1", "R2", "R3", "R1"], "annotator_id": ["A1", "A1", "A1", "A2"]})
        raw = pd.DataFrame({"review_id": ["R1", "R1", "R2", "R2", "R3", "R9", "R1", "R3"],
                            "annotator_id": ["A1", "A1", "A1", "A1", "A1", "A1", "A2", "A2"],
                            "label": ["Match", "match", "match", "different", "maybe", "match", "cant_tell", "match"],
                            "confidence": ["high", "high", "", "", "", "", "", ""], "reason": [""] * 8,
                            "labelled_at": [""] * 8, "timestamp_source": "file", "source_file": "t.csv",
                            "source_sheet": "", "source_row": range(2, 10)})
        v, iss, pend = b2_validate_labels(raw, key, asg)
        got = set(iss["issue"])
        self.assertTrue({"duplicate", "contradictory", "invalid_label", "unknown_review_id", "not_assigned", "reason_missing"} <= got)
        self.assertEqual(sorted(zip(v["review_id"], v["annotator_id"], v["label"])),
                         [("R1", "A1", "match"), ("R1", "A2", "unsure")])

    def test_resolution_and_adjudication(self):
        key = pd.DataFrame({"review_id": ["R1", "R2", "R3"], "pair_id": ["a", "b", "c"], "split": "train"})
        v = pd.DataFrame({"review_id": ["R1", "R1", "R2", "R2", "R3"], "annotator_id": ["A1", "A2", "A1", "A2", "A1"],
                          "label": ["match", "match", "match", "non-match", "unsure"]})
        q = b2_adjudication_queue(v)
        self.assertEqual(sorted(q["review_id"]), ["R2", "R3"])
        f = b2_final_labels(v, pd.DataFrame({"review_id": ["R2"], "final_label": ["non-match"], "adjudicator_id": ["J"],
                                             "adjudication_reason": ["x"]}), key).set_index("review_id")
        self.assertEqual(f.loc["R1", "label_source"], "unanimous"); self.assertEqual(f.loc["R2", "final_label"], "non-match")
        self.assertEqual(f.loc["R3", "final_label"], "unsure"); self.assertTrue(np.isnan(f.loc["R3", "y"]))

    def test_io_roundtrip(self):
        import tempfile
        d = pd.DataFrame({"review_id": ["R1"], "annotator_id": ["A1"], "label": ["match"], "confidence": ["high"],
                          "reason": ["same SSN"], "labelled_at": ["2026-10-01"]})
        with tempfile.TemporaryDirectory() as tmp:
            for ext in (".csv", ".parquet"):
                b2_write_table(d, Path(tmp) / f"l{ext}")
                r = b2_read_labels([Path(tmp) / f"l{ext}"])
                self.assertEqual(r.loc[0, "label"], "match"); self.assertEqual(r.loc[0, "timestamp_source"], "cell")


class B2TestOnRun(_b2_unittest.TestCase):
    """Tests on the pipeline's own tables (the synthetic run when the truth file is there)."""

    @classmethod
    def setUpClass(cls):
        g = globals()
        cls.rrun = g.get("B2RUN") or b2_load_run()
        cls.cfg = B2Config(bootstrap_reps=300)
        cls.feats = b2_features(cls.rrun)
        cls.truth = b2_truth_keys(cls.rrun.truth_path) if cls.rrun.truth_path is not None else None

    def _labelled(self, cfg=None, seed=None):
        cfg = cfg or self.cfg
        key, asg = b2_assign(b2_draw_sample(self.rrun, cfg, 0), cfg)
        sim = b2_simulate_labels(key, asg, self.truth, self.feats, seed or cfg.seed)
        sim = sim.assign(timestamp_source="cell", source_file="sim", source_sheet="", source_row=0)
        valid, _, _ = b2_validate_labels(sim, key, asg, self.rrun)
        q = b2_adjudication_queue(valid)
        final = b2_final_labels(valid, b2_simulate_adjudication(q, key, self.truth, cfg.seed), key)
        return key, asg, valid, final

    def test_rescore_reproduces_pipeline_bits(self):
        p, total, _ = b2_rescore(self.feats)
        self.assertLess(float(np.abs(total - self.rrun.candidates["bits"].to_numpy()).max()), 1e-6)
        ok = ~self.feats.vetoed
        self.assertLess(float(np.abs(p[ok] - self.feats.p[ok]).max()), 1e-9)

    def test_sample_weights_and_blind_packet(self):
        import tempfile, openpyxl
        key, asg = b2_assign(b2_draw_sample(self.rrun, self.cfg, 0), self.cfg)
        fr = b2_frame(self.rrun, self.cfg)
        est = key.groupby("stratum")["weight"].sum()
        pop = fr.groupby("stratum").size().reindex(est.index)
        self.assertLess(float((est - pop).abs().max()), 1e-6)            # HT weights recover N_h
        self.assertEqual(key.groupby("group")["split"].nunique().max(), 1)
        self.assertEqual(key["pair_id"].nunique(), len(key))
        with tempfile.TemporaryDirectory() as tmp:
            small = key.head(40)
            k2, a2 = b2_assign(small, self.cfg)
            b2_write_packet(self.rrun, k2, a2, tmp, self.cfg)
            forbidden = ["p_at_sampling", "stratum", "sealed_test", "inclusion_prob", "name_only", "bits", "basis"]
            ids = set(k2["extracted_record_id"]) | set(k2["watchlist_record_id"])
            for f in Path(tmp, "to-annotators").rglob("*"):
                if f.suffix == ".xlsx":
                    wb = openpyxl.load_workbook(f)
                    for ws in wb.worksheets:
                        if ws.title in B2_SHEETS.values():
                            self.assertTrue(ws.data_validations.dataValidation, "label drop-downs missing")
                            self.assertTrue(ws.protection.sheet)
                        for row in ws.iter_rows(values_only=True):
                            for v in row:
                                s = str(v or "")
                                self.assertFalse(any(t in s for t in forbidden), f"{f.name}: '{s}' reveals hidden data")
                                self.assertNotIn(s, ids)
                elif f.suffix in (".csv", ".md"):
                    txt = f.read_text(encoding="utf-8")
                    self.assertFalse(any(t in txt for t in forbidden[:5]), f.name)
            ck = pd.read_csv(Path(tmp, "coordinator", "sample-key_round00.csv"))
            self.assertIn("stratum", ck.columns); self.assertIn("split", ck.columns)

    def test_workbook_roundtrip(self):
        import tempfile
        key, asg = b2_assign(b2_draw_sample(self.rrun, self.cfg, 0).head(30), self.cfg)
        labs = pd.DataFrame({"review_id": key["review_id"], "label": "match", "confidence": "high", "reason": "r",
                             "labelled_at": "2026-10-01T10:00:00"})
        with tempfile.TemporaryDirectory() as tmp:
            p = b2_write_review_workbook(Path(tmp) / "w.xlsx", b2_display(self.rrun, key, self.cfg), "A1", labels=labs)
            r = b2_read_labels([p])
        self.assertEqual(sorted(r["review_id"]), sorted(key["review_id"]))
        self.assertTrue((r["label"] == "match").all()); self.assertTrue((r["timestamp_source"] == "cell").all())

    def test_active_learning_respects_splits(self):
        key, _ = b2_assign(b2_draw_sample(self.rrun, self.cfg, 0), self.cfg)
        src = b2_parameter_spread(self.rrun, self.feats, self.cfg)
        k1 = b2_select_active(self.rrun, self.feats, key, src, self.cfg, 1)
        self.assertTrue((k1["split"] == "train").all())
        self.assertFalse(k1["pair_id"].isin(key["pair_id"]).any())
        self.assertLessEqual(k1.groupby("group").size().max(), self.cfg.al_per_entity)
        self.assertTrue(k1["weight"].isna().all())

    def test_gb_contributions_are_exact(self):
        if self.truth is None:
            self.skipTest("no truth file")
        _, _, _, final = self._labelled()
        sup, _ = b2_fit_supervised(self.feats, B2Gold(final).tuning_view(), self.cfg)
        X, _ = b2_gb_matrix(self.feats)
        bias, con = b2_gb_contributions(sup, X[:500])
        self.assertLess(float(np.ptp(bias)), 1e-8)
        self.assertLess(float(np.abs(bias + con.sum(1) - sup.model.decision_function(X[:500])).max()), 1e-8)
        self.assertTrue((b2_predict_supervised(self.feats, sup)[self.feats.vetoed] == 0).all())

    def test_sealed_split_never_reaches_tuning(self):
        if self.truth is None:
            self.skipTest("no truth file")
        key, asg, valid, final = self._labelled()
        gold = B2Gold(final)
        for fn in (lambda t: b2_fit_calibration(self.feats, t, self.cfg), lambda t: b2_fit_supervised(self.feats, t, self.cfg),
                   lambda t: b2_fit_semisup_em(self.rrun, self.feats, t, self.cfg),
                   lambda t: b2_fit_cost_threshold(self.feats, t, self.cfg), lambda t: b2_fit_all(self.rrun, self.feats, t, self.cfg)):
            with self.assertRaises(B2SealedAccessError):
                fn(gold.sealed_view("test: must be refused"))
            leak = gold.tuning_view()
            leak = pd.concat([leak, gold.sealed_view("test").head(1)]); leak.attrs["b2_view"] = "tuning"
            with self.assertRaises(B2SealedAccessError):
                fn(leak)
        # changing every sealed label leaves every fitted parameter unchanged
        flipped = final.copy()
        s = flipped["split"] == "sealed_test"
        flipped.loc[s, "y"] = 1 - flipped.loc[s, "y"]
        flipped.loc[s, "final_label"] = flipped.loc[s, "final_label"].map({"match": "non-match", "non-match": "match"}).fillna("unsure")
        fa, _ = b2_fit_all(self.rrun, self.feats, B2Gold(final).tuning_view(), self.cfg)
        fb, _ = b2_fit_all(self.rrun, self.feats, B2Gold(flipped).tuning_view(), self.cfg)
        sa, _ = b2_scores(self.rrun, self.feats, fa); sb, _ = b2_scores(self.rrun, self.feats, fb)
        for m in sa:
            self.assertTrue(np.array_equal(sa[m], sb[m]), m)
        self.assertEqual(fa["cost_threshold"], fb["cost_threshold"])
        self.assertEqual(fa["semisup_em"].m_lookup, fb["semisup_em"].m_lookup)

    def test_end_to_end_adoption_rule(self):
        if self.truth is None:
            self.skipTest("no truth file")
        key, asg, valid, final = self._labelled()
        gold = B2Gold(final)
        fits, notes = b2_fit_all(self.rrun, self.feats, gold.tuning_view(), self.cfg)
        scores, th = b2_scores(self.rrun, self.feats, fits)
        t = b2_evaluate_sealed(self.rrun, self.feats, gold, scores, th, self.cfg, "e2e test")
        ad = t["adoption"]
        self.assertEqual(set(ad["improvement"]), set(B2_IMPROVEMENTS))
        for _, r in ad.iterrows():
            if r["adopted"]:
                self.assertGreater(r["ci_low"], 0); self.assertGreater(r["improvement_estimate"], 0)
            else:
                self.assertTrue(r["reason"])
        self.assertEqual([a["purpose"] for a in gold.access_log], ["e2e test"])
        # a deliberately miscalibrated B1 (p squeezed into 0.3-0.7) must be fixed by calibration
        bad = B2Features(**{**self.feats.__dict__, "p": np.where(self.feats.vetoed, 0.0, 0.3 + 0.4 * self.feats.p)})
        g2 = B2Gold(final)
        f2 = {"calibration": b2_fit_calibration(bad, g2.tuning_view(), self.cfg)[0]}
        s2, th2 = b2_scores(self.rrun, bad, f2)
        a2 = b2_evaluate_sealed(self.rrun, bad, g2, s2, th2, self.cfg, "e2e miscalibrated")["adoption"].set_index("improvement")
        self.assertTrue(bool(a2.loc["calibration", "adopted"]), a2.loc["calibration"].to_dict())


def b2_run_selftests(verbosity=1):
    import io as _io
    buf = _io.StringIO()
    suite = _b2_unittest.TestSuite()
    for c in (B2TestUnits, B2TestOnRun):
        suite.addTests(_b2_unittest.defaultTestLoader.loadTestsFromTestCase(c))
    res = _b2_unittest.TextTestRunner(stream=buf, verbosity=verbosity).run(suite)
    return res, buf.getvalue()


if __name__ == "__main__" or "CAND_OUT" in globals():
    B2_SELFTEST, B2_SELFTEST_TEXT = b2_run_selftests()
    print(B2_SELFTEST_TEXT[-3000:])
    print(f"B2 self-tests: {B2_SELFTEST.testsRun} run, {len(B2_SELFTEST.failures)} failures, {len(B2_SELFTEST.errors)} errors, "
          f"{len(B2_SELFTEST.skipped)} skipped")
    assert B2_SELFTEST.wasSuccessful(), "B2 self-tests failed"
