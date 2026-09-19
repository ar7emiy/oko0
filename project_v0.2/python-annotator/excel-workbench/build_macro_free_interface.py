"""Build the macro-free annotation workbench from workbench-schema.json.

Port of build_macro_free_interface.mjs off the Codex-only @oai/artifact-tool
and onto the openpyxl vendored in this repository, so the delivered workbook
can be regenerated on any machine with Python and without a Codex runtime.

Layout, sheet order, column widths, row heights and colours follow the .mjs
original. What is NOT reproduced: wb.render() sheet PNGs, which needed the
Office rendering service. Use preview_workbook.py for a no-Excel substitute.

This writes a .xlsx only. It does not execute VBA and does not open Excel.
"""
import json
import re
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

ROOT = Path(__file__).resolve().parent
try:
    import openpyxl  # noqa: F401
except ModuleNotFoundError:
    sys.path.insert(0, str(ROOT.parents[1] / 'offline-python-annotator' / 'vendor'))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table

NAVY = 'FF18344C'
AMBER = 'FFFFF4CF'
NOTE_FILL = 'FFEDF3F7'
WHITE = 'FFFFFFFF'
OUT = ROOT / 'Entity-Gold-Review-Macro-Free.xlsx'

schemas = json.loads((ROOT / 'workbench-schema.json').read_text(encoding='utf-8'))
wb = Workbook()
wb.remove(wb.active)


def base_font(size=11, bold=False, color=NAVY):
    return Font(name='Arial', size=size, bold=bold, color=color)


def page(name, widths):
    """New sheet: gridlines off, given column widths, Arial 11 navy over A1:J60."""
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    font = base_font()
    for row in range(1, 61):
        for col in range(1, 11):
            ws.cell(row=row, column=col).font = font
    return ws


def text(ws, ref, value):
    """Write a literal string. A leading '=' stays literal text, never a formula."""
    cell = ws[ref]
    cell.value = value
    if isinstance(value, str) and value.startswith('='):
        cell.data_type = 's'
    return cell


def title(ws, value):
    cell = text(ws, 'B2', value)
    cell.font = base_font(size=16, bold=True)
    ws.row_dimensions[2].height = 30


def field(ws, row, label, value=''):
    """Label in column B, amber text-formatted entry cell in column C."""
    text(ws, f'B{row}', label)
    cell = text(ws, f'C{row}', value)
    cell.fill = PatternFill('solid', start_color=AMBER, end_color=AMBER)
    cell.number_format = '@'
    cell.alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[row].height = 27


def button(ws, ref, label):
    """Navy caption cell that SetupWorkbench anchors a real button over."""
    cell = text(ws, ref, label)
    cell.fill = PatternFill('solid', start_color=NAVY, end_color=NAVY)
    cell.font = base_font(bold=True, color=WHITE)
    cell.alignment = Alignment(wrap_text=True)
    ws.row_dimensions[cell.row].height = 30


def dropdown(ws, ref, values):
    joined = ','.join(values)
    assert len(joined) <= 253, f'inline list validation too long for {ref}'
    dv = DataValidation(type='list', formula1=f'"{joined}"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(ws[ref])


# --- Review Desk -----------------------------------------------------------
d = page('Review Desk', [3, 23, 59, 3, 27, 3, 76])
title(d, 'Annotation workbench')
text(d, 'C3', 'Start: load a note. Yellow cells are editable.')
for row, label, value in [
    (4, 'Claim number', ''), (5, 'Note ID', ''), (6, 'Source version', ''),
    (7, 'Reviewer', ''), (9, 'Entry kind', 'entity'), (10, 'Exact source quote', ''),
    (11, 'Occurrence (1, 2...)', '1'), (12, 'Entity reference', ''),
    (13, 'Display name', ''), (14, 'Entity type', ''), (15, 'Reference form', ''),
    (16, 'Field kind', ''), (17, 'Field value', ''), (18, 'Second entity', ''),
    (19, 'Optional open label', ''), (20, 'Context / reason', ''),
]:
    field(d, row, label, value)
d.row_dimensions[10].height = 90
d.row_dimensions[20].height = 65
dropdown(d, 'C9', ['entity', 'reference', 'field', 'context', 'statement', 'uncertain'])
dropdown(d, 'C15', ['name', 'alias', 'pronoun', 'description'])
dropdown(d, 'C16', ['entity_name', 'address', 'city', 'state', 'zip_code', 'phone', 'TIN', 'other'])
text(d, 'B22', 'Saved state')
text(d, 'C22', 'Not wired. Follow Instructions to activate buttons.').alignment = Alignment(wrap_text=True)
text(d, 'B24', 'Attestation')
text(d, 'C24', 'I confirm that I have read enough of the source note to validate this entry, '
               'its supporting evidence, and its links to other entities. I have corrected any '
               'errors I identified.').alignment = Alignment(wrap_text=True)
d.row_dimensions[24].height = 85
text(d, 'G4', 'Note text and source evidence')
text(d, 'G5', 'Load a note to read it here. Next text page reaches every passage.')
for row in range(5, 21):
    cell = d.cell(row=row, column=7)
    cell.fill = PatternFill('solid', start_color=NOTE_FILL, end_color=NOTE_FILL)
    cell.alignment = Alignment(wrap_text=True, vertical='top')
text(d, 'G23', 'Entities in this claim')
text(d, 'G24', 'Choose entity lists current claim entities.')
for ref, label in [
    ('E4', 'Select note'), ('E5', 'New manual entry'), ('E7', 'Load next AI draft'),
    ('E9', 'Attest & Save'), ('E11', 'Save draft'), ('E12', 'Previous entry'),
    ('E13', 'Next reviewed entry'), ('E15', 'Delete entry'), ('E17', 'Choose entity'),
    ('E18', 'Locate evidence'), ('E20', 'Complete note'), ('E22', 'Previous text page'),
    ('E23', 'Next text page'), ('E25', 'Watchlist review'),
]:
    button(d, ref, label)

# --- Notes paste page ------------------------------------------------------
n = page('Notes', [3, 23, 75, 3, 28])
title(n, 'Add claim notes')
field(n, 4, 'Claim number')
field(n, 5, 'Note ID')
field(n, 6, 'Reading order', '1')
text(n, 'C8', 'Copy the complete note in Notepad. Click Paste whole note. '
              'Do not paste a long note into a single cell.').alignment = Alignment(wrap_text=True)
n.row_dimensions[8].height = 55
for ref, label in [('E4', 'Paste whole note'), ('E6', 'Import TXT files'), ('E8', 'Open selected note')]:
    text(n, ref, label)

# --- AI Instructions -------------------------------------------------------
ai = page('AI Instructions', [3, 24, 110])
title(ai, 'Optional AI drafting')
prompt_lines = [ln for ln in (ROOT / 'COPILOT-PROMPT.md').read_text(encoding='utf-8').split('\n') if ln.strip()]
for i, line in enumerate(prompt_lines):
    cell = text(ai, f'C{i + 4}', line)
    cell.alignment = Alignment(wrap_text=True)
    ai.row_dimensions[i + 4].height = max(28, -(-len(line) // 110) * 18)

# --- AI Analysis Input -----------------------------------------------------
a = page('AI Analysis Input', [3, 24, 105])
title(a, 'AI analysis for current note')
for row, label, value in [(4, 'Run label', ''), (5, 'Processed parts', ''),
                          (6, 'Run complete?', 'no'), (7, 'Brief analysis', '')]:
    field(a, row, label, value)
a.row_dimensions[7].height = 100
text(a, 'C9', 'Save queue snapshot archives analysis for the current note. '
              'Analysis retention: last five distinct notes pasted.').alignment = Alignment(wrap_text=True)
a.row_dimensions[9].height = 45

# --- AI Queue (UI header above the tInput table at row 5) ------------------
q = page('AI Queue', [3])
title(q, 'AI proposals')
text(q, 'B3', 'Run SaveQueueSnapshot after Copilot finishes. Drafts are not accepted annotations.')

# --- Watchlist Desk --------------------------------------------------------
w = page('Watchlist Desk', [3, 23, 60, 3, 25, 3, 76])
title(w, 'Watchlist identity review')
for row, label in [(4, 'Client row ID'), (5, 'Identity decision'), (6, 'Reason / evidence')]:
    field(w, row, label)
w.row_dimensions[6].height = 100
dropdown(w, 'C5', ['same_entity', 'different_entity', 'insufficient_evidence'])
for ref, label in [('E4', 'Load next pair'), ('E6', 'Save decision'), ('E8', 'Import client output')]:
    text(w, ref, label)
text(w, 'G4', 'System record and watchlist record')

# --- Instructions ----------------------------------------------------------
ins = page('Instructions', [3, 24, 115])
title(ins, 'Start here')
STEPS = [
    'Save a copy as Excel Macro-Enabled Workbook (.xlsm). Keep this original blank template.',
    'Alt+F11 > File > Import File: import all eight supplied .bas modules.',
    'Separately, double-click ThisWorkbook under Microsoft Excel Objects and paste the whole of ThisWorkbook.txt into its empty code pane. This file cannot be imported, only pasted. Skipping it leaves the workbook without its reopen and close guards.',
    'Alt+F8 > SetupWorkbench > Run creates the buttons. No manual shape wiring.',
    'Notes: enter IDs, copy the full note, then Paste whole note. Or import UTF-8 TXT files named CLAIM_NOTE.txt.',
    'Review Desk: choose entry kind, paste exact evidence, select occurrence and fill relevant fields. Choose entity lists accepted entities.',
    'Attest & Save validates and saves. Save draft preserves unfinished work. Manual entry needs no Copilot.',
    'Optional AI: follow AI Instructions. SaveQueueSnapshot validates drafts, then Load next AI draft fills the form.',
    'Previous entry restores your latest revision. Changed entries require new attestations; Delete preserves history.',
    'Complete note requires reading all text, adding missing annotations and resolving drafts. Watchlist review follows.',
    'Only analysis for the last five distinct notes pasted is retained. Sources, pending queues and accepted records persist.',
    'One SME writer per workbook. Do not edit with Copilot while macros save. Do not edit backend tables.',
    'Keep older workbooks untouched. Read WIRE-UP-IN-EXCEL.md before bringing historical annotations into this layout.',
    'Licensed Windows Excel runtime qualification is still required. Macros do not run in Excel for the web.',
]
for i, line in enumerate(STEPS):
    text(ins, f'B{4 + i}', f'{i + 1}.')
    text(ins, f'C{4 + i}', line).alignment = Alignment(wrap_text=True)
    ins.row_dimensions[4 + i].height = 45

# --- Backend tables --------------------------------------------------------
WIDE = re.compile(r'text|quote|reason|analysis|value|label')
for name, definition in schemas.items():
    ws = q if name == 'AI Queue' else page(name, [3])
    row = 5 if name == 'AI Queue' else 1
    headers = definition['headers']
    for i, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=i, value=header)
        cell.fill = PatternFill('solid', start_color=NAVY, end_color=NAVY)
        cell.font = base_font(bold=True, color=WHITE)
        cell.alignment = Alignment(wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = 48 if WIDE.search(header) else 22
    ws.row_dimensions[row].height = 42
    for r in range(row + 1, row + 11):
        for i in range(1, len(headers) + 1):
            ws.cell(row=r, column=i).number_format = '@'
    last = get_column_letter(len(headers))
    ws.add_table(Table(displayName=definition['table'], ref=f'A{row}:{last}{row + 1}'))
    ws.freeze_panes = f'A{row + 1}'

wb.save(OUT)

# --- Post-write self-check on the file that was actually written -----------
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
with zipfile.ZipFile(OUT) as z:
    sheet_names = [s.attrib['name'] for s in ET.fromstring(z.read('xl/workbook.xml')).find('s:sheets', NS)]
    tables = {}
    errors = merged = 0
    for entry in z.namelist():
        if re.match(r'xl/tables/table\d+\.xml$', entry):
            t = ET.fromstring(z.read(entry))
            tables[t.attrib['name']] = [c.attrib['name'] for c in t.find('s:tableColumns', NS)]
        if re.match(r'xl/worksheets/sheet\d+\.xml$', entry):
            sheet = ET.fromstring(z.read(entry))
            merged += sheet.find('s:mergeCells', NS) is not None
            errors += sum(c.attrib.get('t') == 'e' for c in sheet.findall('.//s:c', NS))
    assert set(schemas).issubset(sheet_names), 'a backend page is missing from the workbook'
    assert tables == {v['table']: v['headers'] for v in schemas.values()}, 'table columns drifted from the schema'
    assert not merged, 'merged cells present; VBA writes into these ranges'
    assert not errors, 'formula error cells present'
    assert not any('vbaProject' in e for e in z.namelist()), 'workbook is not macro-free'

print(f'Built {OUT.name}: {len(sheet_names)} sheets, {len(tables)} tables, '
      f'no merges, no error cells, macro-free. VBA and Excel rendering remain unverified.')
