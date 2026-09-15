"""Render the built workbook to a single offline HTML page for eyeball review.

Substitute for the .mjs pipeline's wb.render() sheet PNGs, which required the
Codex Office rendering service. This reads the .xlsx that was actually written
and reproduces cell text, fills, bold/white header runs, column widths and
merges (there should be none) closely enough to spot layout mistakes without
a licensed Excel.

It is a reading aid, not a fidelity claim: no fonts, wrapping, row heights or
data-validation dropdowns are simulated. Layout accepted here still has to
pass the QA-ACCEPTANCE.md run in Windows Excel.
"""
import html
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent
try:
    import openpyxl  # noqa: F401
except ModuleNotFoundError:
    sys.path.insert(0, str(ROOT.parents[1] / 'offline-python-annotator' / 'vendor'))

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

SRC = ROOT / 'Entity-Gold-Review-Macro-Free.xlsx'
OUT = ROOT / 'workbook-preview.html'
MAX_ROWS, MAX_COLS = 30, 14


def rgb(color):
    """openpyxl colour -> css, ignoring theme/indexed colours we do not resolve."""
    if color is None or color.type != 'rgb' or not isinstance(color.rgb, str):
        return None
    value = color.rgb[-6:]
    return None if value.upper() in ('000000', 'FFFFFF') and color.rgb.startswith('00') else '#' + value


wb = load_workbook(SRC)
parts = ['<title>Workbench preview</title>', '''<style>
:root{color-scheme:light}body{margin:0;padding:24px;background:#eef3f6;color:#18344c;
font:13px/1.5 Arial,sans-serif}h1{font-size:22px}h2{font-size:16px;margin:28px 0 6px}
.wrap{overflow-x:auto;background:#fff;border:1px solid #cad9df;border-radius:8px;padding:10px}
table{border-collapse:collapse}td,th{border:1px solid #dce5ea;padding:3px 6px;vertical-align:top;
max-width:420px;overflow:hidden;text-overflow:ellipsis;white-space:pre-wrap}
th{background:#f4f8fa;color:#506772;font-weight:normal;font-size:11px;text-align:center}
.note{color:#506772;margin:4px 0 0}</style>''',
         '<h1>Workbench preview</h1>',
         f'<p class="note">Generated from {html.escape(SRC.name)}. '
         f'First {MAX_ROWS} rows and {MAX_COLS} columns of each sheet. '
         'Approximate rendering for layout review only; not an Excel fidelity check.</p>']

for ws in wb.worksheets:
    parts.append(f'<h2>{html.escape(ws.title)}</h2><div class="wrap"><table>')
    header = ''.join(f'<th>{get_column_letter(c)}</th>' for c in range(1, MAX_COLS + 1))
    parts.append(f'<tr><th></th>{header}</tr>')
    for r in range(1, MAX_ROWS + 1):
        cells = [f'<th>{r}</th>']
        for c in range(1, MAX_COLS + 1):
            cell = ws.cell(row=r, column=c)
            style = []
            fill = rgb(cell.fill.fgColor) if cell.fill and cell.fill.patternType else None
            if fill:
                style.append(f'background:{fill}')
            if cell.font and cell.font.bold:
                style.append('font-weight:bold')
            font_color = rgb(cell.font.color) if cell.font else None
            if font_color:
                style.append(f'color:{font_color}')
            width = ws.column_dimensions[get_column_letter(c)].width
            if width:
                style.append(f'min-width:{min(round(width * 7), 300)}px')
            attr = f' style="{";".join(style)}"' if style else ''
            value = '' if cell.value is None else html.escape(str(cell.value))
            cells.append(f'<td{attr}>{value}</td>')
        parts.append('<tr>' + ''.join(cells) + '</tr>')
    parts.append('</table></div>')
    if ws.merged_cells.ranges:
        parts.append(f'<p class="note">MERGED CELLS PRESENT: {ws.merged_cells.ranges}</p>')

OUT.write_text('\n'.join(parts), encoding='utf-8')
print(f'Wrote {OUT.name} for {len(wb.worksheets)} sheets. Layout aid only; Excel rendering unverified.')
