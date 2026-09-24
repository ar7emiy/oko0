"""Offline, reader-first workbench for building a claim-level gold dataset.

The reviewer only reads source notes and makes evidence decisions. This module
creates the traceable workbook in the background. It has no network calls and
does not modify a source note.
"""
from __future__ import annotations

import copy
import csv
import json
import re
import sys
import tempfile
import tkinter as tk
from pathlib import Path
from tkinter import filedialog, messagebox, simpledialog, ttk

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

REQUIRED = [
    "claim_number", "recordType", "entity_category_name", "entity_subcategory_name",
    "entity_name", "entity_address", "entity_city", "entity_state", "entity_zip_code",
    "entity_phone", "entity_TIN", "entity_NER_tag", "Entity_Watchlist_flag",
    "Exact_search_Note_ID", "Exact_search_matched_watchlist_entity_name",
    "GenAI_search_Note_ID", "GenAI_entityNameClearned", "GenAI_tok_sort_similarity",
    "Watchlist_entity_id", "Watchlist_entity_name", "watchlist_address", "watchlist_city",
    "watchlist_state", "watchlist_zip_code", "watchlist_phone", "watchlist_TIN",
]
FIELD_NAMES = ["entity_name", "address", "city", "state", "zip_code", "phone", "TIN", "other"]
ENTITY_TYPES = ["person", "organization", "location", "other", "unknown"]
SESSION_NAME = ".gold-workbench-session.json"


class ToolTip:
    """Native hover reminder for controls whose labels need an example."""
    def __init__(self, widget, text: str) -> None:
        self.widget, self.text, self.window = widget, text, None
        widget.bind("<Enter>", self.show, add=True)
        widget.bind("<Leave>", self.hide, add=True)
        widget.bind("<ButtonPress>", self.hide, add=True)

    def show(self, _event=None) -> None:
        if self.window or not self.text:
            return
        x, y = self.widget.winfo_rootx() + 10, self.widget.winfo_rooty() + self.widget.winfo_height() + 6
        self.window = tk.Toplevel(self.widget)
        self.window.wm_overrideredirect(True)
        self.window.wm_geometry(f"+{x}+{y}")
        ttk.Label(self.window, text=self.text, justify="left", wraplength=340, padding=8, relief="solid", borderwidth=1).pack()

    def hide(self, _event=None) -> None:
        if self.window:
            self.window.destroy()
            self.window = None


def read_client_rows(path: Path) -> list[dict]:
    """Read and validate the client export without changing its content."""
    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as source:
            rows = list(csv.DictReader(source))
    else:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook["01_Raw_client_output"] if "01_Raw_client_output" in workbook.sheetnames else workbook.active
        data = list(sheet.values)
        if not data:
            rows = []
        else:
            header = [str(value or "") for value in data[0]]
            rows = [
                dict(zip(header, ["" if value is None else str(value) for value in row]))
                for row in data[1:] if any(value is not None for value in row)
            ]
    missing = [column for column in REQUIRED if not rows or column not in rows[0]]
    if missing:
        raise ValueError("Missing client-export columns: " + ", ".join(missing))
    return [{**row, "client_row_id": row.get("client_row_id") or f"CR{index + 1}"} for index, row in enumerate(rows)]


class GoldStore:
    """Non-UI annotation model. It makes UI behaviour and exports testable."""
    def __init__(self) -> None:
        self.notes: dict[tuple[str, str], dict] = {}
        self.client_rows: list[dict] = []
        self.entities: list[dict] = []
        self.mentions: list[dict] = []
        self.fields: list[dict] = []
        self.categories: list[dict] = []
        self.statements: list[dict] = []
        self.uncertain: list[dict] = []
        self.watchlist_reviews: list[dict] = []
        self.reviewed_notes: set[tuple[str, str]] = set()
        self.sealed: set[str] = set()
        self._history: list[dict] = []
        self.note_folder: Path | None = None

    def load_note_paths(self, raw_paths: tuple[str, ...] | list[str]) -> None:
        for raw_path in raw_paths:
            path = Path(raw_path)
            parts = path.stem.split("_")
            claim = parts[0] if len(parts) > 1 else "UNASSIGNED"
            note_id = parts[-1]
            self.notes[(claim, note_id)] = {"filename": path.name, "text": path.read_text(encoding="utf-8")}
            self.note_folder = path.parent

    def claims(self) -> list[str]:
        return sorted({claim for claim, _ in self.notes})

    def note_ids(self, claim: str) -> list[str]:
        return sorted(note for note_claim, note in self.notes if note_claim == claim)

    def _snapshot(self) -> None:
        self._history.append(copy.deepcopy({
            "entities": self.entities, "mentions": self.mentions, "fields": self.fields, "categories": self.categories, "statements": self.statements,
            "uncertain": self.uncertain, "watchlist_reviews": self.watchlist_reviews,
            "reviewed_notes": self.reviewed_notes, "sealed": self.sealed,
        }))

    def undo(self) -> bool:
        if not self._history:
            return False
        state = self._history.pop()
        self.entities = state["entities"]
        self.mentions = state["mentions"]
        self.fields = state["fields"]
        self.categories = state["categories"]
        self.statements = state["statements"]
        self.uncertain = state["uncertain"]
        self.watchlist_reviews = state["watchlist_reviews"]
        self.reviewed_notes = state["reviewed_notes"]
        self.sealed = state["sealed"]
        return True

    def span(self, claim: str, note: str, start: int, end: int) -> dict:
        text = self.notes[(claim, note)]["text"]
        if start < 0 or end <= start or end > len(text):
            raise ValueError("The selected source span is outside this note.")
        return {"claim": claim, "note": note, "start": start, "end": end, "text": text[start:end]}

    def claim_entities(self, claim: str) -> list[dict]:
        return [entity for entity in self.entities if entity["claim"] == claim]

    def _next_id(self, prefix: str, rows: list[dict]) -> str:
        return f"{prefix}{len(rows) + 1}"

    def create_entity(self, span: dict, entity_type: str, name: str) -> dict:
        self._snapshot()
        entity = {"id": self._next_id("E", self.entities), "claim": span["claim"], "name": name.strip() or span["text"], "type": entity_type}
        self.entities.append(entity)
        self.mentions.append({**span, "id": self._next_id("M", self.mentions), "entity": entity["id"], "form": "name", "reason": "Reviewer identified named entity"})
        return entity

    def link_mention(self, span: dict, entity_id: str, form: str) -> dict:
        if entity_id not in {entity["id"] for entity in self.claim_entities(span["claim"])}:
            raise ValueError("Choose an entity in this claim.")
        self._snapshot()
        mention = {**span, "id": self._next_id("M", self.mentions), "entity": entity_id, "form": form, "reason": "Reviewer linked reference"}
        self.mentions.append(mention)
        return mention

    def add_field(self, span: dict, entity_id: str, field_name: str, value: str) -> dict:
        if entity_id not in {entity["id"] for entity in self.claim_entities(span["claim"])}:
            raise ValueError("Choose an entity in this claim.")
        self._snapshot()
        field = {**span, "id": self._next_id("GF", self.fields), "entity": entity_id, "field": field_name, "value": value.strip() or span["text"], "method": "manual_selection", "decision": "observed", "reason": "Reviewer selected source text"}
        self.fields.append(field)
        return field

    def add_context(self, span: dict, entity_id: str, optional_open_label: str = "") -> dict:
        if entity_id not in {entity["id"] for entity in self.claim_entities(span["claim"])}:
            raise ValueError("Choose an entity in this claim.")
        self._snapshot()
        record = {**span, "id": self._next_id("GC", self.categories), "entity": entity_id, "source_characterization": span["text"], "optional_open_label": optional_open_label.strip(), "method": "manual_selection", "decision": "observed", "reason": "Reviewer selected source characterization"}
        self.categories.append(record)
        return record

    def add_statement(self, span: dict, primary_entity: str, second_entity: str = "", optional_open_label: str = "", notes: str = "") -> dict:
        valid_entities = {entity["id"] for entity in self.claim_entities(span["claim"])}
        if primary_entity not in valid_entities:
            raise ValueError("Choose the entity that is the primary participant in this statement.")
        if second_entity and second_entity not in valid_entities:
            raise ValueError("Choose a second entity in this claim, or leave it blank.")
        self._snapshot()
        record = {**span, "id": self._next_id("GS", self.statements), "primary_entity": primary_entity, "second_entity": second_entity, "optional_open_label": optional_open_label.strip(), "notes": notes.strip(), "method": "manual_selection", "decision": "observed", "reason": "Reviewer selected statement wording"}
        self.statements.append(record)
        return record

    def add_uncertain(self, span: dict, reason: str) -> dict:
        self._snapshot()
        record = {**span, "id": self._next_id("U", self.uncertain), "reason": reason.strip() or "Reviewer marked uncertain"}
        self.uncertain.append(record)
        return record

    def set_note_reviewed(self, claim: str, note: str, reviewed: bool) -> None:
        self._snapshot()
        key = (claim, note)
        if reviewed:
            self.reviewed_notes.add(key)
        else:
            self.reviewed_notes.discard(key)

    def missing_reviewed_notes(self, claim: str) -> list[str]:
        return [note for note in self.note_ids(claim) if (claim, note) not in self.reviewed_notes]

    def seal(self, claim: str) -> None:
        missing = self.missing_reviewed_notes(claim)
        if missing:
            raise ValueError("Mark every note reviewed before sealing: " + ", ".join(missing))
        self._snapshot()
        self.sealed.add(claim)

    def watchlist_rows_for_note(self, claim: str, note: str) -> list[dict]:
        rows = []
        for row in self.client_rows:
            if row.get("claim_number") != claim or str(row.get("Entity_Watchlist_flag", "")).strip().lower() not in {"1", "true", "y", "yes"}:
                continue
            cited = {str(row.get("Exact_search_Note_ID", "")).strip(), str(row.get("GenAI_search_Note_ID", "")).strip()}
            if note in cited:
                rows.append(row)
        return rows

    def watchlist_review(self, client_row_id: str, note: str) -> dict | None:
        return next((review for review in self.watchlist_reviews if review["client_row_id"] == client_row_id and review["note_id"] == note), None)

    def save_watchlist_review(self, row: dict, note: str, decision: str, source_support: str, reason: str) -> None:
        self._snapshot()
        record = {"pair_review_id": self._next_id("WR", self.watchlist_reviews), "client_row_id": row["client_row_id"], "claim_number": row["claim_number"], "note_id": note, "similarity_score": row.get("GenAI_tok_sort_similarity", ""), "identity_decision": decision, "source_support": source_support, "reason": reason.strip(), "review_status": "complete"}
        self.watchlist_reviews = [review for review in self.watchlist_reviews if not (review["client_row_id"] == record["client_row_id"] and review["note_id"] == note)]
        self.watchlist_reviews.append(record)

    def to_json(self) -> dict:
        return {
            "notes": [{"claim": claim, "note": note, **record} for (claim, note), record in self.notes.items()],
            "client_rows": self.client_rows, "entities": self.entities, "mentions": self.mentions,
            "fields": self.fields, "categories": self.categories, "statements": self.statements, "uncertain": self.uncertain, "watchlist_reviews": self.watchlist_reviews,
            "reviewed_notes": [list(key) for key in sorted(self.reviewed_notes)], "sealed": sorted(self.sealed),
        }

    def load_json(self, data: dict) -> None:
        self.notes = {(row["claim"], row["note"]): {"filename": row["filename"], "text": row["text"]} for row in data.get("notes", [])}
        self.client_rows = data.get("client_rows", [])
        self.entities = data.get("entities", [])
        self.mentions = data.get("mentions", [])
        self.fields = data.get("fields", [])
        self.categories = data.get("categories", [])
        self.statements = data.get("statements", [])
        self.uncertain = data.get("uncertain", [])
        self.watchlist_reviews = data.get("watchlist_reviews", [])
        self.reviewed_notes = {tuple(key) for key in data.get("reviewed_notes", [])}
        self.sealed = set(data.get("sealed", []))
        self._history = []


def write_sheet(workbook: Workbook, name: str, headers: list[str], rows: list[list]) -> None:
    sheet = workbook.create_sheet(name)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True)
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(42, max(14, max(len(str(cell.value or "")) for cell in column) + 2))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def export_gold_workbook(store: GoldStore, target: str | Path) -> None:
    """Generate the same workbook contract as the Excel fallback, invisibly."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    write_sheet(workbook, "00_Instructions", ["Step", "Instruction"], [
        ["Review", "The reviewer selected exact source text in the offline workbench. These tabs are generated evidence records."],
        ["Gold isolation", "Client output is hidden while the reviewer creates blind Gold annotations."],
        ["Traceability", "Every mention and field value retains its exact UTF-8 source text and character offsets."],
    ])
    write_sheet(workbook, "01_Raw_client_output", ["client_row_id", *REQUIRED], [[row.get(column, "") for column in ["client_row_id", *REQUIRED]] for row in store.client_rows])
    write_sheet(workbook, "02_Note_inventory", ["claim_number", "note_id", "note_filename"], [[claim, note, record["filename"]] for (claim, note), record in sorted(store.notes.items())])
    write_sheet(workbook, "03_Assisted_entity_review", ["review_id", "client_row_id", "claim_number", "recordType", "decision", "entity_ref", "source_note_id", "source_match_method", "source_text", "reason"], [])
    write_sheet(workbook, "04_Entity_field_review", ["field_review_id", "client_row_id", "claim_number", "entity_ref", "field_name", "proposed_value", "source_note_id", "source_text", "source_match_method", "decision", "normalized_or_corrected_value", "reason"], [])
    write_sheet(workbook, "05_Category_review", ["review_id", "client_row_id", "claim_number", "proposed_category", "proposed_subcategory", "proposed_ner_tag", "decision", "reason"], [])
    entities = {(entity["claim"], entity["id"]): entity for entity in store.entities}
    write_sheet(workbook, "06_Coreference_links", ["link_id", "claim_number", "note_id", "start_character", "end_character", "exact_text", "entity_ref", "decision", "origin"], [[mention["id"], mention["claim"], mention["note"], mention["start"], mention["end"], mention["text"], mention["entity"], "resolved", "blind_gold"] for mention in store.mentions if mention["form"] in {"pronoun", "description", "alias"}])
    write_sheet(workbook, "07_Blind_gold_mentions", ["mention_id", "claim_number", "note_id", "start_character", "end_character", "exact_text", "entity_ref", "mention_form", "entity_type", "decision", "reason"], [[mention["id"], mention["claim"], mention["note"], mention["start"], mention["end"], mention["text"], mention["entity"], mention["form"], entities[(mention["claim"], mention["entity"])]["type"], "resolved", mention["reason"]] for mention in store.mentions])
    write_sheet(workbook, "08_Blind_gold_entity_fields", ["gold_field_id", "claim_number", "entity_ref", "field_name", "value", "note_id", "start_character", "end_character", "source_text", "source_match_method", "decision", "reason"], [[field["id"], field["claim"], field["entity"], field["field"], field["value"], field["note"], field["start"], field["end"], field["text"], field["method"], field["decision"], field["reason"]] for field in store.fields])
    write_sheet(workbook, "09_Phase1_pair_review", ["pair_review_id", "client_row_id", "claim_number", "similarity_score", "identity_decision", "source_support", "reason", "review_status"], [[review["pair_review_id"], review["client_row_id"], review["claim_number"], review["similarity_score"], review["identity_decision"], review["source_support"], review["reason"], review["review_status"]] for review in store.watchlist_reviews])
    write_sheet(workbook, "10_Phase1_summary", ["Metric", "Value"], [["Completed reviews", "=COUNTIF('09_Phase1_pair_review'!H2:H10000,\"complete\")"], ["Determinate decisions", "=COUNTIF('09_Phase1_pair_review'!E2:E10000,\"same_entity\")+COUNTIF('09_Phase1_pair_review'!E2:E10000,\"different_entity\")"], ["Confirmed same entity", "=COUNTIF('09_Phase1_pair_review'!E2:E10000,\"same_entity\")"], ["Returned-match precision", "=IF(B3=0,\"n.a.\",B4/B3)"]])
    write_sheet(workbook, "11_Phase2_candidate_input", ["candidate_id", "claim_number", "client_row_id", "candidate_status", "similarity_score", "watchlist_entity_id", "watchlist_entity_name", "note_id", "source_note_text"], [])
    write_sheet(workbook, "12_Phase2_pair_review", ["pair_review_id", "candidate_id", "claim_number", "candidate_status", "similarity_score", "identity_decision", "source_support", "reason", "review_status"], [])
    write_sheet(workbook, "13_Phase2_summary", ["Metric", "Value"], [["Completed reviews", "=COUNTIF('12_Phase2_pair_review'!I2:I10000,\"complete\")"], ["Determinate decisions", "=COUNTIF('12_Phase2_pair_review'!F2:F10000,\"same_entity\")+COUNTIF('12_Phase2_pair_review'!F2:F10000,\"different_entity\")"], ["Confirmed same entity", "=COUNTIF('12_Phase2_pair_review'!F2:F10000,\"same_entity\")"], ["Candidate precision", "=IF(B3=0,\"n.a.\",B4/B3)"]])
    write_sheet(workbook, "14_Uncertain_spans", ["uncertain_id", "claim_number", "note_id", "start_character", "end_character", "exact_text", "reason"], [[record["id"], record["claim"], record["note"], record["start"], record["end"], record["text"], record["reason"]] for record in store.uncertain])
    write_sheet(workbook, "15_Blind_gold_context", ["context_id", "claim_number", "entity_ref", "source_characterization", "optional_open_label", "note_id", "start_character", "end_character", "source_text", "source_match_method", "decision", "reason"], [[record["id"], record["claim"], record["entity"], record.get("source_characterization", record.get("category_text", record["text"])), record.get("optional_open_label", ""), record["note"], record["start"], record["end"], record["text"], record["method"], record["decision"], record["reason"]] for record in store.categories])
    write_sheet(workbook, "16_Blind_gold_statements", ["statement_id", "claim_number", "note_id", "start_character", "end_character", "predicate_source_text", "primary_entity_ref", "second_entity_ref_optional", "optional_open_action_or_relationship_label", "notes", "source_match_method", "decision", "reason"], [[record["id"], record["claim"], record["note"], record["start"], record["end"], record["text"], record["primary_entity"], record.get("second_entity", ""), record.get("optional_open_label", record.get("predicate", "")), record.get("notes", ""), record["method"], record["decision"], record["reason"]] for record in store.statements])
    workbook.save(target)


class GoldWorkbench(tk.Tk):
    def __init__(self) -> None:
        super().__init__()
        self.title("Gold Dataset Workbench")
        self.minsize(1120, 740)
        self.geometry("1420x900")
        self.store = GoldStore()
        self.claim = tk.StringVar()
        self.current_note = tk.StringVar()
        self.status = tk.StringVar(value="Start by loading a folder of UTF-8 note files. You only read and annotate; the workbook is created for you.")
        self.selection_label = tk.StringVar(value="1. Select words in the note.  2. Choose what those words mean.")
        self._selected_span: dict | None = None
        self._watchlist_window = None
        self._build()

    def _build(self) -> None:
        self.columnconfigure(0, weight=1)
        self.rowconfigure(3, weight=1)
        header = ttk.Frame(self, padding=(12, 10, 12, 4)); header.grid(row=0, column=0, sticky="ew"); header.columnconfigure(1, weight=1)
        ttk.Label(header, text="Gold Dataset Workbench", font=("Segoe UI", 16, "bold")).grid(row=0, column=0, sticky="w")
        ttk.Label(header, textvariable=self.status, wraplength=950, foreground="#24546f").grid(row=0, column=1, sticky="w", padx=(18, 0))
        controls = ttk.Frame(self, padding=(12, 4)); controls.grid(row=1, column=0, sticky="ew")
        controls_data = (("Load notes", self.load_notes, "Select every UTF-8 .txt note in the packet before reviewing."), ("Restore saved work", self.restore_session, "Resume a previous local annotation session."), ("Load client output", self.load_client, "Load the client export now; it remains hidden until Gold is sealed."), ("Undo last change", self.undo, "Reverse the last annotation, reviewed-note change, or watchlist decision."), ("Export workbook", self.export, "Create the traceable Excel workbook from your saved annotations."), ("Seal this claim", self.seal_claim, "Use only after every note is marked read. It unlocks post-Gold review."))
        for column, (text, command, hint) in enumerate(controls_data):
            button = ttk.Button(controls, text=text, command=command); button.grid(row=0, column=column, padx=(0 if column == 0 else 6, 0), pady=2); ToolTip(button, hint)
        selectors = ttk.Frame(self, padding=(12, 2, 12, 8)); selectors.grid(row=2, column=0, sticky="ew")
        ttk.Label(selectors, text="Claim").grid(row=0, column=0, sticky="w")
        self.claim_box = ttk.Combobox(selectors, textvariable=self.claim, state="readonly", width=18); self.claim_box.grid(row=0, column=1, padx=(5, 18)); self.claim_box.bind("<<ComboboxSelected>>", lambda _: self.refresh_claim())
        ttk.Label(selectors, text="Note").grid(row=0, column=2, sticky="w")
        self.note_box = ttk.Combobox(selectors, textvariable=self.current_note, state="readonly", width=18); self.note_box.grid(row=0, column=3, padx=5); self.note_box.bind("<<ComboboxSelected>>", lambda _: self.render_note())
        self.reviewed = tk.BooleanVar(value=False); reviewed = ttk.Checkbutton(selectors, text="I have read this whole note", variable=self.reviewed, command=self.set_note_reviewed); reviewed.grid(row=0, column=4, padx=(18, 0)); ToolTip(reviewed, "Check only after reading the complete note and recording every resolvable entity, reference, and stated field.")
        self.proposals = tk.BooleanVar(value=False); proposals = ttk.Checkbutton(selectors, text="Show client proposals after sealing", variable=self.proposals, command=self.show_proposals); proposals.grid(row=0, column=5, padx=(18, 0)); ToolTip(proposals, "Client records remain hidden until Gold is sealed, to protect independent annotation.")
        body = ttk.PanedWindow(self, orient="horizontal"); body.grid(row=3, column=0, sticky="nsew", padx=12, pady=(0, 12)); self.body = body
        left, centre, right = ttk.Frame(body, padding=8), ttk.Frame(body, padding=8), ttk.Frame(body, padding=8)
        body.add(left, weight=2); body.add(centre, weight=7); body.add(right, weight=3)
        self._build_entity_panel(left); self._build_reader(centre); self._build_proposal_panel(right)
        self.after(80, self.set_initial_panes)

    def set_initial_panes(self) -> None:
        """Give reading space priority while retaining draggable pane dividers."""
        width = self.body.winfo_width()
        if width < 800:
            return
        self.body.sashpos(0, 235)
        self.body.sashpos(1, max(760, width - 290))

    def _build_entity_panel(self, panel: ttk.Frame) -> None:
        ttk.Label(panel, text="Entities you have identified", font=("Segoe UI", 11, "bold")).pack(anchor="w")
        ttk.Label(panel, text="Click one to inspect its evidence.", foreground="#596b76").pack(anchor="w", pady=(0, 6))
        self.entity_list = tk.Listbox(panel, height=16, exportselection=False, activestyle="none"); self.entity_list.pack(fill="both", expand=True); self.entity_list.bind("<<ListboxSelect>>", lambda _: self.show_entity_evidence())
        ttk.Label(panel, text="Evidence for the selected entity", font=("Segoe UI", 10, "bold")).pack(anchor="w", pady=(10, 2))
        self.evidence = tk.Text(panel, height=13, wrap="word", state="disabled", background="#f7fafc", relief="solid", borderwidth=1); self.evidence.pack(fill="both", expand=True)

    def _build_reader(self, panel: ttk.Frame) -> None:
        action = ttk.LabelFrame(panel, text="After selecting text, choose one action", padding=7); action.pack(fill="x")
        for column in range(3): action.columnconfigure(column, weight=1)
        buttons = (("New entity", self.new_entity, "Use for a direct name: Maya Chen, Harbor Legal, or a newly introduced named organization."), ("Link reference", self.link_selection, "Use for a later name, alias, pronoun, or description that refers to an entity you already created: She, the firm, Northstar."), ("Record field", self.add_field, "Use only when the note directly states a value such as an address, city, phone number, ZIP, or TIN."), ("Record context", self.add_context, "Use for exact words that describe an entity: orthopedic surgeon, the medical provider, counsel. No category or subcategory bucket is required."), ("Record statement", self.add_statement, "Use for the exact wording of an action or relationship: called regarding the claim, represents, treated. Add a second entity only when the wording supports one."), ("Mark uncertain", self.mark_uncertain, "Use when the note does not support one resolution. Preserve the ambiguity instead of guessing."))
        for column, (text, command, hint) in enumerate(buttons):
            button = ttk.Button(action, text=text, command=command); button.grid(row=column // 3, column=column % 3, sticky="ew", padx=3, pady=2); ToolTip(button, hint)
        ttk.Label(panel, textvariable=self.selection_label, foreground="#24546f", wraplength=700).pack(anchor="w", pady=(8, 4))
        frame = ttk.Frame(panel); frame.pack(fill="both", expand=True)
        scrollbar = ttk.Scrollbar(frame, orient="vertical")
        self.reader = tk.Text(frame, wrap="word", font=("Segoe UI", 12), background="#ffffff", relief="solid", borderwidth=1, yscrollcommand=scrollbar.set, exportselection=False, state="disabled")
        scrollbar.configure(command=self.reader.yview); self.reader.pack(side="left", fill="both", expand=True); scrollbar.pack(side="right", fill="y")
        for name, color in (("entity", "#fff0a8"), ("reference", "#dff5ed"), ("field", "#e6efff"), ("context", "#eadcff"), ("statement", "#ffd9c7"), ("uncertain", "#ffe2b7")): self.reader.tag_configure(name, background=color)
        # Capture after Tk's Text binding has completed. This includes a span at 1.0.
        self.reader.bind("<B1-Motion>", lambda _event: self.after(1, self.capture_selection), add=True)
        self.reader.bind("<ButtonRelease-1>", lambda _event: self.after(1, self.capture_selection), add=True)
        self.reader.bind("<KeyRelease>", lambda _event: self.after(1, self.capture_selection), add=True)
        self.reader.bind("<<Selection>>", lambda _event: self.after(1, self.capture_selection), add=True)

    def _build_proposal_panel(self, panel: ttk.Frame) -> None:
        ttk.Label(panel, text="Client proposals", font=("Segoe UI", 11, "bold")).pack(anchor="w")
        ttk.Label(panel, text="These stay hidden until you seal your independent Gold review for the claim.", wraplength=250, foreground="#596b76").pack(anchor="w", pady=(0, 6))
        self.watchlist_notice = tk.StringVar(value="No current-note watchlist review is required.")
        ttk.Label(panel, textvariable=self.watchlist_notice, wraplength=260, foreground="#9a4d00").pack(anchor="w", pady=(4, 2))
        self.watchlist_button = ttk.Button(panel, text="Review watchlist matches for this note", command=self.open_watchlist_review, state="disabled")
        self.watchlist_button.pack(fill="x", pady=(0, 6))
        ToolTip(self.watchlist_button, "After Gold is sealed, review each flagged client watchlist match using the current source note and its Gold evidence.")
        self.client_view = tk.Text(panel, wrap="word", state="disabled", background="#f7fafc", relief="solid", borderwidth=1); self.client_view.pack(fill="both", expand=True)

    def load_notes(self) -> None:
        paths = filedialog.askopenfilenames(title="Select every UTF-8 .txt note in the packet", filetypes=[("UTF-8 notes", "*.txt")])
        if not paths: return
        try:
            self.store.load_note_paths(paths); self._refresh_claim_choices(); self.autosave()
            self.status.set(f"Loaded {len(self.store.notes)} notes across {len(self.store.claims())} claims. Read every note, select exact text, and annotate what it supports.")
        except UnicodeDecodeError: messagebox.showerror("Cannot read note", "A selected file is not UTF-8 text. Convert it to UTF-8 before loading it.")
        except Exception as exc: messagebox.showerror("Cannot load notes", str(exc))

    def _refresh_claim_choices(self) -> None:
        claims = self.store.claims(); self.claim_box["values"] = claims
        if claims: self.claim.set(claims[0]); self.refresh_claim()

    def load_client(self) -> None:
        raw = filedialog.askopenfilename(title="Select the client CSV or workbook", filetypes=[("Client data", "*.csv *.xlsx")])
        if not raw: return
        try:
            self.store.client_rows = read_client_rows(Path(raw)); self.autosave(); self.status.set("Client output loaded. It will remain hidden until independent Gold review is sealed for a claim."); self.show_proposals()
        except Exception as exc: messagebox.showerror("Client output", str(exc))

    def refresh_claim(self) -> None:
        note_ids = self.store.note_ids(self.claim.get()); self.note_box["values"] = note_ids
        if note_ids: self.current_note.set(note_ids[0])
        self.render_note(); self.refresh_entities(); self.show_proposals()

    def render_note(self) -> None:
        self._selected_span = None; self.selection_label.set("1. Select words in the note.  2. Choose what those words mean.")
        record = self.store.notes.get((self.claim.get(), self.current_note.get()))
        self.reader.configure(state="normal"); self.reader.delete("1.0", "end")
        if record:
            self.reader.insert("1.0", record["text"])
            for mention in [item for item in self.store.mentions if item["claim"] == self.claim.get() and item["note"] == self.current_note.get()]:
                self.reader.tag_add("entity" if mention["form"] in {"name", "alias"} else "reference", f"1.0+{mention['start']}c", f"1.0+{mention['end']}c")
            for field in [item for item in self.store.fields if item["claim"] == self.claim.get() and item["note"] == self.current_note.get()]: self.reader.tag_add("field", f"1.0+{field['start']}c", f"1.0+{field['end']}c")
            for category in [item for item in self.store.categories if item["claim"] == self.claim.get() and item["note"] == self.current_note.get()]: self.reader.tag_add("context", f"1.0+{category['start']}c", f"1.0+{category['end']}c")
            for statement in [item for item in self.store.statements if item["claim"] == self.claim.get() and item["note"] == self.current_note.get()]: self.reader.tag_add("statement", f"1.0+{statement['start']}c", f"1.0+{statement['end']}c")
            for item in [item for item in self.store.uncertain if item["claim"] == self.claim.get() and item["note"] == self.current_note.get()]: self.reader.tag_add("uncertain", f"1.0+{item['start']}c", f"1.0+{item['end']}c")
        self.reader.configure(state="disabled")
        self.reviewed.set((self.claim.get(), self.current_note.get()) in self.store.reviewed_notes); self.update_progress(); self.refresh_watchlist_status(); self.after(120, self.maybe_open_watchlist_review)

    def capture_selection(self, _event=None) -> None:
        ranges = self.reader.tag_ranges("sel")
        if len(ranges) != 2 or not self.claim.get() or not self.current_note.get(): return
        start_index, end_index = ranges; text = self.reader.get(start_index, end_index)
        if not text: return
        # Text.count returns None when its start and endpoint coincide at 1.0.
        # Reading the prefix gives a stable Python-character offset for every span,
        # including the first word, and matches the UTF-8 text held in GoldStore.
        start = len(self.reader.get("1.0", start_index))
        end = len(self.reader.get("1.0", end_index))
        self._selected_span = self.store.span(self.claim.get(), self.current_note.get(), start, end)
        shown = re.sub(r"\s+", " ", text).strip(); self.selection_label.set(f"Selected: “{shown[:130]}{'…' if len(shown) > 130 else ''}”. Choose an action above.")

    def selected(self) -> dict | None:
        # Button focus can arrive before a scheduled mouse-release callback.
        # Always take one final direct reading of the Text selection first.
        self.capture_selection()
        if self.claim.get() in self.store.sealed:
            messagebox.showinfo("Gold is sealed", "Gold annotations are locked for this claim. Complete the watchlist review using the current note, or undo the seal before correcting Gold.")
            return None
        if not self._selected_span or self._selected_span["claim"] != self.claim.get() or self._selected_span["note"] != self.current_note.get():
            messagebox.showinfo("Select source text", "Drag across the exact words in the note first. The selection will stay active while you click an action."); return None
        return self._selected_span

    def chooser(self, title: str, controls: list[tuple[str, list[str] | None, str]], on_save) -> None:
        window = tk.Toplevel(self); window.title(title); window.transient(self); window.grab_set(); window.resizable(False, False); values: dict[str, tk.StringVar] = {}
        for row, (label, choices, default) in enumerate(controls):
            ttk.Label(window, text=label).grid(row=row, column=0, sticky="w", padx=14, pady=7); value = tk.StringVar(value=default); values[label] = value
            widget = ttk.Combobox(window, textvariable=value, values=choices, width=42, state="readonly" if choices else "normal") if choices else ttk.Entry(window, textvariable=value, width=45)
            widget.grid(row=row, column=1, padx=(0, 14), pady=7)
            if row == 0: widget.focus_set()
        def save() -> None:
            try: on_save(values)
            except Exception as exc: messagebox.showerror(title, str(exc), parent=window); return
            window.destroy()
        buttons = ttk.Frame(window); buttons.grid(row=len(controls), column=0, columnspan=2, sticky="e", padx=14, pady=(4, 12)); ttk.Button(buttons, text="Cancel", command=window.destroy).pack(side="right", padx=(6, 0)); ttk.Button(buttons, text="Save annotation", command=save).pack(side="right")

    def new_entity(self) -> None:
        span = self.selected()
        if span: self.chooser("Create named entity", [("Entity type", ENTITY_TYPES, "person"), ("Display name", None, span["text"])], lambda values: self._create_entity(span, values))

    def _create_entity(self, span: dict, values: dict[str, tk.StringVar]) -> None:
        self.store.create_entity(span, values["Entity type"].get(), values["Display name"].get()); self.after_annotation("Saved a named entity and its source wording.")

    def _entity_choices(self) -> tuple[list[str], list[dict]]:
        entities = self.store.claim_entities(self.claim.get()); return [f"{entity['id']} · {entity['name']}" for entity in entities], entities

    def link_selection(self) -> None:
        span = self.selected(); labels, entities = self._entity_choices()
        if not span: return
        if not entities: messagebox.showinfo("Create an entity first", "First select a named entity and use New named entity. Then you can link a pronoun, alias, or description to it."); return
        current = self.entity_list.curselection(); default = labels[current[0]] if current else labels[0]
        self.chooser("Link this reference", [("Entity", labels, default), ("Reference form", ["pronoun", "description", "alias", "name"], "pronoun")], lambda values: self._link(span, values, labels, entities))

    def _link(self, span: dict, values: dict[str, tk.StringVar], labels: list[str], entities: list[dict]) -> None:
        entity = entities[labels.index(values["Entity"].get())]; self.store.link_mention(span, entity["id"], values["Reference form"].get()); self.after_annotation("Linked this wording to the selected entity.")

    def add_field(self) -> None:
        span = self.selected(); labels, entities = self._entity_choices()
        if not span: return
        if not entities: messagebox.showinfo("Create an entity first", "Identify the named entity before recording a field value for it."); return
        current = self.entity_list.curselection(); default = labels[current[0]] if current else labels[0]
        self.chooser("Record a source-supported field", [("Entity", labels, default), ("Field", FIELD_NAMES, self.guess_field(span["text"])), ("Value", None, span["text"])], lambda values: self._field(span, values, labels, entities))

    def guess_field(self, text: str) -> str:
        compact = text.strip()
        if re.search(r"\b\d{3}-\d{2}-\d{4}\b|\bTIN\b", compact, re.I): return "TIN"
        if re.search(r"\(?\d{3}\)?[-.\s]\d{3}[-.\s]\d{4}", compact): return "phone"
        if re.fullmatch(r"\d{5}(?:-\d{4})?", compact): return "zip_code"
        if re.fullmatch(r"[A-Z]{2}", compact): return "state"
        if re.search(r"\b(?:Street|St\.|Avenue|Ave\.|Road|Rd\.|Boulevard|Blvd\.|Suite|#)\b", compact, re.I): return "address"
        return "other"

    def _field(self, span: dict, values: dict[str, tk.StringVar], labels: list[str], entities: list[dict]) -> None:
        entity = entities[labels.index(values["Entity"].get())]; self.store.add_field(span, entity["id"], values["Field"].get(), values["Value"].get()); self.after_annotation("Recorded the observed field and its exact source wording.")

    def add_context(self) -> None:
        span = self.selected(); labels, entities = self._entity_choices()
        if not span: return
        if not entities: messagebox.showinfo("Create an entity first", "Identify the entity before recording source context for it."); return
        current = self.entity_list.curselection()
        default = labels[current[0]] if current else labels[0]
        self.chooser("Record source context", [("Entity", labels, default), ("Optional open label", None, "")], lambda values: self._context(span, values, labels, entities))

    def _context(self, span: dict, values: dict[str, tk.StringVar], labels: list[str], entities: list[dict]) -> None:
        entity = entities[labels.index(values["Entity"].get())]
        self.store.add_context(span, entity["id"], values["Optional open label"].get())
        self.after_annotation("Recorded the exact source context for the chosen entity.")

    def add_statement(self) -> None:
        span = self.selected(); labels, entities = self._entity_choices()
        if not span: return
        if not entities: messagebox.showinfo("Create participants first", "Create the named participants before recording a statement or relationship."); return
        current = self.entity_list.curselection(); default = labels[current[0]] if current else labels[0]
        self.chooser("Record statement wording", [("Primary entity", labels, default), ("Second entity (optional)", ["", *labels], ""), ("Optional open label", None, ""), ("Notes (optional)", None, "")], lambda values: self._statement(span, values, labels, entities))

    def _statement(self, span: dict, values: dict[str, tk.StringVar], labels: list[str], entities: list[dict]) -> None:
        primary = entities[labels.index(values["Primary entity"].get())]
        second_label = values["Second entity (optional)"].get()
        second = entities[labels.index(second_label)] if second_label else None
        self.store.add_statement(span, primary["id"], second["id"] if second else "", values["Optional open label"].get(), values["Notes (optional)"].get())
        self.after_annotation("Recorded the exact statement wording and its linked entity or entities.")

    def mark_uncertain(self) -> None:
        span = self.selected()
        if not span: return
        reason = simpledialog.askstring("Mark uncertain", "Why can this wording not be resolved now?", parent=self)
        if reason is not None: self.store.add_uncertain(span, reason); self.after_annotation("Saved this wording as uncertain. It will not be silently resolved.")

    def after_annotation(self, message: str) -> None:
        self._selected_span = None; self.selection_label.set("Saved. Continue reading, or select the next relevant words."); self.refresh_entities(); self.render_note(); self.autosave(); self.status.set(message + " Your work was saved locally.")

    def refresh_entities(self) -> None:
        current = self.entity_list.curselection(); self.entity_list.delete(0, "end")
        for entity in self.store.claim_entities(self.claim.get()):
            count = sum(item["entity"] == entity["id"] for item in self.store.mentions); self.entity_list.insert("end", f"{entity['id']} · {entity['name']}  ({entity['type']}; {count} mentions)")
        if current and current[0] < self.entity_list.size(): self.entity_list.selection_set(current[0]); self.show_entity_evidence()

    def show_entity_evidence(self) -> None:
        self.evidence.configure(state="normal"); self.evidence.delete("1.0", "end"); selection = self.entity_list.curselection()
        if not selection: self.evidence.configure(state="disabled"); return
        entity = self.store.claim_entities(self.claim.get())[selection[0]]; self.evidence.insert("end", f"{entity['id']} · {entity['name']}\nType: {entity['type']}\n\nMENTIONS\n")
        for mention in [item for item in self.store.mentions if item["entity"] == entity["id"]]: self.evidence.insert("end", f"{mention['note']}: “{mention['text']}” — {mention['form']}\n")
        self.evidence.insert("end", "\nOBSERVED FIELDS\n")
        for field in [item for item in self.store.fields if item["entity"] == entity["id"]]: self.evidence.insert("end", f"{field['field']}: {field['value']}\n  {field['note']}: “{field['text']}”\n")
        self.evidence.insert("end", "\nSOURCE CONTEXT\n")
        for category in [item for item in self.store.categories if item["entity"] == entity["id"]]:
            label = category.get("optional_open_label", "")
            self.evidence.insert("end", f"{category.get('source_characterization', category.get('category_text', category['text']))}{f' → {label}' if label else ''}\n  {category['note']}: “{category['text']}”\n")
        self.evidence.insert("end", "\nSTATEMENTS\n")
        for statement in [item for item in self.store.statements if item["primary_entity"] == entity["id"] or item.get("second_entity") == entity["id"]]:
            label = statement.get("optional_open_label", "")
            self.evidence.insert("end", f"{statement['text']}{f' → {label}' if label else ''}\n  {statement['note']}: “{statement['text']}”\n")
        self.evidence.configure(state="disabled")

    def set_note_reviewed(self) -> None:
        if not self.claim.get() or not self.current_note.get(): return
        if self.claim.get() in self.store.sealed:
            self.reviewed.set(True)
            messagebox.showinfo("Gold is sealed", "Undo the seal before changing note-completion state.")
            return
        self.store.set_note_reviewed(self.claim.get(), self.current_note.get(), self.reviewed.get()); self.autosave(); self.update_progress(); self.refresh_watchlist_status()

    def update_progress(self) -> None:
        claim = self.claim.get()
        if not claim: return
        total = len(self.store.note_ids(claim)); done = total - len(self.store.missing_reviewed_notes(claim)); entities = len(self.store.claim_entities(claim)); mentions = sum(item["claim"] == claim for item in self.store.mentions); fields = sum(item["claim"] == claim for item in self.store.fields); contexts = sum(item["claim"] == claim for item in self.store.categories); statements = sum(item["claim"] == claim for item in self.store.statements); sealed = " · Gold sealed" if claim in self.store.sealed else ""
        self.status.set(f"{claim}: {done}/{total} notes read · {entities} entities · {mentions} mentions · {fields} fields · {contexts} context spans · {statements} statements{sealed}")

    def undo(self) -> None:
        if not self.store.undo(): messagebox.showinfo("Nothing to undo", "There is no saved change to undo in this session."); return
        self.autosave(); self.refresh_claim(); self.status.set("Undid the last annotation or review-state change. Your updated session was saved.")

    def seal_claim(self) -> None:
        claim = self.claim.get()
        if not claim: return
        missing = self.store.missing_reviewed_notes(claim)
        if missing: messagebox.showwarning("Finish reading the claim", "Mark every note as read before sealing Gold. Still to read: " + ", ".join(missing)); return
        if not messagebox.askyesno("Seal Gold", "You are about to make client proposals available for this claim. Confirm that your annotations were made from the notes, without using the client output."): return
        try: self.store.seal(claim)
        except ValueError as exc: messagebox.showwarning("Cannot seal", str(exc)); return
        self.autosave(); self.update_progress(); self.show_proposals(); self.refresh_watchlist_status(); self.after(120, self.maybe_open_watchlist_review)

    def show_proposals(self) -> None:
        self.client_view.configure(state="normal"); self.client_view.delete("1.0", "end"); claim = self.claim.get()
        if not self.proposals.get(): self.client_view.insert("end", "Client output is hidden while you create Gold annotations.\n\nOnce every note is marked read, seal the claim. Then you may display these records for a separate reconciliation pass.")
        elif claim not in self.store.sealed: self.client_view.insert("end", "Seal your Gold review for this claim before displaying client records.")
        else:
            rows = [row for row in self.store.client_rows if row["claim_number"] == claim]
            if not rows: self.client_view.insert("end", "No client records are loaded for this claim.")
            for row in rows:
                self.client_view.insert("end", f"{row['client_row_id']} · {row['recordType']}\n{row['entity_name']}\nCategory: {row['entity_category_name']}\n")
                for label, value in (("Address", row.get("entity_address", "")), ("Phone", row.get("entity_phone", "")), ("TIN", row.get("entity_TIN", "")), ("Watchlist", row.get("Watchlist_entity_name", ""))):
                    if value: self.client_view.insert("end", f"{label}: {value}\n")
                self.client_view.insert("end", "\n")
        self.client_view.configure(state="disabled")

    def refresh_watchlist_status(self) -> None:
        if not hasattr(self, "watchlist_button"):
            return
        claim, note = self.claim.get(), self.current_note.get()
        rows = self.store.watchlist_rows_for_note(claim, note) if claim and note else []
        if not rows:
            self.watchlist_notice.set("No current-note watchlist review is required.")
            self.watchlist_button.configure(state="disabled")
            return
        if claim not in self.store.sealed:
            self.watchlist_notice.set(f"{len(rows)} watchlist match(es) cite this note. Additional review unlocks after Gold is sealed.")
            self.watchlist_button.configure(state="disabled")
            return
        if (claim, note) not in self.store.reviewed_notes:
            self.watchlist_notice.set(f"{len(rows)} watchlist match(es) cite this note. Mark the note read before review.")
            self.watchlist_button.configure(state="disabled")
            return
        completed = sum(self.store.watchlist_review(row["client_row_id"], note) is not None for row in rows)
        if completed == len(rows):
            self.watchlist_notice.set(f"Watchlist review complete for this note ({completed}/{len(rows)}).")
        else:
            self.watchlist_notice.set(f"Additional step required: review {len(rows) - completed} watchlist match(es) for this note.")
        self.watchlist_button.configure(state="normal")

    def maybe_open_watchlist_review(self) -> None:
        claim, note = self.claim.get(), self.current_note.get()
        if not claim or not note or claim not in self.store.sealed or (claim, note) not in self.store.reviewed_notes:
            return
        rows = self.store.watchlist_rows_for_note(claim, note)
        if rows and any(self.store.watchlist_review(row["client_row_id"], note) is None for row in rows):
            self.open_watchlist_review()

    def open_watchlist_review(self) -> None:
        claim, note = self.claim.get(), self.current_note.get()
        if not claim or not note:
            return
        if claim not in self.store.sealed:
            messagebox.showinfo("Seal Gold first", "Watchlist validation is a post-Gold step. Seal this claim after every note is marked read.")
            return
        rows = self.store.watchlist_rows_for_note(claim, note)
        if not rows:
            messagebox.showinfo("No watchlist match", "No flagged client watchlist match explicitly cites this note.")
            return
        if self._watchlist_window and self._watchlist_window.winfo_exists():
            self._watchlist_window.lift(); self._watchlist_window.focus_force(); return
        window = tk.Toplevel(self); self._watchlist_window = window
        window.title(f"Watchlist review · {claim} · {note}")
        window.geometry("540x670")
        window.minsize(480, 540)
        window.protocol("WM_DELETE_WINDOW", lambda: (window.destroy(), setattr(self, "_watchlist_window", None)))
        ttk.Label(window, text="Review this returned watchlist match with the source note still visible behind this window.", wraplength=500, padding=(12, 10), foreground="#24546f").pack(anchor="w")
        index = tk.IntVar(value=0)
        position = tk.StringVar()
        detail = tk.Text(window, height=20, wrap="word", state="disabled", background="#f7fafc", relief="solid", borderwidth=1)
        detail.pack(fill="both", expand=True, padx=12, pady=(0, 8))
        form = ttk.Frame(window, padding=(12, 0, 12, 8)); form.pack(fill="x")
        ttk.Label(form, textvariable=position, font=("Segoe UI", 10, "bold")).grid(row=0, column=0, columnspan=3, sticky="w", pady=(0, 5))
        decision = tk.StringVar(value="same_entity")
        for column, (label, value) in enumerate((("Same entity", "same_entity"), ("Different entity", "different_entity"), ("Insufficient evidence", "insufficient_evidence"))):
            ttk.Radiobutton(form, text=label, variable=decision, value=value).grid(row=1, column=column, sticky="w", padx=(0, 10))
        ttk.Label(form, text="Source support").grid(row=2, column=0, sticky="w", pady=(8, 2))
        support = tk.StringVar(value="direct_note_support")
        ttk.Combobox(form, textvariable=support, values=["direct_note_support", "partial_note_support", "no_note_support", "ambiguous_note_support"], state="readonly", width=30).grid(row=2, column=1, columnspan=2, sticky="w", pady=(8, 2))
        ttk.Label(form, text="Reason").grid(row=3, column=0, sticky="nw", pady=(8, 2))
        reason = tk.Text(form, height=3, wrap="word", width=46)
        reason.grid(row=3, column=1, columnspan=2, sticky="ew", pady=(8, 2))

        def render() -> None:
            row = rows[index.get()]
            prior = self.store.watchlist_review(row["client_row_id"], note)
            position.set(f"Match {index.get() + 1} of {len(rows)} · {row['client_row_id']}")
            detail.configure(state="normal"); detail.delete("1.0", "end")
            detail.insert("end", "SYSTEM ENTITY FOUND IN THIS NOTE\n")
            detail.insert("end", f"Name: {row.get('entity_name', '')}\nCategory: {row.get('entity_category_name', '')}\nAddress: {row.get('entity_address', '')}\nCity/state/ZIP: {row.get('entity_city', '')}, {row.get('entity_state', '')} {row.get('entity_zip_code', '')}\nPhone: {row.get('entity_phone', '')}\nTIN: {row.get('entity_TIN', '')}\n\n")
            detail.insert("end", "WATCHLIST CANDIDATE\n")
            detail.insert("end", f"Name: {row.get('Watchlist_entity_name', '')}\nAddress: {row.get('watchlist_address', '')}\nCity/state/ZIP: {row.get('watchlist_city', '')}, {row.get('watchlist_state', '')} {row.get('watchlist_zip_code', '')}\nPhone: {row.get('watchlist_phone', '')}\nTIN: {row.get('watchlist_TIN', '')}\nSimilarity score: {row.get('GenAI_tok_sort_similarity', '')}\n\n")
            detail.insert("end", "Decide whether the entity identified in the current source note is the watchlist candidate. Use the full note and the Gold evidence, not similarity alone.")
            detail.configure(state="disabled")
            decision.set(prior["identity_decision"] if prior else "same_entity")
            support.set(prior["source_support"] if prior else "direct_note_support")
            reason.delete("1.0", "end"); reason.insert("1.0", prior["reason"] if prior else "")

        def save() -> None:
            if not reason.get("1.0", "end-1c").strip():
                messagebox.showwarning("Explain the decision", "Record why the note evidence supports this identity decision.", parent=window); return
            self.store.save_watchlist_review(rows[index.get()], note, decision.get(), support.get(), reason.get("1.0", "end-1c"))
            self.autosave(); self.refresh_watchlist_status(); self.status.set("Saved the watchlist decision for this note. It will export in Phase 1 pair review.")

        navigation = ttk.Frame(window, padding=(12, 0, 12, 12)); navigation.pack(fill="x")
        ttk.Button(navigation, text="Previous", command=lambda: (index.set(max(0, index.get() - 1)), render())).pack(side="left")
        ttk.Button(navigation, text="Next", command=lambda: (index.set(min(len(rows) - 1, index.get() + 1)), render())).pack(side="left", padx=6)
        ttk.Button(navigation, text="Save decision", command=save).pack(side="right")
        render()

    def autosave(self) -> None:
        if self.store.note_folder: (self.store.note_folder / SESSION_NAME).write_text(json.dumps(self.store.to_json(), ensure_ascii=False, indent=2), encoding="utf-8")

    def restore_session(self) -> None:
        raw = filedialog.askopenfilename(title="Select a saved Gold Workbench session", initialfile=SESSION_NAME, filetypes=[("Gold Workbench session", "*.json")])
        if not raw: return
        try:
            self.store.load_json(json.loads(Path(raw).read_text(encoding="utf-8"))); self.store.note_folder = Path(raw).parent; self._refresh_claim_choices(); self.status.set("Restored your local annotation session. The source notes and all annotations are available again.")
        except Exception as exc: messagebox.showerror("Cannot restore session", str(exc))

    def export(self) -> None:
        if not self.store.notes: messagebox.showwarning("Load notes", "Load the note packet first."); return
        target = filedialog.asksaveasfilename(defaultextension=".xlsx", initialfile="gold-dataset.xlsx", filetypes=[("Excel workbook", "*.xlsx")])
        if not target: return
        try: export_gold_workbook(self.store, target); self.status.set("Exported the traceable Gold workbook. Your offline workbench session remains saved separately.")
        except Exception as exc: messagebox.showerror("Cannot export", str(exc))


def self_test() -> None:
    """Automated regression checks runnable without Tk or a desktop display."""
    base = Path(__file__).parent / "sample-data"; store = GoldStore(); store.load_note_paths([str(path) for path in sorted(base.glob("C104_*.txt"))])
    assert len(store.notes) == 2 and store.claims() == ["C104"]
    source = store.notes[("C104", "N01")]["text"]; start = source.index("Maya Chen"); maya = store.create_entity(store.span("C104", "N01", start, start + len("Maya Chen")), "person", "Maya Chen")
    phone = source.index("555-0199"); store.add_field(store.span("C104", "N01", phone, phone + len("555-0199")), maya["id"], "phone", "555-0199")
    store.add_context(store.span("C104", "N01", start, start + len("Maya Chen")), maya["id"], "claimant")
    called = source.index("called about her claim"); store.add_statement(store.span("C104", "N01", called, called + len("called about her claim")), maya["id"], optional_open_label="called")
    pronoun = source.index("She"); store.link_mention(store.span("C104", "N01", pronoun, pronoun + len("She")), maya["id"], "pronoun"); store.add_uncertain(store.span("C104", "N01", pronoun, pronoun + len("She")), "Test unresolved alternate reading")
    assert len(store.mentions) == 2 and len(store.fields) == 1 and len(store.categories) == 1 and len(store.statements) == 1 and len(store.uncertain) == 1
    store.set_note_reviewed("C104", "N01", True)
    try: store.seal("C104"); raise AssertionError("seal should require every note")
    except ValueError: pass
    store.set_note_reviewed("C104", "N02", True); store.seal("C104"); store.client_rows = read_client_rows(base / "client-export.csv")
    with tempfile.TemporaryDirectory() as temporary:
        output = Path(temporary) / "gold.xlsx"; export_gold_workbook(store, output); workbook = load_workbook(output, data_only=False)
        expected = ["00_Instructions", "01_Raw_client_output", "02_Note_inventory", "03_Assisted_entity_review", "04_Entity_field_review", "05_Category_review", "06_Coreference_links", "07_Blind_gold_mentions", "08_Blind_gold_entity_fields", "09_Phase1_pair_review", "10_Phase1_summary", "11_Phase2_candidate_input", "12_Phase2_pair_review", "13_Phase2_summary", "14_Uncertain_spans", "15_Blind_gold_context", "16_Blind_gold_statements"]
        assert workbook.sheetnames == expected; mention = list(workbook["07_Blind_gold_mentions"].values)[1]; assert mention[5] == "Maya Chen" and mention[3:5] == (start, start + len("Maya Chen")); assert len(list(workbook["06_Coreference_links"].values)) == 2; assert list(workbook["15_Blind_gold_context"].values)[1][3] == "Maya Chen"; assert list(workbook["16_Blind_gold_statements"].values)[1][8] == "called"
    n01_watchlist = store.watchlist_rows_for_note("C104", "N01"); n02_watchlist = store.watchlist_rows_for_note("C104", "N02")
    assert len(n01_watchlist) == 1 and len(n02_watchlist) == 2
    store.save_watchlist_review(n01_watchlist[0], "N01", "same_entity", "direct_note_support", "Name and phone agree in the note.")
    with tempfile.TemporaryDirectory() as temporary:
        output = Path(temporary) / "watchlist.xlsx"; export_gold_workbook(store, output); workbook = load_workbook(output, data_only=False)
        phase_one = list(workbook["09_Phase1_pair_review"].values)
        assert len(phase_one) == 2 and phase_one[1][4] == "same_entity"
    restored = GoldStore(); restored.load_json(store.to_json()); assert restored.mentions[0]["text"] == "Maya Chen"
    unicode_store = GoldStore(); unicode_store.notes[("U1", "N1")] = {"filename": "U1_N1.txt", "text": "Zoë met 李雷."}
    unicode_start = unicode_store.notes[("U1", "N1")]["text"].index("李雷")
    assert unicode_store.span("U1", "N1", unicode_start, unicode_start + 2)["text"] == "李雷"
    assert store.undo() and not store.watchlist_reviews
    assert store.undo() and "C104" not in store.sealed
    stress = GoldStore(); stress_base = base / "stress-packet"; stress.load_note_paths([str(path) for path in sorted((stress_base / "notes").glob("*.txt"))])
    assert len(stress.notes) == 60 and len(stress.claims()) == 6 and all(len(stress.note_ids(claim)) == 10 for claim in stress.claims())
    stress.client_rows = read_client_rows(stress_base / "client-export.csv")
    assert len(stress.client_rows) == 12 and {row["claim_number"] for row in stress.client_rows} == set(stress.claims())
    print("PASS: Gold Workbench self-test completed (source spans, field evidence, coreference, seal gate, export contract, undo).")


if __name__ == "__main__":
    self_test() if "--self-test" in sys.argv else GoldWorkbench().mainloop()
